#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kretschmann-geometry SPRM reflectivity and propagation calculator.

The program (A) compares p-polarized R(theta) curves for several metals at
fixed film thickness, (B) scans thickness to report theta_min(d), R_min(d),
and propagation length Lx(d), and (C) exports figures and an XLSX workbook
with matching JSON metadata.

The ideal optical stack is semi-infinite prism / metal film of thickness d /
semi-infinite dielectric. With time convention exp(-i omega t), passive
absorption has Im(epsilon) > 0. The multilayer Fresnel calculation uses

    k_x   = k0 sqrt(eps_p) sin(theta),
    k_z,j = sqrt(eps_j k0^2 - k_x^2), with Im(k_z,j) >= 0,
    r_ij  = (eps_j k_z,i - eps_i k_z,j)
          / (eps_j k_z,i + eps_i k_z,j),
    r     = (r01 + r12 exp(2i k_z,1 d))
          / (1 + r01 r12 exp(2i k_z,1 d)),  R = |r|^2.

Theta is the internal incidence angle at the metal interface, measured on
the prism side. For a triangular prism, a reported external-face angle needs
an additional refraction and geometry conversion. The critical angle is
theta_c = asin(n_dielectric/n_prism); the refined SPR search is restricted
to [theta_c, 90 degrees) to avoid subcritical Fresnel minima.

The intrinsic, semi-infinite bound-SPP propagation length is

    k_spp = k0 sqrt(eps_m eps_d/(eps_m + eps_d)),
    Lx_int = 1/(2 Im(k_spp)).

For a finite Kretschmann film, radiation leakage into the prism gives
Gamma_tot(d) = Gamma_int + Gamma_rad(d) and Lx = 1/(2 Gamma_tot).
The complex pole is found with the numerically stable entire equation

    F(kx) = D01 D12 + N01 N12 exp(2i k_z,1 d) = 0,
    N_ij = eps_j k_z,i - eps_i k_z,j,
    D_ij = eps_j k_z,i + eps_i k_z,j.

The leaky sheet has Im(k_z,prism) < 0; the bound dielectric sheet has
Im(k_z,dielectric) > 0. Thickness continuation starts from a thick-film
bound-SPP solution, and the physical constraint Lx_leaky <= Lx_int is used.
For Au/water/SF10 at 632.8 nm, example ratios Lx/Lx_int are 0.253 at
30 nm, 0.546 at 48 nm, and approximately 1 at 500 nm.

Literature checks include the Nat. Rev. Methods Primers SPR Box 3 stack
(prism eps = 3.084; Au eps = -12.08 + 1.52i; d = 43.1 nm; solution
eps = 1.773; lambda = 632.8 nm): calculated theta_min = 55.339 degrees,
compared with an estimated 55.3-55.5 degrees in the cited figure.
The intrinsic Au/water length at 632.8 nm is 3.86 micrometers, near the
roughly 3 micrometers given in that source.

Optical constants are based on measured n and k tables from
refractiveindex.info (CC0), with PCHIP interpolation; these include Johnson
and Christy for Au, Ag, Cu, and Pd, and a Brendel-Bormann fit by Rakic et al.
for Pt. Prism dispersion uses SCHOTT Sellmeier coefficients and published
fused-silica and sapphire formulas. Water uses Daimon and Masumura at
specified temperatures; standard air uses Ciddor. A Lorentz-Drude fit is
not used for Au or Cu because of large interband-region errors near
632.8 nm. Source citations and wavelength limits are recorded below.

Interpretation limits
---------------------
* A deep R_min alone does not establish a usable SPP. Even Pt can show
  near-zero reflectivity by critical coupling in a thin film while its dip
  remains broad. Use FWHM, |Re(eps)|/Im(eps), and Lx/lambda_spp to assess
  whether the mode propagates.
* Metal interband absorption matters strongly at 532 nm, particularly for
  Au and Cu. Ag is generally the better of these modeled materials there.
  Optical constants for Cu vary substantially with film preparation and
  oxidation; select a source appropriate to the deposited film.
* Use the experimental water temperature and standard-air refractive index
  rather than replacing air with vacuum. The ideal three-layer model omits
  adhesion layers and roughness, so fitted eps and d are effective values.

The TMM kernel was compared with the independently implemented MIT-licensed
tmm package (S. J. Byrnes, arXiv:1603.02720): across five metals, four
wavelengths, three thicknesses, and 400 angles, the maximum reported
absolute difference was 6.97e-14.

Dependencies: NumPy, SciPy, Matplotlib, openpyxl, and Tkinter.
"""


def _configure_tcltk_environment():
    """Ensure Tkinter can find Tcl/Tk runtime files on Windows."""
    import os
    import sys
    from pathlib import Path

    if os.environ.get("TCL_LIBRARY") and os.environ.get("TK_LIBRARY"):
        return

    def _set_if_exists(env_name, path):
        if path and path.exists():
            os.environ.setdefault(env_name, str(path))
            return True
        return False

    roots = [Path(sys.base_prefix), Path(sys.prefix), Path(sys.executable).resolve().parent.parent]
    candidates = []
    for root in roots:
        if not root:
            continue
        candidates.extend([
            root / "tcl" / "tcl8.6",
            root / "tcl" / "tk8.6",
            root / "Lib" / "tcl8.6",
            root / "Lib" / "tk8.6",
            root / "tcl8.6",
            root / "tk8.6",
        ])

    for candidate in candidates:
        if candidate.name == "tcl8.6":
            _set_if_exists("TCL_LIBRARY", candidate)
        elif candidate.name == "tk8.6":
            _set_if_exists("TK_LIBRARY", candidate)
        if os.environ.get("TCL_LIBRARY") and os.environ.get("TK_LIBRARY"):
            break


_configure_tcltk_environment()

import datetime
import json
from pathlib import Path
import numpy as np
from scipy.interpolate import PchipInterpolator
from scipy.optimize import minimize_scalar, root

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import matplotlib
matplotlib.use("TkAgg")
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from matplotlib.lines import Line2D

# openpyxl is optional for calculation and plotting; an export request reports how to install it if missing.
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - only for environments without openpyxl
    Workbook = None

# Use English plot labels to avoid missing CJK glyphs.
matplotlib.rcParams.update({
    "font.size": 9, "axes.labelsize": 10, "axes.titlesize": 10,
    "legend.fontsize": 8, "xtick.labelsize": 9, "ytick.labelsize": 9,
    "axes.unicode_minus": True, "figure.autolayout": False,
})

# ============================================================================
# ==== Built-in dispersion data extracted from refractiveindex.info (CC0), 380-1100 nm ====
# Format: METAL -> {source_label: (lambda_nm[], n[], k[])}
# The first source for each metal is the default; all were checked against literature values at 632.8 nm.
NK_DATA = {
    "Au": {
        "Johnson & Christy 1972": (
            [381.5, 397.4, 413.3, 430.5, 450.9, 471.4, 495.9, 520.9, 548.6, 582.1, 616.8, 659.5,
            704.5, 756.0, 821.1, 892.0, 984.0, 1088.0],
            [1.4600, 1.4700, 1.4600, 1.4500, 1.3800, 1.3100, 1.0400, 0.6200, 0.4300, 0.2900,
            0.2100, 0.1400, 0.1300, 0.1400, 0.1600, 0.1700, 0.2200, 0.2700],
            [1.9330, 1.9520, 1.9580, 1.9480, 1.9140, 1.8490, 1.8330, 2.0810, 2.4550, 2.8630,
            3.2720, 3.6970, 4.1030, 4.5420, 5.0830, 5.6630, 6.3500, 7.1500],
        ),
        "Olmon 2012 (evaporated)": (
            [380.0, 390.0, 400.0, 410.0, 420.0, 430.0, 440.0, 450.0, 460.0, 470.0, 480.0, 490.0,
            500.0, 510.0, 520.0, 530.0, 540.0, 550.0, 560.0, 570.0, 580.0, 590.0, 600.0, 610.0,
            620.0, 630.0, 640.0, 650.0, 660.0, 670.0, 680.0, 690.0, 700.0, 710.0, 720.0, 730.0,
            740.0, 750.0, 760.0, 770.0, 780.0, 790.0, 800.0, 810.0, 820.0, 830.0, 840.0, 850.0,
            860.0, 870.0, 880.0, 890.0, 900.0, 910.0, 920.0, 930.0, 940.0, 950.0, 960.0, 970.0,
            980.0, 990.0, 1000.0, 1010.0, 1020.0, 1030.0, 1040.0, 1050.0, 1060.0, 1070.0,
            1080.0, 1090.0, 1100.0],
            [1.5880, 1.5850, 1.5880, 1.5850, 1.5670, 1.5370, 1.5050, 1.4720, 1.4220, 1.3330,
            1.1930, 1.0110, 0.8197, 0.6527, 0.5263, 0.4363, 0.3724, 0.3256, 0.2899, 0.2617,
            0.2388, 0.2199, 0.2041, 0.1908, 0.1794, 0.1698, 0.1616, 0.1546, 0.1487, 0.1436,
            0.1394, 0.1358, 0.1328, 0.1304, 0.1284, 0.1268, 0.1257, 0.1248, 0.1242, 0.1239,
            0.1239, 0.1240, 0.1244, 0.1249, 0.1256, 0.1265, 0.1275, 0.1286, 0.1298, 0.1312,
            0.1326, 0.1342, 0.1358, 0.1375, 0.1393, 0.1412, 0.1431, 0.1451, 0.1471, 0.1493,
            0.1514, 0.1536, 0.1559, 0.1582, 0.1606, 0.1630, 0.1655, 0.1680, 0.1705, 0.1731,
            0.1757, 0.1783, 0.1810],
            [1.8770, 1.9040, 1.9150, 1.9110, 1.8980, 1.8870, 1.8770, 1.8550, 1.8120, 1.7580,
            1.7160, 1.7200, 1.7870, 1.9080, 2.0580, 2.2130, 2.3640, 2.5070, 2.6410, 2.7690,
            2.8910, 3.0090, 3.1220, 3.2320, 3.3400, 3.4440, 3.5470, 3.6470, 3.7460, 3.8430,
            3.9380, 4.0330, 4.1260, 4.2170, 4.3080, 4.3980, 4.4870, 4.5750, 4.6620, 4.7490,
            4.8340, 4.9200, 5.0040, 5.0880, 5.1720, 5.2550, 5.3380, 5.4200, 5.5010, 5.5830,
            5.6640, 5.7450, 5.8250, 5.9050, 5.9850, 6.0640, 6.1430, 6.2220, 6.3010, 6.3790,
            6.4570, 6.5350, 6.6130, 6.6900, 6.7680, 6.8450, 6.9220, 6.9990, 7.0750, 7.1520,
            7.2280, 7.3040, 7.3800],
        ),
        "McPeak 2015": (
            [380.0, 390.0, 400.0, 410.0, 420.0, 430.0, 440.0, 450.0, 460.0, 470.0, 480.0, 490.0,
            500.0, 510.0, 520.0, 530.0, 540.0, 550.0, 560.0, 570.0, 580.0, 590.0, 600.0, 610.0,
            620.0, 630.0, 640.0, 650.0, 660.0, 670.0, 680.0, 690.0, 700.0, 710.0, 720.0, 730.0,
            740.0, 750.0, 760.0, 770.0, 780.0, 790.0, 800.0, 810.0, 820.0, 830.0, 840.0, 850.0,
            860.0, 870.0, 880.0, 890.0, 900.0, 910.0, 920.0, 930.0, 940.0, 950.0, 960.0, 970.0,
            980.0, 990.0, 1000.0, 1010.0, 1020.0, 1030.0, 1040.0, 1050.0, 1060.0, 1070.0,
            1080.0, 1090.0, 1100.0],
            [1.6781, 1.6725, 1.6656, 1.6539, 1.6374, 1.6123, 1.5815, 1.5383, 1.4753, 1.3863,
            1.2529, 1.0644, 0.8485, 0.6616, 0.5291, 0.4380, 0.3725, 0.3239, 0.2850, 0.2541,
            0.2289, 0.2071, 0.1888, 0.1727, 0.1601, 0.1464, 0.1354, 0.1255, 0.1178, 0.1111,
            0.1057, 0.1015, 0.0986, 0.0978, 0.0986, 0.0969, 0.0995, 0.0985, 0.0992, 0.1017,
            0.1013, 0.1031, 0.1042, 0.1059, 0.1043, 0.1045, 0.1070, 0.1110, 0.1108, 0.1120,
            0.1142, 0.1160, 0.1168, 0.1170, 0.1195, 0.1206, 0.1226, 0.1240, 0.1239, 0.1262,
            0.1280, 0.1299, 0.1323, 0.1339, 0.1354, 0.1381, 0.1379, 0.1392, 0.1340, 0.1357,
            0.1404, 0.1387, 0.1384],
            [1.9484, 1.9665, 1.9739, 1.9752, 1.9703, 1.9595, 1.9403, 1.9107, 1.8723, 1.8277,
            1.7821, 1.7679, 1.8283, 1.9645, 2.1297, 2.2950, 2.4517, 2.5972, 2.7390, 2.8719,
            2.9995, 3.1213, 3.2417, 3.3568, 3.4656, 3.5776, 3.6860, 3.7923, 3.8961, 3.9995,
            4.1027, 4.2050, 4.3051, 4.4038, 4.4992, 4.5958, 4.6872, 4.7811, 4.8724, 4.9589,
            5.0516, 5.1481, 5.2237, 5.3130, 5.4074, 5.4924, 5.5770, 5.6593, 5.7492, 5.8306,
            5.9130, 5.9929, 6.0820, 6.1638, 6.2416, 6.3286, 6.4104, 6.4918, 6.5740, 6.6629,
            6.7381, 6.8192, 6.9045, 6.9783, 7.0616, 7.1416, 7.2225, 7.3021, 7.4066, 7.4828,
            7.5568, 7.6370, 7.7193],
        ),
    },
    "Ag": {
        "Johnson & Christy 1972": (
            [381.5, 397.4, 413.3, 430.5, 450.9, 471.4, 495.9, 520.9, 548.6, 582.1, 616.8, 659.5,
            704.5, 756.0, 821.1, 892.0, 984.0, 1088.0],
            [0.0500, 0.0500, 0.0500, 0.0400, 0.0400, 0.0500, 0.0500, 0.0500, 0.0600, 0.0500,
            0.0600, 0.0500, 0.0400, 0.0300, 0.0400, 0.0400, 0.0400, 0.0400],
            [1.8640, 2.0700, 2.2750, 2.4620, 2.6570, 2.8690, 3.0930, 3.3240, 3.5860, 3.8580,
            4.1520, 4.4830, 4.8380, 5.2420, 5.7270, 6.3120, 6.9920, 7.7950],
        ),
        "Babar & Weaver 2015": (
            [387.5, 399.9, 413.3, 427.5, 442.8, 459.2, 476.9, 495.9, 516.6, 539.1, 563.6, 590.4,
            619.9, 652.5, 688.8, 729.3, 774.9, 826.6, 885.6, 953.7, 1033.0],
            [0.0500, 0.0540, 0.0500, 0.0510, 0.0520, 0.0520, 0.0530, 0.0520, 0.0520, 0.0517,
            0.0501, 0.0498, 0.0480, 0.0496, 0.0508, 0.0515, 0.0527, 0.0550, 0.0590, 0.0654,
            0.0729],
            [1.9810, 2.1380, 2.2920, 2.4480, 2.6040, 2.7650, 2.9300, 3.1050, 3.2880, 3.4830,
            3.6940, 3.9180, 4.1640, 4.4320, 4.7250, 5.0480, 5.4090, 5.8140, 6.2760, 6.8020,
            7.4120],
        ),
        "McPeak 2015": (
            [380.0, 390.0, 400.0, 410.0, 420.0, 430.0, 440.0, 450.0, 460.0, 470.0, 480.0, 490.0,
            500.0, 510.0, 520.0, 530.0, 540.0, 550.0, 560.0, 570.0, 580.0, 590.0, 600.0, 610.0,
            620.0, 630.0, 640.0, 650.0, 660.0, 670.0, 680.0, 690.0, 700.0, 710.0, 720.0, 730.0,
            740.0, 750.0, 760.0, 770.0, 780.0, 790.0, 800.0, 810.0, 820.0, 830.0, 840.0, 850.0,
            860.0, 870.0, 880.0, 890.0, 900.0, 910.0, 920.0, 930.0, 940.0, 950.0, 960.0, 970.0,
            980.0, 990.0, 1000.0, 1010.0, 1020.0, 1030.0, 1040.0, 1050.0, 1060.0, 1070.0,
            1080.0, 1090.0, 1100.0],
            [0.0508, 0.0487, 0.0457, 0.0445, 0.0433, 0.0420, 0.0411, 0.0409, 0.0410, 0.0407,
            0.0408, 0.0412, 0.0414, 0.0416, 0.0424, 0.0422, 0.0434, 0.0438, 0.0445, 0.0460,
            0.0468, 0.0468, 0.0474, 0.0490, 0.0496, 0.0511, 0.0516, 0.0509, 0.0534, 0.0525,
            0.0542, 0.0544, 0.0549, 0.0565, 0.0574, 0.0577, 0.0596, 0.0601, 0.0590, 0.0638,
            0.0635, 0.0622, 0.0640, 0.0675, 0.0676, 0.0682, 0.0675, 0.0701, 0.0705, 0.0725,
            0.0739, 0.0749, 0.0748, 0.0776, 0.0788, 0.0801, 0.0800, 0.0811, 0.0840, 0.0823,
            0.0853, 0.0878, 0.0880, 0.0905, 0.0912, 0.0932, 0.0910, 0.0931, 0.0875, 0.0901,
            0.0914, 0.0895, 0.0876],
            [1.8587, 1.9955, 2.1229, 2.2418, 2.3572, 2.4653, 2.5733, 2.6758, 2.7770, 2.8738,
            2.9722, 3.0663, 3.1594, 3.2511, 3.3421, 3.4332, 3.5222, 3.6101, 3.6970, 3.7839,
            3.8707, 3.9572, 4.0419, 4.1263, 4.2090, 4.2931, 4.3768, 4.4602, 4.5431, 4.6237,
            4.7059, 4.7884, 4.8691, 4.9508, 5.0342, 5.1132, 5.1953, 5.2744, 5.3568, 5.4362,
            5.5130, 5.5988, 5.6781, 5.7533, 5.8362, 5.9113, 5.9955, 6.0765, 6.1553, 6.2356,
            6.3136, 6.3898, 6.4692, 6.5476, 6.6280, 6.7061, 6.7841, 6.8651, 6.9423, 7.0244,
            7.1001, 7.1771, 7.2542, 7.3352, 7.4110, 7.4884, 7.5716, 7.6475, 7.7382, 7.8126,
            7.8964, 7.9690, 8.0559],
        ),
    },
    "Cu": {
        "Johnson & Christy 1972": (
            [381.5, 397.4, 413.3, 430.5, 450.9, 471.4, 495.9, 520.9, 548.6, 582.1, 616.8, 659.5,
            704.5, 756.0, 821.1, 892.0, 984.0, 1088.0],
            [1.3300, 1.3200, 1.2800, 1.2500, 1.2400, 1.2500, 1.2200, 1.1800, 1.0200, 0.7000,
            0.3000, 0.2200, 0.2100, 0.2400, 0.2600, 0.3000, 0.3200, 0.3600],
            [2.0450, 2.1160, 2.2070, 2.3050, 2.3970, 2.4830, 2.5640, 2.6080, 2.5770, 2.7040,
            3.2050, 3.7470, 4.2050, 4.6650, 5.1800, 5.7680, 6.4210, 7.2170],
        ),
        "Babar & Weaver 2015": (
            [387.5, 399.9, 413.3, 427.5, 442.8, 459.2, 476.9, 495.9, 516.6, 539.1, 563.6, 590.4,
            619.9, 652.5, 688.8, 729.3, 774.9, 826.6, 885.6, 953.7, 1033.0],
            [1.2760, 1.2390, 1.1980, 1.1860, 1.1680, 1.1470, 1.1310, 1.1210, 1.0890, 0.9717,
            0.5166, 0.1268, 0.0946, 0.0878, 0.0728, 0.0643, 0.0600, 0.0625, 0.0667, 0.0734,
            0.0839],
            [2.0370, 2.0780, 2.1580, 2.2380, 2.2990, 2.3720, 2.4450, 2.4850, 2.5010, 2.4130,
            2.4100, 2.9180, 3.3820, 3.7590, 4.1230, 4.5070, 4.9170, 5.3610, 5.8500, 6.4000,
            7.0280],
        ),
        "McPeak 2015": (
            [380.0, 390.0, 400.0, 410.0, 420.0, 430.0, 440.0, 450.0, 460.0, 470.0, 480.0, 490.0,
            500.0, 510.0, 520.0, 530.0, 540.0, 550.0, 560.0, 570.0, 580.0, 590.0, 600.0, 610.0,
            620.0, 630.0, 640.0, 650.0, 660.0, 670.0, 680.0, 690.0, 700.0, 710.0, 720.0, 730.0,
            740.0, 750.0, 760.0, 770.0, 780.0, 790.0, 800.0, 810.0, 820.0, 830.0, 840.0, 850.0,
            860.0, 870.0, 880.0, 890.0, 900.0, 910.0, 920.0, 930.0, 940.0, 950.0, 960.0, 970.0,
            980.0, 990.0, 1000.0, 1010.0, 1020.0, 1030.0, 1040.0, 1050.0, 1060.0, 1070.0,
            1080.0, 1090.0, 1100.0],
            [1.1658, 1.1399, 1.1193, 1.0977, 1.0829, 1.0672, 1.0563, 1.0482, 1.0441, 1.0408,
            1.0404, 1.0356, 1.0292, 1.0160, 0.9955, 0.9575, 0.8964, 0.7975, 0.6499, 0.4677,
            0.3081, 0.2065, 0.1534, 0.1297, 0.1167, 0.1101, 0.1072, 0.1042, 0.1025, 0.1024,
            0.1012, 0.1016, 0.1012, 0.1016, 0.1011, 0.1008, 0.1009, 0.1012, 0.1018, 0.1017,
            0.1042, 0.1015, 0.1051, 0.1056, 0.1048, 0.1081, 0.1063, 0.1068, 0.1081, 0.1094,
            0.1092, 0.1107, 0.1115, 0.1106, 0.1144, 0.1157, 0.1162, 0.1165, 0.1172, 0.1176,
            0.1200, 0.1207, 0.1233, 0.1235, 0.1259, 0.1267, 0.1255, 0.1279, 0.1230, 0.1241,
            0.1272, 0.1261, 0.1270],
            [2.0463, 2.0901, 2.1422, 2.1935, 2.2512, 2.3068, 2.3619, 2.4136, 2.4641, 2.5090,
            2.5496, 2.5777, 2.6010, 2.6106, 2.6139, 2.6036, 2.5841, 2.5642, 2.5666, 2.6337,
            2.7745, 2.9531, 3.1248, 3.2808, 3.4222, 3.5466, 3.6668, 3.7757, 3.8796, 3.9818,
            4.0823, 4.1751, 4.2706, 4.3654, 4.4537, 4.5415, 4.6328, 4.7186, 4.8069, 4.8903,
            4.9858, 5.0588, 5.1413, 5.2257, 5.3144, 5.3990, 5.4717, 5.5584, 5.6436, 5.7181,
            5.7994, 5.8733, 5.9579, 6.0421, 6.1147, 6.1944, 6.2722, 6.3524, 6.4336, 6.5070,
            6.5884, 6.6653, 6.7449, 6.8194, 6.8948, 6.9775, 7.0487, 7.1286, 7.2144, 7.2909,
            7.3683, 7.4471, 7.5258],
        ),
    },
    "Pt": {
        "Rakic 1998 (BB model)": (
            [382.1, 389.7, 397.5, 405.3, 413.4, 421.6, 430.0, 438.5, 447.2, 456.1, 465.2, 474.4,
            483.8, 493.4, 503.2, 513.2, 523.4, 533.8, 544.4, 555.2, 566.2, 577.4, 588.9, 600.6,
            612.5, 624.7, 637.1, 649.7, 662.6, 675.8, 689.2, 702.9, 716.8, 731.1, 745.6, 760.4,
            775.5, 790.9, 806.6, 822.6, 838.9, 855.6, 872.6, 889.9, 907.6, 925.6, 943.9, 962.7,
            981.8, 1001.3, 1021.2, 1041.4, 1062.1, 1083.2],
            [1.6838, 1.7032, 1.7225, 1.7418, 1.7611, 1.7804, 1.7998, 1.8193, 1.8390, 1.8591,
            1.8796, 1.9006, 1.9222, 1.9445, 1.9676, 1.9916, 2.0166, 2.0427, 2.0699, 2.0984,
            2.1281, 2.1592, 2.1917, 2.2255, 2.2607, 2.2974, 2.3354, 2.3748, 2.4155, 2.4574,
            2.5006, 2.5449, 2.5903, 2.6366, 2.6837, 2.7317, 2.7802, 2.8294, 2.8789, 2.9288,
            2.9790, 3.0294, 3.0800, 3.1306, 3.1814, 3.2324, 3.2835, 3.3350, 3.3871, 3.4398,
            3.4934, 3.5483, 3.6049, 3.6634],
            [2.7414, 2.7867, 2.8323, 2.8785, 2.9253, 2.9727, 3.0210, 3.0702, 3.1204, 3.1717,
            3.2240, 3.2774, 3.3320, 3.3878, 3.4447, 3.5027, 3.5618, 3.6219, 3.6830, 3.7451,
            3.8079, 3.8715, 3.9358, 4.0007, 4.0661, 4.1318, 4.1979, 4.2641, 4.3305, 4.3969,
            4.4633, 4.5297, 4.5958, 4.6618, 4.7276, 4.7932, 4.8586, 4.9238, 4.9889, 5.0540,
            5.1191, 5.1843, 5.2498, 5.3158, 5.3823, 5.4496, 5.5178, 5.5871, 5.6577, 5.7297,
            5.8033, 5.8785, 5.9554, 6.0341],
        ),
        "Tselin 2024": (
            [380.4, 385.4, 390.4, 395.3, 400.3, 405.3, 410.2, 413.5, 418.5, 425.1, 430.1, 435.1,
            440.0, 445.0, 450.0, 454.9, 461.6, 466.5, 471.5, 478.1, 483.1, 489.7, 494.7, 501.3,
            507.9, 512.9, 519.5, 526.1, 532.8, 537.7, 544.3, 551.0, 557.6, 564.2, 570.8, 577.5,
            585.7, 592.4, 599.0, 605.6, 613.9, 620.5, 628.8, 635.4, 643.7, 652.0, 658.6, 666.9,
            675.2, 683.4, 691.7, 700.0, 708.3, 716.6, 724.8, 733.1, 743.0, 751.3, 759.6, 769.5,
            779.5, 787.8, 797.7, 807.6, 815.9, 825.8, 835.8, 845.7, 857.3, 867.2, 877.2, 887.1,
            898.7, 908.6, 920.2, 930.2, 941.7, 953.3, 964.9, 976.5, 988.1, 999.7, 1011.3,
            1024.5, 1036.1, 1049.4, 1061.0, 1074.2, 1087.5, 1099.1],
            [1.4790, 1.4910, 1.5040, 1.5170, 1.5300, 1.5430, 1.5560, 1.5640, 1.5770, 1.5940,
            1.6060, 1.6190, 1.6320, 1.6450, 1.6570, 1.6700, 1.6870, 1.7000, 1.7130, 1.7310,
            1.7440, 1.7610, 1.7750, 1.7920, 1.8100, 1.8240, 1.8420, 1.8600, 1.8790, 1.8930,
            1.9110, 1.9300, 1.9490, 1.9690, 1.9880, 2.0080, 2.0320, 2.0520, 2.0720, 2.0930,
            2.1180, 2.1390, 2.1640, 2.1850, 2.2110, 2.2380, 2.2590, 2.2850, 2.3120, 2.3390,
            2.3660, 2.3930, 2.4200, 2.4470, 2.4750, 2.5020, 2.5350, 2.5630, 2.5900, 2.6230,
            2.6570, 2.6840, 2.7180, 2.7510, 2.7790, 2.8130, 2.8460, 2.8800, 2.9190, 2.9530,
            2.9870, 3.0210, 3.0610, 3.0950, 3.1350, 3.1690, 3.2090, 3.2500, 3.2900, 3.3310,
            3.3720, 3.4130, 3.4540, 3.5010, 3.5430, 3.5900, 3.6320, 3.6800, 3.7270, 3.7690],
            [3.2260, 3.2650, 3.3040, 3.3430, 3.3820, 3.4200, 3.4580, 3.4840, 3.5220, 3.5730,
            3.6110, 3.6480, 3.6860, 3.7240, 3.7610, 3.7990, 3.8480, 3.8860, 3.9230, 3.9720,
            4.0090, 4.0580, 4.0950, 4.1440, 4.1930, 4.2300, 4.2780, 4.3260, 4.3750, 4.4110,
            4.4590, 4.5060, 4.5540, 4.6010, 4.6480, 4.6950, 4.7530, 4.8000, 4.8460, 4.8920,
            4.9490, 4.9940, 5.0510, 5.0950, 5.1510, 5.2060, 5.2500, 5.3050, 5.3590, 5.4130,
            5.4660, 5.5190, 5.5710, 5.6240, 5.6750, 5.7270, 5.7880, 5.8380, 5.8880, 5.9480,
            6.0070, 6.0560, 6.1150, 6.1720, 6.2200, 6.2770, 6.3340, 6.3900, 6.4540, 6.5090,
            6.5640, 6.6180, 6.6810, 6.7340, 6.7950, 6.8470, 6.9070, 6.9670, 7.0250, 7.0830,
            7.1410, 7.1970, 7.2530, 7.3160, 7.3700, 7.4300, 7.4830, 7.5410, 7.5980, 7.6470],
        ),
    },
    "Pd": {
        "Johnson & Christy 1974": (
            [381.0, 397.0, 413.0, 431.0, 451.0, 471.0, 496.0, 521.0, 549.0, 582.0, 617.0, 659.0,
            704.0, 756.0, 821.0, 892.0, 984.0, 1088.0],
            [1.2600, 1.3000, 1.3300, 1.3700, 1.4100, 1.4600, 1.5200, 1.5700, 1.6400, 1.6800,
            1.7500, 1.8000, 1.8600, 1.9500, 2.0600, 2.2300, 2.3400, 2.5200],
            [2.8300, 2.9300, 3.0300, 3.1400, 3.2600, 3.3900, 3.5400, 3.6800, 3.8400, 4.0200,
            4.2100, 4.4200, 4.6500, 4.8900, 5.1900, 5.5000, 5.8900, 6.3300],
        ),
        "Rakic 1998 (BB model)": (
            [382.1, 389.7, 397.5, 405.3, 413.4, 421.6, 430.0, 438.5, 447.2, 456.1, 465.2, 474.4,
            483.8, 493.4, 503.2, 513.2, 523.4, 533.8, 544.4, 555.2, 566.2, 577.4, 588.9, 600.6,
            612.5, 624.7, 637.1, 649.7, 662.6, 675.8, 689.2, 702.9, 716.8, 731.1, 745.6, 760.4,
            775.5, 790.9, 806.6, 822.6, 838.9, 855.6, 872.6, 889.9, 907.6, 925.6, 943.9, 962.7,
            981.8, 1001.3, 1021.2, 1041.4, 1062.1, 1083.2],
            [1.2443, 1.2602, 1.2771, 1.2948, 1.3132, 1.3324, 1.3522, 1.3726, 1.3935, 1.4149,
            1.4368, 1.4592, 1.4819, 1.5050, 1.5285, 1.5522, 1.5763, 1.6006, 1.6252, 1.6500,
            1.6751, 1.7003, 1.7258, 1.7514, 1.7772, 1.8032, 1.8292, 1.8555, 1.8818, 1.9083,
            1.9348, 1.9615, 1.9882, 2.0150, 2.0418, 2.0687, 2.0957, 2.1227, 2.1498, 2.1768,
            2.2040, 2.2311, 2.2583, 2.2855, 2.3128, 2.3401, 2.3675, 2.3950, 2.4225, 2.4502,
            2.4780, 2.5059, 2.5340, 2.5622],
            [2.8123, 2.8618, 2.9117, 2.9620, 3.0125, 3.0633, 3.1145, 3.1659, 3.2176, 3.2696,
            3.3219, 3.3746, 3.4276, 3.4810, 3.5348, 3.5891, 3.6437, 3.6988, 3.7544, 3.8105,
            3.8671, 3.9243, 3.9820, 4.0404, 4.0993, 4.1589, 4.2192, 4.2802, 4.3419, 4.4044,
            4.4676, 4.5316, 4.5965, 4.6623, 4.7289, 4.7965, 4.8651, 4.9347, 5.0053, 5.0770,
            5.1498, 5.2238, 5.2990, 5.3755, 5.4532, 5.5323, 5.6127, 5.6946, 5.7780, 5.8628,
            5.9492, 6.0373, 6.1269, 6.2183],
        ),
        "Palm 2018": (
            [382.5, 385.7, 390.5, 395.3, 400.1, 404.9, 409.7, 414.4, 419.2, 425.6, 430.4, 435.2,
            440.0, 444.8, 451.1, 455.9, 462.3, 467.1, 473.5, 478.3, 484.6, 489.4, 495.8, 502.2,
            508.5, 513.3, 519.7, 526.1, 532.4, 538.8, 545.2, 551.6, 557.9, 565.9, 572.2, 578.6,
            585.0, 592.9, 599.3, 607.2, 613.6, 621.5, 629.5, 635.8, 643.8, 651.7, 659.7, 667.6,
            675.5, 683.5, 691.4, 700.9, 708.8, 716.8, 726.3, 734.2, 743.7, 751.6, 761.1, 770.6,
            780.1, 788.0, 797.4, 806.9, 817.9, 827.4, 836.9, 846.3, 857.3, 866.8, 877.8, 888.8,
            898.2, 909.2, 920.2, 931.2, 942.2, 954.7, 965.7, 976.6, 989.1, 1011.4, 1014.8,
            1025.0, 1038.6, 1048.7, 1062.3, 1075.9, 1089.5, 1099.7],
            [1.3585, 1.3663, 1.3764, 1.3853, 1.3935, 1.4015, 1.4097, 1.4185, 1.4284, 1.4437,
            1.4567, 1.4707, 1.4856, 1.5009, 1.5219, 1.5378, 1.5587, 1.5741, 1.5938, 1.6078,
            1.6257, 1.6387, 1.6554, 1.6717, 1.6877, 1.6996, 1.7155, 1.7315, 1.7476, 1.7640,
            1.7808, 1.7978, 1.8150, 1.8368, 1.8544, 1.8721, 1.8898, 1.9120, 1.9298, 1.9519,
            1.9694, 1.9912, 2.0128, 2.0299, 2.0510, 2.0721, 2.0930, 2.1138, 2.1346, 2.1553,
            2.1760, 2.2008, 2.2215, 2.2421, 2.2669, 2.2876, 2.3123, 2.3330, 2.3578, 2.3824,
            2.4070, 2.4274, 2.4518, 2.4760, 2.5039, 2.5277, 2.5511, 2.5744, 2.6012, 2.6238,
            2.6498, 2.6754, 2.6970, 2.7217, 2.7460, 2.7698, 2.7932, 2.8193, 2.8418, 2.8638,
            2.8884, 2.9311, 2.9375, 2.9563, 2.9809, 2.9990, 3.0226, 3.0457, 3.0683, 3.0848],
            [2.9643, 2.9839, 3.0138, 3.0447, 3.0769, 3.1105, 3.1453, 3.1812, 3.2180, 3.2676,
            3.3046, 3.3409, 3.3764, 3.4110, 3.4558, 3.4882, 3.5301, 3.5606, 3.6003, 3.6296,
            3.6685, 3.6976, 3.7365, 3.7756, 3.8149, 3.8445, 3.8841, 3.9240, 3.9639, 4.0039,
            4.0438, 4.0834, 4.1228, 4.1716, 4.2102, 4.2485, 4.2864, 4.3333, 4.3703, 4.4162,
            4.4525, 4.4975, 4.5421, 4.5775, 4.6216, 4.6654, 4.7090, 4.7524, 4.7956, 4.8385,
            4.8812, 4.9321, 4.9743, 5.0162, 5.0661, 5.1074, 5.1565, 5.1971, 5.2453, 5.2931,
            5.3403, 5.3792, 5.4255, 5.4713, 5.5241, 5.5688, 5.6130, 5.6569, 5.7074, 5.7503,
            5.7998, 5.8489, 5.8905, 5.9387, 5.9864, 6.0338, 6.0809, 6.1343, 6.1807, 6.2269,
            6.2794, 6.3725, 6.3865, 6.4286, 6.4845, 6.5263, 6.5819, 6.6374, 6.6927, 6.7341],
        ),
    },
}
# ============================================================================

# ---- Analytic prism and medium dispersion formulas from refractiveindex.info (CC0) ----
# Source references, DOIs, and symbol definitions are documented in this module.
# Entry format: (form, coefficients, lam_lo_nm, lam_hi_nm); reject wavelengths outside the source range.
# form=1:  n^2 - 1 = sum_i  B_i * L^2 / (L^2 - C_i^2)      (Sellmeier form with C squared)
# form=2:  n^2 - 1 = sum_i  B_i * L^2 / (L^2 - C_i)        (Sellmeier form with already-squared C)
# form=6:  n   - 1 = sum_i  B_i / (C_i - L^-2)             (Ciddor gas dispersion)
# L = vacuum wavelength (um)
SELLMEIER = {
    # --- Prism and incident-side glass (SCHOTT Sellmeier coefficients) ---
    "N-BK7 (coverslip/objective coupling)": (2, [1.03961212, 0.00600069867, 0.231792344,
                                    0.0200179144, 1.01046945, 103.560653], 300, 2500),
    "SF10":                    (2, [1.61625977, 0.0127534559, 0.259229334,
                                    0.0581983954, 1.07762317, 116.60768], 380, 2500),
    "SF11":                    (2, [1.73848403, 0.0136068604, 0.311168974,
                                    0.0615960463, 1.17490871, 121.922711], 390, 2500),
    "SF14":                    (2, [1.69182538, 0.0133151542, 0.285919934,
                                    0.0612647445, 1.12595145, 118.405242], 390, 2500),
    "LaSFN9":                  (2, [1.97888194, 0.0118537266, 0.320435298,
                                    0.052738177, 1.92900751, 166.25654], 365, 2500),
    "Fused silica (SiO2)":     (1, [0.6961663, 0.0684043, 0.4079426,
                                    0.1162414, 0.8974794, 9.896161], 210, 6700),
    "Sapphire (o-ray)":        (1, [1.4313493, 0.0726631, 0.65054713,
                                    0.1193242, 5.3414021, 18.028251], 200, 5000),

    # --- Ambient medium: water (Daimon and Masumura 2007, temperature-specific data) ---
    # SPR is temperature sensitive: dn/dT is about -1e-4 /degC, and 1e-6 RIU corresponds to about 0.01 degC.
    # The values in Nat. Rev. Methods Primers Box 3 (eps=1.773) and Fig. 4b (eps=1.7734) correspond to 24 degC.
    "Water 19.0°C":   (2, [0.5672526103, 0.005085550461, 0.1736581125, 0.01814938654,
                           0.02121531502, 0.02617260739, 0.1138493213, 10.73888649], 182, 1129),
    "Water 20.0°C":   (2, [0.5684027565, 0.005101829712, 0.1726177391, 0.01821153936,
                           0.02086189578, 0.02620722293, 0.1130748688, 10.69792721], 182, 1129),
    "Water 21.5°C":   (2, [0.5689093832, 0.005110301794, 0.1719708856, 0.01825180155,
                           0.02062501582, 0.02624158904, 0.1123965424, 10.67505178], 182, 1129),
    "Water 24.0°C":   (2, [0.566695982, 0.005084151894, 0.1731900098, 0.01818488474,
                           0.02095951857, 0.02625439472, 0.1125228406, 10.73842352], 182, 1129),

    # --- Ambient medium: standard dry air (Ciddor 1996; 15 degC, 101.325 kPa, 450 ppm CO2) ---
    # n(632.8nm) = 1.0002765. Using n=1 (vacuum) shifts theta_min systematically by about 0.01 degree.
    "Air (15°C, 101.325 kPa)": (6, [0.05792105, 238.0185, 0.00167917, 57.362], 230, 1690),
}

PRISM_LIST = ["N-BK7 (coverslip/objective coupling)", "SF10", "SF11", "SF14", "LaSFN9",
              "Fused silica (SiO2)", "Sapphire (o-ray)", "Custom (constant n)"]
MEDIUM_LIST = ["Water 20.0°C", "Water 19.0°C", "Water 21.5°C", "Water 24.0°C",
               "Air (15°C, 101.325 kPa)", "Vacuum (n=1)", "Custom (constant n)"]

# ---- Excitation wavelengths: real laboratory laser lines (vacuum wavelengths, nm) ----
# Why use a dropdown: HeNe is 632.8 nm, not 633 nm. The 0.2 nm difference shifts theta_min by about 4 mdeg,
# four times the program's 1 mdeg numerical resolution. The dropdown preserves the exact value.
# Why include more than 532/633 nm: Au can have better modeled SPR performance at 785/850 nm (more negative Re(eps), higher Q,
# narrower dip); commercial SPR instruments may use 760 nm. Multiple wavelengths permit laser-choice comparisons.
WAVELENGTHS = [
    "405.0 nm — violet diode",
    "488.0 nm — Ar+ / blue DPSS",
    "532.0 nm — green DPSS (frequency-doubled Nd:YAG)",
    "594.1 nm — yellow HeNe",
    "632.8 nm — red HeNe",
    "660.0 nm — diode",
    "785.0 nm — NIR diode",
    "850.0 nm — NIR diode",
    "1064.0 nm — Nd:YAG fundamental",
    "Custom...",
]
WL_DEFAULT = "632.8 nm — red HeNe"

METALS = ["Au", "Ag", "Cu", "Pt", "Pd"]
CUSTOM = "Custom"          # pseudo-material for manually entered eps (literature or ellipsometric fits)

# Okabe-Ito colorblind-safe palette plus redundant line styles
STYLE = {
    "Au": ("#E69F00", "-"),
    "Ag": ("#56B4E9", "--"),
    "Cu": ("#D55E00", "-."),
    "Pt": ("#009E73", ":"),
    "Pd": ("#CC79A7", (0, (3, 1, 1, 1))),
    CUSTOM: ("#000000", (0, (5, 2, 1, 2, 1, 2))),
}

Q_OVERDAMPED = 3.0        # below this |Re(eps)|/Im(eps) threshold, the SPP is overdamped
NK_LO, NK_HI = 380.0, 1100.0

THETA_DEF_CURVE = "Tab A sampled curve minimum (current theta range and step)"
THETA_DEF_SPR = "Refined SPR minimum above theta_c"
THETA_DEFINITIONS = [THETA_DEF_CURVE, THETA_DEF_SPR]


# ============================================================================
#  Physical core
# ============================================================================

def parse_wl(label):
    """Return 632.8 for a HeNe label or None for a custom wavelength."""
    if label.startswith("Custom"):
        return None
    return float(label.split()[0])


def sellmeier_n(name, lam_nm):
    """Evaluate dispersion to obtain n for a transparent medium (k=0, so eps=n^2 is real).
    Reject wavelengths outside the source range; do not extrapolate."""
    form, c, lo, hi = SELLMEIER[name]
    if not (lo - 1e-6 <= lam_nm <= hi + 1e-6):
        raise ValueError(f"{name} dispersion formula is valid from {lo:.0f}–{hi:.0f} nm; "
                         f"it does not cover {lam_nm:.1f} nm (extrapolation rejected)")
    L = lam_nm / 1000.0                      # um
    if form == 6:                            # Ciddor gas: n - 1 = sum B/(C - L^-2)
        s2 = 1.0 / L ** 2
        n = 1.0
        for i in range(0, len(c), 2):
            n += c[i] / (c[i + 1] - s2)
        return float(n)
    L2 = L ** 2                              # Sellmeier: n^2 - 1 = sum B*L^2/(L^2 - C)
    n2 = 1.0
    for i in range(0, len(c), 2):
        B, C = c[i], c[i + 1]
        n2 += B * L2 / (L2 - (C if form == 2 else C ** 2))
    return float(np.sqrt(n2))


_INTERP_CACHE = {}


def eps_metal(metal, source, lam_nm):
    """Return metal eps=(n+ik)^2 using shape-preserving PCHIP interpolation without ringing."""
    key = (metal, source)
    if key not in _INTERP_CACHE:
        w, n, k = NK_DATA[metal][source]
        w = np.asarray(w, float)
        _INTERP_CACHE[key] = (PchipInterpolator(w, np.asarray(n, float)),
                              PchipInterpolator(w, np.asarray(k, float)),
                              w[0], w[-1])
    fn, fk, lo, hi = _INTERP_CACHE[key]
    if not (lo - 1e-6 <= lam_nm <= hi + 1e-6):
        raise ValueError(f"{metal} / {source} data range is {lo:.0f}-{hi:.0f} nm; "
                         f"does not cover {lam_nm:.1f} nm")
    return complex((float(fn(lam_nm)) + 1j * float(fk(lam_nm))) ** 2)


def spp_quality(eps_m):
    """SPP quality proxy |Re(eps)|/Im(eps); values above about 3 can support a usable resonance."""
    return abs(eps_m.real) / eps_m.imag if eps_m.imag > 0 else np.inf


def spp_cycles(eps_m, eps_d):
    """Lx/lambda_spp is the number of SPP wavelengths traveled before intensity decay.
    = Re(n_eff) / (4*pi*Im(n_eff)),  n_eff = sqrt(eps_m*eps_d/(eps_m+eps_d))

    This is a useful propagation-quality metric:
        Au @633nm -> 8.8   (well-propagating SPP)
        Ag @633nm -> 55.0  (excellent)
        Au @532nm -> 0.7   (decays within one SPP wavelength; localized absorption dominates)
        Pt @633nm -> 2.3   (marginal)
    """
    s = complex(eps_m) + eps_d
    if s.real >= 0:
        return np.nan
    ns = np.sqrt(complex(eps_m) * eps_d / s)
    if ns.imag <= 0 or ns.real <= 0:
        return np.nan
    return float(ns.real / (4.0 * np.pi * ns.imag))


def critical_angle(eps_p, eps_d):
    """Total-internal-reflection critical angle (degrees)"""
    s = np.sqrt(eps_d / eps_p)
    return float(np.rad2deg(np.arcsin(s))) if s < 1 else 0.0


def reflectivity_p(eps_list, d_list, lam_nm, theta_deg):
    """N-layer p-polarized reflectivity, vectorized over theta.
    eps_list = [eps_in(semi-infinite), ..., eps_out(semi-infinite)]; d_list = internal-layer thicknesses (nm)"""
    k0 = 2.0 * np.pi / lam_nm
    th = np.deg2rad(np.atleast_1d(np.asarray(theta_deg, float)))
    kx = k0 * np.sqrt(complex(eps_list[0])) * np.sin(th)
    kz = []
    for e in eps_list:
        q = np.sqrt(np.asarray(complex(e) * k0 ** 2 - kx ** 2, dtype=complex))
        kz.append(np.where(q.imag < 0, -q, q))          # enforce Im(kz) >= 0
    def rij(i, j):
        return ((eps_list[j] * kz[i] - eps_list[i] * kz[j]) /
                (eps_list[j] * kz[i] + eps_list[i] * kz[j]))
    N = len(eps_list) - 1
    r = rij(N - 1, N)
    for i in range(N - 2, -1, -1):                      # bottom-up Airy recursion
        ph = np.exp(2j * kz[i + 1] * d_list[i])
        r = (rij(i, i + 1) + r * ph) / (1.0 + rij(i, i + 1) * r * ph)
    return np.abs(r) ** 2


def find_curve_minimum(eps_p, eps_m, eps_d, d_nm, lam_nm, theta_deg,
                       r_min_max=0.60, depth_min=0.05):
    """Return the lowest sampled point on a discrete angular curve and its detectability.

    This matches the R(theta) samples plotted in Tab A: no inter-sample refinement and no automatic
    exclusion of subcritical Fresnel minima. The separate ``valid`` flag identifies usable SPR minima.
    """
    theta = np.atleast_1d(np.asarray(theta_deg, float))
    if theta.size < 2 or np.any(~np.isfinite(theta)) or np.any(np.diff(theta) <= 0):
        raise ValueError("A sampled-curve minimum requires at least two strictly increasing finite theta values")

    R = reflectivity_p([eps_p, eps_m, eps_d], [d_nm], lam_nm, theta)
    i0 = int(np.argmin(R))
    theta_min = float(theta[i0])
    R_min = float(R[i0])

    right = R[i0:]
    base = float(right.max()) if right.size else np.nan
    depth = base - R_min
    half = (base + R_min) / 2.0
    li, ri = i0, i0
    while li > 0 and R[li] < half:
        li -= 1
    while ri < theta.size - 1 and R[ri] < half:
        ri += 1
    fwhm = float(theta[ri] - theta[li]) if (li > 0 and ri < theta.size - 1) else np.nan

    notes = []
    dip_ok = True
    if R_min > r_min_max:
        dip_ok = False; notes.append(f"R_min={R_min:.2f} is too high")
    if not np.isfinite(depth) or depth < depth_min:
        dip_ok = False; notes.append("dip is too shallow")
    step = float(np.median(np.diff(theta)))
    edge_margin = max(0.05, 1.5 * step)
    if theta_min < theta[0] + edge_margin or theta_min > theta[-1] - edge_margin:
        dip_ok = False; notes.append("minimum touches the sampled-curve boundary")

    tc = critical_angle(eps_p, eps_d)
    above_critical = theta_min > tc + 0.02
    if not above_critical:
        notes.append(f"below critical angle theta_c={tc:.2f} deg; not SPR")
    Q = spp_quality(eps_m)
    cyc = spp_cycles(eps_m, eps_d)
    overdamped = Q < Q_OVERDAMPED
    if overdamped:
        notes.append(f"overdamped: Q={Q:.2f}, Lx/lambda_spp={cyc:.1f}")

    return dict(theta_min=theta_min, R_min=R_min, fwhm=fwhm, depth=depth,
                dip_ok=dip_ok, overdamped=overdamped,
                valid=(dip_ok and above_critical and not overdamped),
                above_critical=above_critical,
                note="; ".join(notes) if notes else "")


def find_dip(eps_p, eps_m, eps_d, d_nm, lam_nm,
             r_min_max=0.60, depth_min=0.05):
    """Locate the reflectivity minimum. Return:
        theta_min, R_min, fwhm, depth   —— numerical values returned even when not usable
        dip_ok      —— whether the dip is detectable (deep enough and away from boundaries)
        overdamped  —— whether the SPP is overdamped (Q < 3), independently of dip depth
        valid       —— dip_ok AND (not overdamped), usable as a resonance angle

    R_min alone cannot establish an SPP resonance.
    A sufficiently thin film of almost any metal can reach critical coupling (Gamma_rad = Gamma_int), driving R_min toward zero.
    For SF10/water at 632.8 nm, Pt at d=16 nm has R_min=3e-5, as deep as Au, but its FWHM
    is 31.8 deg versus 3.7 deg for Au. FWHM and Lx distinguish propagation quality.
    The overdamping check (Q; see spp_quality) is therefore necessary and independent of dip depth.
    """
    tc = critical_angle(eps_p, eps_d)
    lo, hi = tc + 0.02, 89.90
    if lo >= hi:
        return dict(theta_min=np.nan, R_min=np.nan, fwhm=np.nan, depth=np.nan,
                    dip_ok=False, overdamped=False, valid=False,
                    note="no total internal reflection (n_prism <= n_medium)")

    # Use a 0.02-degree coarse search; Ag dips can be as narrow as 0.3 degree.
    ng = max(2000, int((hi - lo) / 0.02))
    tg = np.linspace(lo, hi, ng)
    Rg = reflectivity_p([eps_p, eps_m, eps_d], [d_nm], lam_nm, tg)
    i0 = int(np.argmin(Rg))

    # Refine the neighboring interval with bounded Brent optimization to <0.001 degree.
    a = tg[max(i0 - 1, 0)]
    b = tg[min(i0 + 1, ng - 1)]
    if b - a < 1e-9:
        a, b = lo, hi
    res = minimize_scalar(
        lambda t: float(reflectivity_p([eps_p, eps_m, eps_d], [d_nm], lam_nm, t)[0]),
        bounds=(a, b), method="bounded", options={"xatol": 1e-5})
    th_min, R_min = float(res.x), float(res.fun)

    # Measure dip depth against the local maximum to the right; the left side is too close to theta_c.
    right = Rg[i0:]
    base = float(right.max()) if right.size else np.nan
    depth = base - R_min

    # Angular FWHM is NaN when a half-depth crossing lies beyond the search window.
    half = (base + R_min) / 2.0
    li, ri = i0, i0
    while li > 0 and Rg[li] < half:
        li -= 1
    while ri < ng - 1 and Rg[ri] < half:
        ri += 1
    fwhm = (tg[ri] - tg[li]) if (li > 0 and ri < ng - 1) else np.nan

    notes = []
    dip_ok = True
    if R_min > r_min_max:
        dip_ok = False; notes.append(f"R_min={R_min:.2f} is too high")
    if not np.isfinite(depth) or depth < depth_min:
        dip_ok = False; notes.append("dip is too shallow")
    if th_min < lo + 0.05 or th_min > hi - 0.05:
        dip_ok = False; notes.append("minimum touches the search boundary")

    Q = spp_quality(eps_m)
    cyc = spp_cycles(eps_m, eps_d)
    overdamped = Q < Q_OVERDAMPED
    if overdamped:
        notes.append(f"overdamped: Q={Q:.2f}, Lx/lambda_spp={cyc:.1f}")

    return dict(theta_min=th_min, R_min=R_min, fwhm=fwhm, depth=depth,
                dip_ok=dip_ok, overdamped=overdamped,
                valid=(dip_ok and not overdamped),
                note="; ".join(notes) if notes else "")


def lx_intrinsic(eps_m, eps_d, lam_nm):
    """Intrinsic propagation length of a semi-infinite bound SPP (nm); NaN means no bound SPP."""
    s = complex(eps_m) + eps_d
    if s.real >= 0:
        return np.nan                       # Re(eps_m) > -eps_d; no bound SPP
    ns = np.sqrt(complex(eps_m) * eps_d / s)
    if ns.imag <= 0:
        return np.nan
    return float(1.0 / (2.0 * (2.0 * np.pi / lam_nm) * ns.imag))


def _F_pole(kx, ep, em, ed, d, k0):
    """Pole-free entire form of the SPP pole equation for numerical stability."""
    s = lambda e: np.sqrt(complex(e) * k0 ** 2 - kx ** 2)
    kz0, kz1, kz2 = s(ep), s(em), s(ed)
    if kz0.imag > 0:
        kz0 = -kz0                          # leaky sheet: radiation into the prism
    if kz2.imag < 0:
        kz2 = -kz2                          # bound sheet: confined to the interface
    N01 = em * kz0 - ep * kz1
    D01 = em * kz0 + ep * kz1
    N12 = ed * kz1 - em * kz2
    D12 = ed * kz1 + em * kz2
    return D01 * D12 + N01 * N12 * np.exp(2j * kz1 * d)


def lx_leaky_sweep(eps_p, eps_m, eps_d, d_arr, lam_nm):
    """Leaky-SPP propagation length Lx(d) in nm from thickness continuation of the complex pole.
    Return (Lx array, Lx_int); failed pole solves are NaN."""
    k0 = 2.0 * np.pi / lam_nm
    Lint = lx_intrinsic(eps_m, eps_d, lam_nm)
    d_arr = np.atleast_1d(np.asarray(d_arr, float))
    if not np.isfinite(Lint):
        return np.full(d_arr.shape, np.nan), np.nan

    im_int = 1.0 / (2.0 * Lint)             # Gamma_int; physically Gamma_tot >= Gamma_int
    # Start continuation well above the target thickness range, where Gamma_rad is nearly zero.
    pad = max(d_arr.max() * 2.0, 500.0)
    grid = np.unique(np.concatenate([
        d_arr, np.linspace(d_arr.min(), pad, 250)]))[::-1]

    guess = k0 * np.sqrt(complex(eps_m) * eps_d / (complex(eps_m) + eps_d))
    sol = {}
    for d in grid:
        fv = lambda v: [_F_pole(v[0] + 1j * v[1], eps_p, eps_m, eps_d, d, k0).real,
                        _F_pole(v[0] + 1j * v[1], eps_p, eps_m, eps_d, d, k0).imag]
        got = None
        for sc in (1.0, 1.3, 2.0):          # use continued roots first; perturb and retry after failure
            s = root(fv, [guess.real, guess.imag * sc], tol=1e-13)
            if s.success and s.x[1] > im_int * 0.98:
                got = s.x[0] + 1j * s.x[1]
                break
        if got is not None:
            guess = got
            sol[d] = 1.0 / (2.0 * got.imag)
    return np.array([sol.get(d, np.nan) for d in d_arr]), Lint


# ============================================================================
#  XLSX export
# ============================================================================

def _xlsx_value(value):
    """Convert NumPy scalars to native Excel types; write NaN and infinity as empty cells."""
    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        return float(value) if np.isfinite(value) else None
    return value


def _xlsx_common_metadata(d, export_type):
    """Traceable metadata shared by all worksheets."""
    return [
        ("export_type", export_type),
        ("generated", f"{datetime.datetime.now():%Y-%m-%d %H:%M:%S}"),
        ("wavelength_nm", float(d["lam"])),
        ("prism_n", float(d["n_p"])),
        ("prism_eps", float(d["n_p"] ** 2)),
        ("medium_n", float(d["n_d"])),
        ("medium_eps", float(d["n_d"] ** 2)),
        ("critical_angle_deg", float(d["tc"])),
        ("angle_convention", "internal angle at metal interface (glass side)"),
        ("model", "prism / metal / medium; no adhesion layer; no roughness"),
    ]


def _xlsx_number_format(header):
    """Assign Excel number formats according to the physical quantity."""
    if header in ("dip_measurable", "overdamped", "valid_SPR"):
        return "General"
    if header == "R_min" or header.startswith("R_"):
        return "0.00000000"
    if "theta" in header.lower() or header.endswith("_deg"):
        return "0.00000"
    if header in ("eps_real", "eps_imag", "Q", "Lx_over_lambda_spp"):
        return "0.0000"
    if header.endswith("_nm"):
        return "0.0000"
    if header.endswith("_um"):
        return "0.000000"
    return "General"


def _write_xlsx_sheet(ws, headers, rows, metadata):
    """Write a flat data table with metadata to the right of the importable data."""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    meta_fill = PatternFill("solid", fgColor="D9EAF7")
    meta_font = Font(name="Calibri", size=10, bold=True, color="1F1F1F")
    thin_blue = Side(style="thin", color="9EBCD1")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin_blue)

    for row in rows:
        ws.append([_xlsx_value(v) for v in row])

    data_end = max(1, len(rows) + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{data_end}"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 22

    # Apply number formats and suitable widths only to populated cells.
    for col, header in enumerate(headers, 1):
        fmt = _xlsx_number_format(header)
        for row_idx in range(2, data_end + 1):
            ws.cell(row=row_idx, column=col).number_format = fmt
        sample = [str(header)]
        sample.extend("" if r[col - 1] is None else str(r[col - 1]) for r in rows[:200])
        width = min(max(max(len(v) for v in sample) + 2, 11), 28)
        ws.column_dimensions[get_column_letter(col)].width = width

    meta_col = len(headers) + 2
    meta_key_col = get_column_letter(meta_col)
    meta_val_col = get_column_letter(meta_col + 1)
    meta_head = ws.cell(row=1, column=meta_col, value="Metadata")
    meta_head.fill = meta_fill
    meta_head.font = meta_font
    meta_head.border = Border(bottom=thin_blue)
    ws.cell(row=1, column=meta_col + 1).fill = meta_fill
    ws.cell(row=1, column=meta_col + 1).border = Border(bottom=thin_blue)
    for row_idx, (key, value) in enumerate(metadata, 2):
        ws.cell(row=row_idx, column=meta_col, value=key).font = Font(bold=True, color="404040")
        ws.cell(row=row_idx, column=meta_col + 1, value=_xlsx_value(value))
    ws.column_dimensions[meta_key_col].width = 25
    ws.column_dimensions[meta_val_col].width = 52


def save_results_xlsx(path, tag, d):
    """Write the latest calculation to a structured XLSX workbook."""
    if Workbook is None:
        raise RuntimeError("XLSX export requires openpyxl. Run: pip install openpyxl")

    wb = Workbook()
    wb.properties.creator = "SPRM Reflectivity Calculator"
    wb.properties.title = "Kretschmann SPR reflectivity results"
    wb.properties.subject = "p-polarized multilayer Fresnel/TMM calculation"

    if tag == "tab1":
        metals = list(d["curves"].keys())
        headers = ["theta_deg"] + [f"R_{m}" for m in metals]
        rows = []
        for i, theta in enumerate(d["theta"]):
            rows.append([theta] + [d["curves"][m][i] for m in metals])
        ws = wb.active
        ws.title = "R_vs_theta"
        _write_xlsx_sheet(ws, headers, rows,
                          _xlsx_common_metadata(d, "Tab A - R(theta) multi-material comparison"))

        summary_headers = [
            "material", "d_nm", "eps_real", "eps_imag", "source", "Q",
            "Lx_over_lambda_spp", "theta_curve_min_deg", "R_curve_min",
            "theta_min_deg", "R_min", "FWHM_deg",
            "Lx_leaky_um", "Lx_intrinsic_um", "dip_measurable", "overdamped",
            "valid_SPR", "note",
        ]
        summary_rows = []
        for r in d["rows"]:
            summary_rows.append([
                r["metal"], r["d"], r["eps"].real, r["eps"].imag, r["src"], r["Q"],
                r["cyc"], r["curve_theta_min"], r["curve_R_min"],
                r["theta_min"], r["R_min"], r["fwhm"], r["Lx"] / 1000,
                r["Lint"] / 1000, r["dip_ok"], r["overdamped"], r["valid"], r["note"],
            ])
        ws_summary = wb.create_sheet("Summary")
        _write_xlsx_sheet(ws_summary, summary_headers, summary_rows,
                          _xlsx_common_metadata(d, "Tab A - material summary"))
    else:
        common_meta = _xlsx_common_metadata(d, "Tab B - thickness scan") + [
            ("metal", d["metal"]),
            ("source", d["src"]),
            ("theta_min_definition", d["theta_definition"]),
            ("theta_scan_start_deg", d.get("theta_scan_start")),
            ("theta_scan_stop_deg", d.get("theta_scan_stop")),
            ("theta_scan_step_deg", d.get("theta_scan_step")),
            ("thickness_axis_scale", d.get("thickness_axis_scale", "linear")),
            ("thickness_sampling", "linear step in d_nm"),
            ("eps_real", float(d["eps"].real)),
            ("eps_imag", float(d["eps"].imag)),
            ("Q", float(d["Q"])),
            ("Lx_over_lambda_spp", float(d["cyc"])),
            ("overdamped", bool(d["overdamped"])),
            ("Lx_intrinsic_um", _xlsx_value(d["Lint"] / 1000)),
            ("note", "R_min follows theta_min_definition; only an above-critical valid SPR "
                     "minimum may be interpreted using critical-coupling language."),
        ]

        theta_rows = []
        rmin_rows = []
        lx_rows = []
        for i, thickness in enumerate(d["d"]):
            theta = _xlsx_value(d["theta_min"][i])
            fwhm = _xlsx_value(d["fwhm"][i])
            band_low = theta - fwhm / 2 if theta is not None and fwhm is not None else None
            band_high = theta + fwhm / 2 if theta is not None and fwhm is not None else None
            theta_rows.append([
                thickness, theta, fwhm, band_low, band_high,
                d["dip_ok"][i], d["valid"][i],
            ])
            rmin_rows.append([thickness, d["R_min"][i], d["dip_ok"][i], d["valid"][i]])
            lx_rows.append([thickness, d["Lx"][i] / 1000, d["Lint"] / 1000])

        ws_theta = wb.active
        ws_theta.title = "Theta_min"
        _write_xlsx_sheet(
            ws_theta,
            ["d_nm", "theta_min_deg", "FWHM_deg", "band_low_deg", "band_high_deg",
             "dip_measurable", "valid_SPR"],
            theta_rows,
            common_meta + [("sheet_note",
                            "theta_min is retained across the full scan; use dip_measurable and "
                            "valid_SPR to identify unreliable minima.")],
        )
        ws_rmin = wb.create_sheet("R_min")
        _write_xlsx_sheet(
            ws_rmin,
            ["d_nm", "R_min", "dip_measurable", "valid_SPR"],
            rmin_rows,
            common_meta,
        )
        ws_lx = wb.create_sheet("Lx")
        _write_xlsx_sheet(
            ws_lx,
            ["d_nm", "Lx_leaky_um", "Lx_intrinsic_um"],
            lx_rows,
            common_meta,
        )

    wb.active = 0
    wb.save(path)


def _json_value(value):
    """Recursively convert values to strict JSON types; encode NaN and infinity as null."""
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [_json_value(v) for v in value]
    if isinstance(value, complex):
        return {"real": _json_value(value.real), "imag": _json_value(value.imag)}
    return _xlsx_value(value)


def _json_common_metadata(d, export_type, paired_xlsx):
    """Machine-readable parameters and units for reproducible JSON output."""
    return {
        "schema_version": "1.1",
        "generator": "SPRM Reflectivity Calculator",
        "script": "SPRM_Reflectivity_GUI.py",
        "export_type": export_type,
        "generated": datetime.datetime.now().astimezone().isoformat(timespec="seconds"),
        "paired_xlsx": Path(paired_xlsx).name,
        "parameters": {
            "wavelength_nm": float(d["lam"]),
            "prism": {"n": float(d["n_p"]), "eps": float(d["n_p"] ** 2)},
            "medium": {"n": float(d["n_d"]), "eps": float(d["n_d"] ** 2)},
            "critical_angle_deg": float(d["tc"]),
            "angle_convention": "internal angle at metal interface (glass side)",
            "model": "prism / metal / medium; no adhesion layer; no roughness",
            "time_convention": "exp(-i*omega*t); eps_imag > 0 means absorption",
        },
    }


def save_results_json(path, tag, d, paired_xlsx=None):
    """Save strict JSON paired with the XLSX file for recalculation and audit."""
    path = Path(path)
    paired_xlsx = Path(paired_xlsx) if paired_xlsx else path.with_suffix(".xlsx")

    if tag == "tab1":
        payload = _json_common_metadata(
            d, "Tab A - R(theta) multi-material comparison", paired_xlsx)
        metals = list(d["curves"].keys())
        curve_columns = ["theta_deg"] + [f"R_{m}" for m in metals]
        curve_records = []
        for i, theta in enumerate(d["theta"]):
            record = {"theta_deg": _json_value(theta)}
            record.update({f"R_{m}": _json_value(d["curves"][m][i]) for m in metals})
            curve_records.append(record)

        summary_columns = [
            "material", "d_nm", "eps_real", "eps_imag", "source", "Q",
            "Lx_over_lambda_spp", "theta_curve_min_deg", "R_curve_min",
            "theta_min_deg", "R_min", "FWHM_deg",
            "Lx_leaky_um", "Lx_intrinsic_um", "dip_measurable", "overdamped",
            "valid_SPR", "note",
        ]
        summary_records = []
        for r in d["rows"]:
            summary_records.append(_json_value({
                "material": r["metal"],
                "d_nm": r["d"],
                "eps_real": r["eps"].real,
                "eps_imag": r["eps"].imag,
                "source": r["src"],
                "Q": r["Q"],
                "Lx_over_lambda_spp": r["cyc"],
                "theta_curve_min_deg": r["curve_theta_min"],
                "R_curve_min": r["curve_R_min"],
                "theta_min_deg": r["theta_min"],
                "R_min": r["R_min"],
                "FWHM_deg": r["fwhm"],
                "Lx_leaky_um": r["Lx"] / 1000,
                "Lx_intrinsic_um": r["Lint"] / 1000,
                "dip_measurable": r["dip_ok"],
                "overdamped": r["overdamped"],
                "valid_SPR": r["valid"],
                "note": r["note"],
            }))
        payload["sheet_order"] = ["R_vs_theta", "Summary"]
        payload["sheets"] = {
            "R_vs_theta": {"columns": curve_columns, "records": curve_records},
            "Summary": {"columns": summary_columns, "records": summary_records},
        }
    else:
        payload = _json_common_metadata(d, "Tab B - thickness scan", paired_xlsx)
        payload["material"] = _json_value({
            "name": d["metal"],
            "source": d["src"],
            "eps": {"real": d["eps"].real, "imag": d["eps"].imag},
            "Q": d["Q"],
            "Lx_over_lambda_spp": d["cyc"],
            "overdamped": d["overdamped"],
            "Lx_intrinsic_um": d["Lint"] / 1000,
        })
        payload["parameters"]["thickness_axis_scale"] = d.get(
            "thickness_axis_scale", "linear")
        payload["parameters"]["thickness_sampling"] = "linear step in d_nm"
        payload["parameters"]["theta_min_definition"] = d["theta_definition"]
        payload["parameters"]["theta_scan_start_deg"] = d.get("theta_scan_start")
        payload["parameters"]["theta_scan_stop_deg"] = d.get("theta_scan_stop")
        payload["parameters"]["theta_scan_step_deg"] = d.get("theta_scan_step")
        theta_columns = [
            "d_nm", "theta_min_deg", "FWHM_deg", "band_low_deg", "band_high_deg",
            "dip_measurable", "valid_SPR",
        ]
        rmin_columns = ["d_nm", "R_min", "dip_measurable", "valid_SPR"]
        lx_columns = ["d_nm", "Lx_leaky_um", "Lx_intrinsic_um"]
        theta_records, rmin_records, lx_records = [], [], []
        for i, thickness in enumerate(d["d"]):
            theta = _xlsx_value(d["theta_min"][i])
            fwhm = _xlsx_value(d["fwhm"][i])
            band_low = theta - fwhm / 2 if theta is not None and fwhm is not None else None
            band_high = theta + fwhm / 2 if theta is not None and fwhm is not None else None
            theta_records.append(_json_value({
                "d_nm": thickness,
                "theta_min_deg": theta,
                "FWHM_deg": fwhm,
                "band_low_deg": band_low,
                "band_high_deg": band_high,
                "dip_measurable": d["dip_ok"][i],
                "valid_SPR": d["valid"][i],
            }))
            rmin_records.append(_json_value({
                "d_nm": thickness,
                "R_min": d["R_min"][i],
                "dip_measurable": d["dip_ok"][i],
                "valid_SPR": d["valid"][i],
            }))
            lx_records.append(_json_value({
                "d_nm": thickness,
                "Lx_leaky_um": d["Lx"][i] / 1000,
                "Lx_intrinsic_um": d["Lint"] / 1000,
            }))
        payload["sheet_order"] = ["Theta_min", "R_min", "Lx"]
        payload["sheets"] = {
            "Theta_min": {"columns": theta_columns, "records": theta_records},
            "R_min": {"columns": rmin_columns, "records": rmin_records},
            "Lx": {"columns": lx_columns, "records": lx_records},
        }
        payload["notes"] = [
            "theta_min_deg follows parameters.theta_min_definition and is retained across the "
            "full scan; use dip_measurable and valid_SPR to identify unreliable minima.",
            "For the Tab A sampled-curve definition, theta_scan_* records the exact angle grid.",
            "R_min follows theta_min_definition; only an above-critical valid SPR minimum may "
            "be interpreted using critical-coupling language.",
            "Overdamped measurable minima are retained but valid_SPR is false.",
        ]

    with path.open("w", encoding="utf-8", newline="\n") as fh:
        json.dump(_json_value(payload), fh, ensure_ascii=False, indent=2, allow_nan=False)
        fh.write("\n")


# ============================================================================
#  GUI
# ============================================================================

class SPRApp(tk.Tk):

    def __init__(self):
        super().__init__()
        self.title("SPRM Reflectivity Calculator — Kretschmann geometry (three-layer p-polarized TMM)")
        self.geometry("1280x860")
        self.last = {"tab1": None, "tab2": None}     # Cache the latest results for export

        self._build_global()
        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=8, pady=(0, 4))
        self.tab1 = ttk.Frame(nb); nb.add(self.tab1, text="  A · R(θ) material comparison  ")
        self.tab2 = ttk.Frame(nb); nb.add(self.tab2, text="  B · θmin / R_min / Lx vs thickness  ")
        self.nb = nb
        nb.bind("<<NotebookTabChanged>>", lambda e: self._refresh_n())
        self._build_tab1()
        self._build_tab2()
        self._refresh_n()          # Refresh after tab creation so availability uses each selected metal data source

        self.status = tk.Text(self, height=7, wrap="none",
                              font=("Consolas", 9), bg="#f7f7f7")
        self.status.pack(fill="x", padx=8, pady=(0, 8))
        self._log(
            "Ready. θ is the internal incidence angle at the metal interface (glass side).\n"
            "Select a preset laser wavelength. HeNe is 632.8 nm, not 633 nm;"
            "a 0.2 nm shift changes θmin by about 4 mdeg (four times the numerical precision).\n"
            "Metal suitability updates with wavelength: Q = |ε′|/ε″ < 3 indicates overdamping;"
            "its reflectivity minimum is not a plasmon resonance but an interband absorption edge.")

    # ---------------- Global parameters ----------------
    def _build_global(self):
        f = ttk.LabelFrame(self, text=" Global parameters ")
        f.pack(fill="x", padx=8, pady=8)

        # Group controls compactly to prevent excessive horizontal spacing.
        controls = ttk.Frame(f)
        controls.pack(anchor="w", padx=10, pady=(3, 1))
        left = ttk.Frame(controls)
        left.grid(row=0, column=0, sticky="nw")
        right = ttk.Frame(controls)
        right.grid(row=0, column=1, sticky="nw", padx=(24, 0))

        # --- Left group, row 1: excitation wavelength ---
        ttk.Label(left, text="Excitation wavelength:").grid(row=0, column=0, sticky="e", padx=(0, 4), pady=4)
        self.v_wl = tk.StringVar(value=WL_DEFAULT)
        cbw = ttk.Combobox(left, textvariable=self.v_wl, values=WAVELENGTHS,
                           width=24, state="readonly")
        cbw.grid(row=0, column=1, padx=2)
        cbw.bind("<<ComboboxSelected>>", lambda e: self._refresh_n())
        self.v_lam = tk.StringVar(value="632.8")
        self.e_lam = ttk.Entry(left, textvariable=self.v_lam, width=8, state="readonly")
        self.e_lam.grid(row=0, column=2, padx=(8, 2))
        self.e_lam.bind("<Return>", lambda e: self._refresh_n())
        self.e_lam.bind("<FocusOut>", lambda e: self._refresh_n())
        ttk.Label(left, text="nm").grid(row=0, column=3, sticky="w")

        # --- Right group, row 1: prism ---
        ttk.Label(right, text="Prism / incidence side:").grid(row=0, column=0, sticky="e", padx=(0, 4), pady=4)
        self.v_prism = tk.StringVar(value="SF10")
        cbp = ttk.Combobox(right, textvariable=self.v_prism, values=PRISM_LIST,
                           width=22, state="readonly")
        cbp.grid(row=0, column=1, padx=2)
        cbp.bind("<<ComboboxSelected>>", lambda e: self._refresh_n())
        self.v_np = tk.StringVar(value="1.723087")
        self.e_np = ttk.Entry(right, textvariable=self.v_np, width=9)
        self.e_np.grid(row=0, column=2, padx=(8, 2))
        ttk.Label(right, text="n_p").grid(row=0, column=3, sticky="w")

        # --- Left group, row 2: theta scan range ---
        ttk.Label(left, text="θ range (°):").grid(row=1, column=0, sticky="e", padx=(0, 4), pady=4)
        fr = ttk.Frame(left); fr.grid(row=1, column=1, columnspan=3, sticky="w", padx=2)
        self.v_t0 = tk.StringVar(value="40")
        self.v_t1 = tk.StringVar(value="85")
        self.v_dt = tk.StringVar(value="0.01")
        ttk.Entry(fr, textvariable=self.v_t0, width=6).pack(side="left")
        ttk.Label(fr, text=" — ").pack(side="left")
        ttk.Entry(fr, textvariable=self.v_t1, width=6).pack(side="left")
        ttk.Label(fr, text="   step ").pack(side="left")
        ttk.Entry(fr, textvariable=self.v_dt, width=6).pack(side="left")

        # --- Right group, row 2: ambient medium ---
        ttk.Label(right, text="Ambient medium:").grid(row=1, column=0, sticky="e", padx=(0, 4), pady=4)
        self.v_med = tk.StringVar(value="Water 20.0°C")
        cbm = ttk.Combobox(right, textvariable=self.v_med, values=MEDIUM_LIST,
                           width=22, state="readonly")
        cbm.grid(row=1, column=1, padx=2)
        cbm.bind("<<ComboboxSelected>>", lambda e: self._refresh_n())
        self.v_nd = tk.StringVar(value="1.332106")
        self.e_nd = ttk.Entry(right, textvariable=self.v_nd, width=9)
        self.e_nd.grid(row=1, column=2, padx=(8, 2))
        ttk.Label(right, text="n_d").grid(row=1, column=3, sticky="w")

        # --- Critical-angle summary ---
        self.lbl_tc = ttk.Label(f, text="", foreground="#0a6d00")
        self.lbl_tc.pack(anchor="w", padx=10, pady=(2, 0))

        # --- Metal suitability at this wavelength ---
        vf = ttk.Frame(f)
        vf.pack(anchor="w", padx=10, pady=(5, 7))
        ttk.Label(vf, text="Metal suitability at this wavelength:", font=("", 9, "bold"))\
            .grid(row=0, column=0, rowspan=2, sticky="nw", padx=(0, 5), pady=2)
        self.lbl_q = {}
        for i, m in enumerate(METALS):
            lb = tk.Label(vf, text=f" {m} — ", font=("Consolas", 9),
                          padx=5, pady=1, bd=1, relief="solid")
            row = 0 if i < 3 else 1
            col = i + 1 if i < 3 else i - 2
            lb.grid(row=row, column=col, sticky="w", padx=3, pady=1)
            self.lbl_q[m] = lb
        self.lbl_note = ttk.Label(vf, text="", foreground="#555", font=("", 8))
        self.lbl_note.grid(row=1, column=3, sticky="w", padx=(12, 0))

        self._refresh_n()

    def _cur_lam(self):
        w = parse_wl(self.v_wl.get())
        return w if w is not None else float(self.v_lam.get())

    def _refresh_n(self):
        """Refresh refractive indices, critical angle, and metal suitability when inputs change."""
        # --- wavelength ---
        w = parse_wl(self.v_wl.get())
        if w is None:                                   # Custom wavelength: enable manual entry
            self.e_lam.config(state="normal")
        else:
            self.v_lam.set(f"{w:g}")
            self.e_lam.config(state="readonly")
        try:
            lam = float(self.v_lam.get())
        except ValueError:
            return

        # --- Prism and medium refractive indices ---
        p, m = self.v_prism.get(), self.v_med.get()
        err = ""
        try:
            if p.startswith("Custom"):
                self.e_np.config(state="normal")
            else:
                self.v_np.set(f"{sellmeier_n(p, lam):.6f}")
                self.e_np.config(state="readonly")
            if m.startswith("Custom"):
                self.e_nd.config(state="normal")
            elif m.startswith("Vacuum"):
                self.v_nd.set("1.000000"); self.e_nd.config(state="readonly")
            else:
                self.v_nd.set(f"{sellmeier_n(m, lam):.6f}"); self.e_nd.config(state="readonly")
        except ValueError as ex:
            err = str(ex)

        # --- critical angle ---
        try:
            n_p, n_d = float(self.v_np.get()), float(self.v_nd.get())
            tc = critical_angle(n_p ** 2, n_d ** 2)
            self.lbl_tc.config(
                text=f"λ = {lam:.1f} nm   n_p = {n_p:.6f}   n_d = {n_d:.6f}   "
                     f"critical angle θc = {tc:.3f}°   (θmin search is above θc)"
                     + (f"    ⚠ {err}" if err else ""),
                foreground="#b00000" if err else "#0a6d00")
            ed = n_d ** 2
        except Exception:
            ed = 1.777

        # --- Evaluate each metal's SPP suitability at the current wavelength ---
        tab2_active = False
        try:
            tab2_active = (hasattr(self, "v_m2") and self.nb.index(self.nb.select()) == 1)
        except Exception:
            pass
        tab2_m = self.v_m2.get() if tab2_active else None

        n_ok = 0
        for mt in METALS:
            try:
                if tab2_active and mt == tab2_m:
                    # For Tab B's current metal, use its selected data source so the displayed Q matches the calculation.
                    src = self.v_src2.get()
                else:
                    src = (self.m_src[mt].get() if hasattr(self, "m_src") and mt in self.m_src
                           else list(NK_DATA[mt])[0])
                e = eps_metal(mt, src, lam)
                Q = spp_quality(e)
                cyc = spp_cycles(e, ed)
                ok = Q >= Q_OVERDAMPED
                n_ok += ok
                self.lbl_q[mt].config(
                    text=f" {mt}  Q={Q:5.1f}  Lx/λspp={cyc:5.1f}  {'✓' if ok else '✗ overdamped'} ",
                    fg="#0a6d00" if ok else "#b00000",
                    bg="#e9f7e9" if ok else "#fdeded")
            except Exception:
                self.lbl_q[mt].config(text=f" {mt}  data unavailable ", fg="#777", bg="#f0f0f0")
        note = ("A usable SPP resonance requires Q = |ε′|/ε″ ≥ 3 and Lx/λspp ≳ 2"
                if n_ok else "⚠ No metal supports a usable SPP at this wavelength")
        if tab2_active and tab2_m in METALS:
            note += f"  |  Current {tab2_m}: {self.v_src2.get()}"
        self.lbl_note.config(text=note)

        # wavelengthchanges also refresh Tab B's read-only ε display.
        if hasattr(self, "v_m2"):
            self._refresh_eps2_display()

    def _params(self):
        lam = self._cur_lam()
        self._refresh_n()
        np_ = float(self.v_np.get()); nd = float(self.v_nd.get())
        if np_ <= nd:
            raise ValueError(f"n_prism ({np_}) must exceed n_medium ({nd}); otherwise there is no total internal reflection")
        if not (NK_LO <= lam <= NK_HI):
            raise ValueError(f"wavelengthmust lie within {NK_LO:.0f}–{NK_HI:.0f} nm (built-in metal data range)")
        return lam, np_ ** 2, nd ** 2, np_, nd

    def _log(self, s, clear=True):
        if clear:
            self.status.delete("1.0", "end")
        self.status.insert("end", s + "\n")
        self.status.see("end")

    # ---------------- Tab A ----------------
    def _build_tab1(self):
        left = ttk.Frame(self.tab1); left.pack(side="left", fill="y", padx=6, pady=6)
        ttk.Label(left, text="Metal / thickness (nm) / data source",
                  font=("", 9, "bold")).grid(row=0, column=0, columnspan=3, pady=(0, 6))

        self.m_on, self.m_d, self.m_src = {}, {}, {}
        for i, m in enumerate(METALS):
            self.m_on[m] = tk.BooleanVar(value=(m in ("Au", "Ag", "Cu")))
            ttk.Checkbutton(left, text=m, variable=self.m_on[m], width=4)\
                .grid(row=i + 1, column=0, sticky="w")
            self.m_d[m] = tk.StringVar(value="47")
            ttk.Entry(left, textvariable=self.m_d[m], width=7)\
                .grid(row=i + 1, column=1, padx=3)
            srcs = list(NK_DATA[m].keys())
            self.m_src[m] = tk.StringVar(value=srcs[0])
            cbs = ttk.Combobox(left, textvariable=self.m_src[m], values=srcs, width=22,
                               state="readonly")
            cbs.grid(row=i + 1, column=2, padx=3, pady=1)
            cbs.bind("<<ComboboxSelected>>",
                     lambda e, metal=m: self._on_src1_changed(metal))

        # --- Manual ε row: published values or fitted ellipsometry values ---
        # Default to the fitted Au values in Nat. Rev. Methods Primers Fig. 4b (-12.900 + 1.35i, d=48.0 nm);
        # these effective values include the 1.2 nm Cr adhesion layer and surface roughness. Comparing them with J&C
        # illustrates the departure of the real device from the ideal model.
        rc = len(METALS) + 1
        self.m_on[CUSTOM] = tk.BooleanVar(value=False)
        ttk.Checkbutton(left, text="Manual ε", variable=self.m_on[CUSTOM], width=7,
                        command=self._toggle_custom_inputs)\
            .grid(row=rc, column=0, sticky="w")
        self.m_d[CUSTOM] = tk.StringVar(value="48.0")
        self.e_custom_d = ttk.Entry(left, textvariable=self.m_d[CUSTOM], width=7,
                                    state="disabled")
        self.e_custom_d.grid(row=rc, column=1, padx=3)
        cf = ttk.Frame(left); cf.grid(row=rc, column=2, sticky="w", padx=3, pady=1)
        ttk.Label(cf, text="ε′").pack(side="left")
        self.c_e1 = tk.StringVar(value="-12.900")
        self.e_custom_e1 = ttk.Entry(cf, textvariable=self.c_e1, width=8,
                                     state="disabled")
        self.e_custom_e1.pack(side="left", padx=(2, 6))
        ttk.Label(cf, text="ε″").pack(side="left")
        self.c_e2 = tk.StringVar(value="1.350")
        self.e_custom_e2 = ttk.Entry(cf, textvariable=self.c_e2, width=7,
                                     state="disabled")
        self.e_custom_e2.pack(side="left", padx=2)

        r = len(METALS) + 2
        ttk.Separator(left, orient="horizontal").grid(row=r, column=0, columnspan=3,
                                                      sticky="ew", pady=8)
        ttk.Label(left, text="Common thickness:").grid(row=r + 1, column=0, sticky="e")
        self.v_all = tk.StringVar(value="47")
        ttk.Entry(left, textvariable=self.v_all, width=7).grid(row=r + 1, column=1)
        ttk.Button(left, text="Apply to all", command=self._sync_d)\
            .grid(row=r + 1, column=2, sticky="w", padx=3)

        ttk.Button(left, text="Calculate", command=self.run_tab1)\
            .grid(row=r + 2, column=0, columnspan=3, sticky="ew", pady=(12, 3))
        ttk.Button(left, text="Export data (XLSX + JSON)…", command=lambda: self.export_xlsx("tab1"))\
            .grid(row=r + 3, column=0, columnspan=3, sticky="ew", pady=2)
        ttk.Button(left, text="Export figure…", command=lambda: self.export_fig("tab1"))\
            .grid(row=r + 4, column=0, columnspan=3, sticky="ew", pady=2)

        right = ttk.Frame(self.tab1); right.pack(side="left", fill="both", expand=True)
        self.fig1 = Figure(figsize=(8.4, 5.6), dpi=100)
        self.ax1 = self.fig1.add_subplot(111)
        self.cv1 = FigureCanvasTkAgg(self.fig1, right)
        self.cv1.get_tk_widget().pack(fill="both", expand=True)
        NavigationToolbar2Tk(self.cv1, right).update()

    def _sync_d(self):
        for m in METALS:
            self.m_d[m].set(self.v_all.get())

    def _on_src1_changed(self, metal):
        """Share the selected data source for each built-in metal between Tabs A and B."""
        if hasattr(self, "v_m2") and self.v_m2.get() == metal:
            self.v_src2.set(self.m_src[metal].get())
        self._refresh_n()

    def _toggle_custom_inputs(self):
        """Enable independent thickness and permittivity edits only for Manual ε."""
        state = "normal" if self.m_on[CUSTOM].get() else "disabled"
        for widget in (self.e_custom_d, self.e_custom_e1, self.e_custom_e2):
            widget.config(state=state)

    def run_tab1(self):
        try:
            lam, ep, ed, n_p, n_d = self._params()
            sel = [m for m in METALS if self.m_on[m].get()]
            if self.m_on[CUSTOM].get():
                sel.append(CUSTOM)
            if not sel:
                raise ValueError("Select at least one material")
            t0, t1 = float(self.v_t0.get()), float(self.v_t1.get())
            dt = float(self.v_dt.get())
            if t1 <= t0 or dt <= 0:
                raise ValueError("Invalid θ range or step")
            th = np.arange(t0, t1 + dt * 0.5, dt)
            tc = critical_angle(ep, ed)

            self.ax1.clear()
            curves, rows = {}, []
            lines = [f"λ = {lam:.1f} nm   prism n_p = {n_p:.6f}   medium n_d = {n_d:.6f}"
                     f"   θc = {tc:.3f}°",
                     f"{'Material':<7}{'d(nm)':>7}{'ε′':>9}{'ε″':>8}{'Q':>7}{'Lx/λspp':>9}"
                     f"{'θcurve(°)':>11}{'θSPR(°)':>10}{'R_SPR':>9}{'FWHM(°)':>9}"
                     f"{'Lx_leaky(µm)':>13}{'Lx_int(µm)':>11}  Notes"]
            for m in sel:
                d = float(self.m_d[m].get())
                if m == CUSTOM:
                    e1, e2 = float(self.c_e1.get()), float(self.c_e2.get())
                    if e2 <= 0:
                        raise ValueError("Manual ε″ must be > 0 (with exp(-iωt), ε″ > 0 indicates absorption)")
                    em, src = complex(e1, e2), "Manual input (literature/ellipsometry fit)"
                else:
                    src = self.m_src[m].get()
                    em = eps_metal(m, src, lam)
                R = reflectivity_p([ep, em, ed], [d], lam, th)
                curves[m] = R
                curve_dip = find_curve_minimum(ep, em, ed, d, lam, th)
                dip = find_dip(ep, em, ed, d, lam)
                Lx, Lint = lx_leaky_sweep(ep, em, ed, [d], lam)
                Q = spp_quality(em)
                cyc = spp_cycles(em, ed)
                c, ls = STYLE[m]
                lab = f"{m}  {d:g} nm  |  {src}"
                if dip["overdamped"]:
                    lab += f"  [overdamped, Lx/$\\lambda_{{spp}}$={cyc:.1f}]"
                elif not dip["dip_ok"]:
                    lab += "  [no usable dip]"
                self.ax1.plot(th, R, color=c, ls=ls, lw=1.7, label=lab)
                self.ax1.plot(curve_dip["theta_min"], curve_dip["R_min"], marker="v",
                              ls="none", ms=6.5, mec=c, mew=1.2,
                              mfc=c if curve_dip["valid"] else "none", zorder=5)
                if t0 <= dip["theta_min"] <= t1:
                    self.ax1.plot(dip["theta_min"], dip["R_min"], marker="o",
                                  ls="none", ms=6, mec=c, mew=1.2, mfc="none", zorder=5)
                lines.append(
                    f"{m:<7}{d:>7.1f}{em.real:>9.2f}{em.imag:>8.2f}{Q:>7.1f}{cyc:>9.1f}"
                    f"{curve_dip['theta_min']:>11.3f}{dip['theta_min']:>10.3f}"
                    f"{dip['R_min']:>9.4f}"
                    f"{dip['fwhm']:>9.2f}"
                    f"{(Lx[0]/1000 if np.isfinite(Lx[0]) else np.nan):>13.2f}"
                    f"{(Lint/1000 if np.isfinite(Lint) else np.nan):>11.2f}"
                    f"  {dip['note']}\n"
                    f"         source: {src}; curve minimum: {curve_dip['note'] or 'measurable'}")
                rows.append(dict(metal=m, d=d, eps=em, src=src, cyc=cyc,
                                 curve_theta_min=curve_dip["theta_min"],
                                 curve_R_min=curve_dip["R_min"],
                                 curve_fwhm=curve_dip["fwhm"],
                                 curve_dip_ok=curve_dip["dip_ok"],
                                 curve_valid=curve_dip["valid"],
                                 curve_note=curve_dip["note"],
                                 **dip, Lx=Lx[0], Lint=Lint, Q=Q))

            if t0 < tc:
                self.ax1.axvspan(t0, min(tc, t1), color="0.85", alpha=0.45, zorder=0,
                                 label="Below TIR — not SPR search region")
            self.ax1.axvline(tc, color="0.55", lw=1.0, ls=(0, (6, 4)))
            self.ax1.annotate(f"$\\theta_c$ = {tc:.2f}°", xy=(tc, 0.03),
                              xytext=(4, 0), textcoords="offset points",
                              color="0.35", fontsize=8, rotation=90, va="bottom")
            self.ax1.set_xlabel("Internal angle of incidence  $\\theta$  (deg)")
            self.ax1.set_ylabel("Reflectivity  $R_p$")
            self.ax1.set_title(f"Kretschmann SPR  |  $\\lambda$ = {lam:.1f} nm,  "
                               f"$n_{{prism}}$ = {n_p:.4f},  $n_{{medium}}$ = {n_d:.4f}")
            self.ax1.set_xlim(t0, t1); self.ax1.set_ylim(-0.02, 1.05)
            self.ax1.grid(alpha=0.25, lw=0.5)
            handles, labels = self.ax1.get_legend_handles_labels()
            handles.extend([
                Line2D([], [], marker="v", ls="none", color="0.2", mfc="0.2", ms=6,
                       label="Tab A sampled curve minimum"),
                Line2D([], [], marker="o", ls="none", color="0.2", mfc="none", ms=6,
                       label=r"Refined SPR-region minimum ($\theta > \theta_c$)"),
            ])
            labels.extend(["Tab A sampled curve minimum",
                           r"Refined SPR-region minimum ($\theta > \theta_c$)"])
            self.ax1.legend(handles, labels, loc="lower right", framealpha=0.92,
                            fontsize=7.3)
            self.fig1.tight_layout(); self.cv1.draw()

            self.last["tab1"] = dict(theta=th, curves=curves, rows=rows, lam=lam,
                                     n_p=n_p, n_d=n_d, tc=tc)
            self._log("\n".join(lines))
        except Exception as e:
            messagebox.showerror("Calculation failed", str(e))

    # ---------------- Tab B ----------------
    def _build_tab2(self):
        left = ttk.Frame(self.tab2); left.pack(side="left", fill="y", padx=6, pady=6)
        ttk.Label(left, text="Material", font=("", 9, "bold")).grid(row=0, column=0,
                                                                sticky="w", pady=(0, 4))
        self.v_m2 = tk.StringVar(value="Au")
        cb = ttk.Combobox(left, textvariable=self.v_m2, values=METALS + [CUSTOM],
                          width=8, state="readonly")
        cb.grid(row=0, column=1, sticky="w")
        cb.bind("<<ComboboxSelected>>", lambda e: self._sync_src2())

        ttk.Label(left, text="Data source").grid(row=1, column=0, sticky="w", pady=4)
        self.v_src2 = tk.StringVar(value=list(NK_DATA["Au"].keys())[0])
        self.cb_src2 = ttk.Combobox(left, textvariable=self.v_src2,
                                    values=list(NK_DATA["Au"].keys()),
                                    width=24, state="readonly")
        self.cb_src2.grid(row=1, column=1, sticky="w")
        self.cb_src2.bind("<<ComboboxSelected>>", lambda e: self._on_src2_changed())

        cf2 = ttk.Frame(left); cf2.grid(row=2, column=0, columnspan=2, sticky="w", pady=3)
        self.lbl_eps2_mode = ttk.Label(cf2, text="Current ε")
        self.lbl_eps2_mode.pack(side="left", padx=(0, 5))
        ttk.Label(cf2, text="ε′").pack(side="left")
        # Keep custom ε inputs separate from the read-only values of built-in metals.
        self.c2_e1 = tk.StringVar(value="-12.900")
        self.c2_e2 = tk.StringVar(value="1.350")
        self.v2_eps_e1 = tk.StringVar(value="")
        self.v2_eps_e2 = tk.StringVar(value="")
        self.e2_e1 = ttk.Entry(cf2, textvariable=self.v2_eps_e1, width=8, state="readonly")
        self.e2_e1.pack(side="left", padx=(2, 6))
        ttk.Label(cf2, text="ε″").pack(side="left")
        self.e2_e2 = ttk.Entry(cf2, textvariable=self.v2_eps_e2, width=7, state="readonly")
        self.e2_e2.pack(side="left", padx=2)
        self.lbl_eps2_state = ttk.Label(cf2, text="(read-only)", foreground="#666")
        self.lbl_eps2_state.pack(side="left", padx=(3, 0))

        ttk.Label(left, text="θmin definition").grid(row=3, column=0, sticky="w", pady=4)
        self.v_theta_def = tk.StringVar(value=THETA_DEF_CURVE)
        ttk.Combobox(left, textvariable=self.v_theta_def, values=THETA_DEFINITIONS,
                     width=31, state="readonly").grid(row=3, column=1, sticky="w")

        ttk.Separator(left, orient="horizontal").grid(row=4, column=0, columnspan=2,
                                                      sticky="ew", pady=8)
        ttk.Label(left, text="Thickness range (nm)", font=("", 9, "bold"))\
            .grid(row=5, column=0, columnspan=2, sticky="w")
        self.v_d0 = tk.StringVar(value="20")
        self.v_d1 = tk.StringVar(value="90")
        self.v_dd = tk.StringVar(value="0.5")
        for i, (t, v) in enumerate([("Start", self.v_d0), ("Stop", self.v_d1), ("step", self.v_dd)]):
            ttk.Label(left, text=t).grid(row=6 + i, column=0, sticky="e", padx=2)
            ttk.Entry(left, textvariable=v, width=9).grid(row=6 + i, column=1, sticky="w")

        self.v_dlog = tk.BooleanVar(value=False)
        ttk.Checkbutton(left, text="Use logarithmic thickness axis", variable=self.v_dlog)\
            .grid(row=9, column=0, columnspan=2, sticky="w", pady=(5, 0))

        ttk.Button(left, text="Calculate", command=self.run_tab2)\
            .grid(row=10, column=0, columnspan=2, sticky="ew", pady=(12, 3))
        ttk.Button(left, text="Export data (XLSX + JSON)…", command=lambda: self.export_xlsx("tab2"))\
            .grid(row=11, column=0, columnspan=2, sticky="ew", pady=2)
        ttk.Button(left, text="Export figure…", command=lambda: self.export_fig("tab2"))\
            .grid(row=12, column=0, columnspan=2, sticky="ew", pady=2)

        ttk.Label(left, text="Unmeasurable θmin values use\ngray dotted lines; overdamped\ncurves are not quantitative.",
                  foreground="#a33", font=("", 8)).grid(row=13, column=0, columnspan=2,
                                                        sticky="w", pady=(14, 0))

        right = ttk.Frame(self.tab2); right.pack(side="left", fill="both", expand=True)
        self.fig2 = Figure(figsize=(8.4, 6.4), dpi=100)
        self.ax2 = self.fig2.subplots(3, 1, sharex=True)
        self.cv2 = FigureCanvasTkAgg(self.fig2, right)
        self.cv2.get_tk_widget().pack(fill="both", expand=True)
        NavigationToolbar2Tk(self.cv2, right).update()
        self._refresh_eps2_display()

    def _sync_src2(self):
        m = self.v_m2.get()
        if m == CUSTOM:
            self.cb_src2.config(values=["Manual ε input (at right)"], state="disabled")
            self.v_src2.set("Manual ε input (at right)")
        else:
            srcs = list(NK_DATA[m].keys())
            self.cb_src2.config(values=srcs, state="readonly")
            shared_src = self.m_src[m].get() if hasattr(self, "m_src") else srcs[0]
            self.v_src2.set(shared_src if shared_src in srcs else srcs[0])
        self._refresh_n()

    def _on_src2_changed(self):
        """Refresh the current ε display and quality indicator after a source change."""
        m = self.v_m2.get()
        if m in METALS and hasattr(self, "m_src"):
            self.m_src[m].set(self.v_src2.get())
        self._refresh_n()

    def _refresh_eps2_display(self):
        """Tab B shows read-only ε for built-in metals and editable ε for Custom."""
        m = self.v_m2.get()
        if m == CUSTOM:
            self.lbl_eps2_mode.config(text="Manual ε")
            self.lbl_eps2_state.config(text="(editable)", foreground="#0a6d00")
            self.e2_e1.config(textvariable=self.c2_e1, state="normal")
            self.e2_e2.config(textvariable=self.c2_e2, state="normal")
            return

        self.lbl_eps2_mode.config(text="Current ε")
        self.lbl_eps2_state.config(text="(read-only)", foreground="#666")
        self.e2_e1.config(textvariable=self.v2_eps_e1, state="readonly")
        self.e2_e2.config(textvariable=self.v2_eps_e2, state="readonly")
        try:
            em = eps_metal(m, self.v_src2.get(), self._cur_lam())
            self.v2_eps_e1.set(f"{em.real:.4f}")
            self.v2_eps_e2.set(f"{em.imag:.4f}")
        except Exception:
            self.v2_eps_e1.set("—")
            self.v2_eps_e2.set("—")

    def run_tab2(self):
        try:
            lam, ep, ed, n_p, n_d = self._params()
            m = self.v_m2.get(); src = self.v_src2.get()
            d0, d1 = float(self.v_d0.get()), float(self.v_d1.get())
            dd = float(self.v_dd.get())
            if d1 <= d0 or dd <= 0:
                raise ValueError("Invalid thickness range or step")
            thickness_scale = "log" if self.v_dlog.get() else "linear"
            if thickness_scale == "log" and d0 <= 0:
                raise ValueError("A logarithmic thickness axis requires a starting thickness > 0 nm")
            ds = np.arange(d0, d1 + dd * 0.5, dd)
            if len(ds) > 4000:
                raise ValueError(f"Thickness point count {len(ds)} is too large; increase the step")

            if m == CUSTOM:
                e1, e2 = float(self.c2_e1.get()), float(self.c2_e2.get())
                if e2 <= 0:
                    raise ValueError("Manual ε″ must be > 0 (with exp(-iωt), ε″ > 0 indicates absorption)")
                em = complex(e1, e2)
                src = "Manual input (literature/ellipsometry fit)"
            else:
                em = eps_metal(m, src, lam)
            Q = spp_quality(em)
            tc = critical_angle(ep, ed)
            theta_definition = self.v_theta_def.get()
            if theta_definition not in THETA_DEFINITIONS:
                raise ValueError("Unknown θmin definition")
            theta_scan = None
            if theta_definition == THETA_DEF_CURVE:
                t0, t1 = float(self.v_t0.get()), float(self.v_t1.get())
                dt = float(self.v_dt.get())
                if t1 <= t0 or dt <= 0:
                    raise ValueError("Invalid Tab A θ range or step")
                theta_scan = np.arange(t0, t1 + dt * 0.5, dt)

            th_min = np.full(len(ds), np.nan)
            R_min = np.full(len(ds), np.nan)
            fwhm = np.full(len(ds), np.nan)
            valid = np.zeros(len(ds), bool)
            dip_ok = np.zeros(len(ds), bool)
            for i, d in enumerate(ds):
                if theta_definition == THETA_DEF_CURVE:
                    dip = find_curve_minimum(ep, em, ed, float(d), lam, theta_scan)
                else:
                    dip = find_dip(ep, em, ed, float(d), lam)
                R_min[i] = dip["R_min"]
                fwhm[i] = dip["fwhm"]
                valid[i] = dip["valid"]
                dip_ok[i] = dip["dip_ok"]
                # Retain numerical minima for all thicknesses; separate flags indicate measurability and SPR validity.
                th_min[i] = dip["theta_min"]
            Lx, Lint = lx_leaky_sweep(ep, em, ed, ds, lam)
            cyc = spp_cycles(em, ed)
            od = Q < Q_OVERDAMPED

            c, ls = STYLE[m]
            if od:
                c, ls = "0.45", (0, (4, 2))          # Gray dashed line for overdamped curves
            for a in self.ax2:
                # Restore a linear axis before clearing a previous log plot to avoid an xlim warning.
                a.set_xscale("linear")
                a.clear()
            a0, a1, a2 = self.ax2
            for a in self.ax2:
                a.set_xscale(thickness_scale)

            # Panel 1 separates valid SPR, measurable non-SPR, and unmeasurable minima.
            th_valid = np.where(valid, th_min, np.nan)
            th_non_spr = np.where(dip_ok & ~valid, th_min, np.nan)
            th_unmeasurable = np.where(~dip_ok, th_min, np.nan)
            non_spr_color = "0.45" if od else "#b36b00"
            non_spr_ls = (0, (4, 2)) if od else (0, (5, 2))
            a0.plot(ds, th_valid, color=c, ls=ls, lw=1.8, zorder=4,
                    label=r"$\theta_{min}$ (valid SPR)")
            a0.plot(ds, th_non_spr, color=non_spr_color, ls=non_spr_ls, lw=1.6,
                    zorder=3, label=r"$\theta_{min}$ (measurable, not valid SPR)")
            a0.plot(ds, th_unmeasurable, color="0.55", ls=(0, (1.5, 2.0)), lw=1.5,
                    zorder=2, label=r"$\theta_{min}$ (unmeasurable minimum)")
            band_ok = dip_ok & np.isfinite(fwhm)
            a0.fill_between(ds, th_min - fwhm / 2, th_min + fwhm / 2,
                            where=band_ok, interpolate=True, color=c,
                            alpha=0.18, lw=0, zorder=1, label=r"$\pm$FWHM/2 (measurable)")
            a0.set_ylabel(r"$\theta_{min}$  (deg)")
            a0.grid(alpha=0.25, lw=0.5)
            a0.legend(loc="upper right", fontsize=8, ncol=2)
            if od:
                a0.text(0.5, 0.05,
                        f"OVERDAMPED:  Q = {Q:.2f},   "
                        f"$L_x/\\lambda_{{spp}}$ = {cyc:.1f}\n"
                        f"the minimum is an absorption edge, NOT a plasmon resonance",
                        transform=a0.transAxes, ha="center", va="bottom",
                        fontsize=8.5, color="#a00000", weight="bold",
                        bbox=dict(fc="#fff2f2", ec="#a00000", lw=0.9, alpha=0.95))

            # Panel 2 uses the selected theta definition for R_min(d); the validity flag determines SPR status.
            a1.plot(ds, R_min, color=c, ls=ls, lw=1.8)
            a1.set_ylabel(r"$R_{min}$")
            a1.set_yscale("log")
            a1.grid(alpha=0.25, lw=0.5, which="both")

            # Panel 3: Lx(d)
            a2.plot(ds, Lx / 1000.0, color=c, ls=ls, lw=1.8, label=r"$L_x$ (leaky)")
            if np.isfinite(Lint):
                a2.axhline(Lint / 1000.0, color="0.5", lw=1.0, ls=(0, (6, 4)),
                           label=rf"$L_x^{{int}}$ = {Lint/1000:.2f} µm (upper bound)")
            a2.set_ylabel(r"$L_x$  (µm)")
            a2.set_xlabel("Metal film thickness  $d$  (nm)")
            a2.grid(alpha=0.25, lw=0.5)
            a2.legend(loc="lower right", fontsize=8)

            info = [f"{m} / {src}   λ = {lam:.1f} nm",
                    f"θmin definition = {theta_definition}",
                    f"ε = {em.real:.3f} {em.imag:+.3f}i    Q = |ε′|/ε″ = {Q:.2f}"
                    f"    Lx/λspp = {cyc:.1f}",
                    f"prism n_p = {n_p:.6f}   medium n_d = {n_d:.6f}   θc = {tc:.3f}°",
                    f"Thickness axis = {'logarithmic' if thickness_scale == 'log' else 'linear'} scale"
                    " (thickness points remain linearly sampled at the selected step)",
                    (f"Intrinsic propagation length Lx_int = {Lint/1000:.2f} µm"
                     if np.isfinite(Lint) else
                     "Intrinsic propagation length: no bound SPP (ε′ > −ε_d)")]

            if dip_ok.any():
                j = int(np.nanargmin(np.where(dip_ok, R_min, np.inf)))
                a1.plot(ds[j], R_min[j], "v", color=c, ms=7)
                a0.plot(ds[j], th_min[j], "v", color=c, ms=7)
                optimum_label = ("Candidate critical-coupling thickness d_opt" if theta_definition == THETA_DEF_SPR
                                 else "Thickness of deepest minimum under selected definition d_min")
                info.append(
                    f"{optimum_label} = {ds[j]:.2f} nm  →  R_min = {R_min[j]:.2e}, "
                    f"θmin = {th_min[j]:.3f}°, FWHM = {fwhm[j]:.2f}°, "
                    f"Lx = {Lx[j]/1000:.2f} µm"
                    + (f" ({Lx[j]/Lint*100:.0f}% of Lx_int)" if np.isfinite(Lint) else ""))
            if od:
                info += [
                    "",
                    f"!! Overdamped (Q = {Q:.2f} < {Q_OVERDAMPED:.0f}). The θmin curve is marked as an invalid SPR.",
                    "   Note: a deep dip does not prove a plasmon resonance. A sufficiently thin film can reach critical coupling",
                    "   (Γ_rad = Γ_int) and make R_min approach zero. FWHM and Lx distinguish propagation quality:",
                    f"   → Lx/λspp = {cyc:.1f}  (the SPP travels only {cyc:.1f} of its own wavelengths;"
                    f" < 2 does not constitute a propagating mode)",
                    "   → This minimum is a Fresnel/interband absorption feature and must not be used as a quantitative resonance angle."]
            elif not dip_ok.any():
                info.append("!! No measurable minimum exists over this thickness range; gray dotted θmin values are diagnostic only.")
            elif not dip_ok.all():
                info.append("Gray dotted lines indicate unmeasurable numerical θmin values and are not quantitative resonance angles.")

            definition_short = ("Tab A sampled curve minimum"
                                if theta_definition == THETA_DEF_CURVE else
                                "refined SPR-region minimum")
            a0.set_title(f"{m} ({src})  |  {definition_short}  |  $\\lambda$ = {lam:.1f} nm, "
                         f"$n_p$ = {n_p:.4f}, $n_d$ = {n_d:.4f}", fontsize=9.5)
            self.fig2.tight_layout(); self.cv2.draw()

            self.last["tab2"] = dict(d=ds, theta_min=th_min, R_min=R_min, fwhm=fwhm,
                                     Lx=Lx, valid=valid, dip_ok=dip_ok, Lint=Lint,
                                     metal=m, src=src, eps=em, lam=lam, n_p=n_p,
                                     n_d=n_d, tc=tc, Q=Q, cyc=cyc, overdamped=od,
                                     thickness_axis_scale=thickness_scale,
                                     theta_definition=theta_definition,
                                     theta_scan_start=(float(theta_scan[0])
                                                       if theta_scan is not None else None),
                                     theta_scan_stop=(float(theta_scan[-1])
                                                      if theta_scan is not None else None),
                                     theta_scan_step=(float(np.median(np.diff(theta_scan)))
                                                      if theta_scan is not None else None))
            self._log("\n".join(info))
        except Exception as e:
            messagebox.showerror("Calculation failed", str(e))

    # ---------------- Export ----------------
    def export_xlsx(self, tag):
        if self.last[tag] is None:
            messagebox.showwarning("No data", "Run a calculation first."); return
        p = filedialog.asksaveasfilename(
            title="Export data to…", defaultextension=".xlsx",
            initialfile=("SPR_R_vs_theta.xlsx" if tag == "tab1" else "SPR_thickness_scan.xlsx"),
            filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")])
        if not p:
            return
        try:
            xlsx_path = Path(p)
            if xlsx_path.suffix.lower() != ".xlsx":
                xlsx_path = xlsx_path.with_suffix(".xlsx")
            json_path = xlsx_path.with_suffix(".json")
            save_results_xlsx(xlsx_path, tag, self.last[tag])
            save_results_json(json_path, tag, self.last[tag], paired_xlsx=xlsx_path)
            self._log(f"Exported XLSX: {xlsx_path}\nPaired JSON: {json_path}", clear=False)
        except Exception as e:
            messagebox.showerror("Export failed", str(e))

    def export_fig(self, tag):
        if self.last[tag] is None:
            messagebox.showwarning("No data", "Run a calculation first."); return
        p = filedialog.asksaveasfilename(
            title="Export figure to…", defaultextension=".pdf",
            filetypes=[("PDF (vector)", "*.pdf"), ("SVG (vector)", "*.svg"),
                       ("PNG (600 dpi)", "*.png"), ("All files", "*.*")])
        if not p:
            return
        fig = self.fig1 if tag == "tab1" else self.fig2
        try:
            fig.savefig(p, dpi=600, bbox_inches="tight")
            self._log(f"Exported: {p}", clear=False)
        except Exception as e:
            messagebox.showerror("Export failed", str(e))


if __name__ == "__main__":
    SPRApp().mainloop()
