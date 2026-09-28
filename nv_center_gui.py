#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate an NV-center ODMR splitting figure and export its plotted data.

This companion script reuses the CW-ODMR Hamiltonian and Lorentzian spectrum
implemented in ``nv_odmr_gui.py``. It plots normalized fluorescence versus
microwave frequency for several magnetic fields, with adjustable vertical
offsets and per-curve dip magnitudes.

The plot window exports a reproducible bundle containing PNG (300 dpi), PDF,
SVG, CSV, XLSX, and JSON files sharing the selected filename stem.
"""

from __future__ import annotations

import csv
import json
import re
import traceback
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

import matplotlib

matplotlib.use("TkAgg")
matplotlib.rcParams.update(
    {
        "font.family": "Arial",
        "font.size": 9,
        "axes.labelsize": 10,
        "axes.titlesize": 10,
        "legend.fontsize": 8,
        "xtick.labelsize": 9,
        "ytick.labelsize": 9,
        "axes.linewidth": 0.9,
        "lines.linewidth": 1.8,
        "savefig.dpi": 300,
        "savefig.facecolor": "white",
        "axes.unicode_minus": False,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    }
)
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from matplotlib.figure import Figure

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import nv_odmr_gui as nv
except ImportError as exc:  # A clearer message when this file is moved elsewhere.
    raise ImportError(
        "nv_center_gui.py must remain in the same directory as nv_odmr_gui.py"
    ) from exc


SCRIPT_NAME = "nv_center_gui.py"
ODMR_CMAP = matplotlib.colormaps["viridis"]


@dataclass(frozen=True)
class ODMRConfig:
    """User-facing parameters for the simplified review-box ODMR panel."""

    B_values_mT: tuple[float, ...] = (0.0, 0.5,1.0,1.5,2.0)
    f_min_GHz: float = 2.76
    f_max_GHz: float = 2.98
    n_frequency: int = 1201
    D_GHz: float = 2.870
    strain_E_MHz: float = 1.0
    linewidth_MHz: float = 7.0
    contrast_percent: float = 8.0
    B_angle_deg: float = 0.0
    vertical_offset: float = 0.05
    magnitude_scales: tuple[float, ...] = (0.5, 0.5, 0.5, 0.5, 0.5)


def parse_B_values(text: str) -> tuple[float, ...]:
    """Parse comma, semicolon, or whitespace separated field values."""

    tokens = [item for item in re.split(r"[\s,;]+", text.strip()) if item]
    if not tokens:
        raise ValueError("Enter at least one magnetic-field value.")
    values = sorted({float(item) for item in tokens})
    if len(values) > 12:
        raise ValueError("Use at most 12 magnetic-field curves for a readable overlay.")
    if values[0] < 0.0 or values[-1] > 50.0:
        raise ValueError("Magnetic-field values must be between 0 and 50 mT.")
    return tuple(values)


def parse_magnitude_scales(text: str, curve_count: int) -> tuple[float, ...]:
    """Parse one shared multiplier or one dip-depth multiplier per B curve."""

    tokens = [item for item in re.split(r"[\s,;]+", text.strip()) if item]
    if not tokens:
        return (1.0,) * curve_count
    values = tuple(float(item) for item in tokens)
    if len(values) == 1:
        values = values * curve_count
    if len(values) != curve_count:
        raise ValueError(
            "Dip magnitude multipliers must contain either one value or exactly "
            f"{curve_count} values (one for each B curve)."
        )
    if any(value < 0.0 or value > 5.0 for value in values):
        raise ValueError("Dip magnitude multipliers must be between 0 and 5.")
    return values


def validate_odmr(config: ODMRConfig) -> None:
    if not config.B_values_mT:
        raise ValueError("At least one B value is required.")
    if config.f_min_GHz >= config.f_max_GHz:
        raise ValueError("Frequency maximum must exceed frequency minimum.")
    if not 201 <= config.n_frequency <= 20_001:
        raise ValueError("Frequency points must be between 201 and 20001.")
    if not 2.0 <= config.D_GHz <= 4.0:
        raise ValueError("D must be between 2 and 4 GHz.")
    if not 0.0 <= config.strain_E_MHz <= 100.0:
        raise ValueError("Strain E must be between 0 and 100 MHz.")
    if not 0.05 <= config.linewidth_MHz <= 100.0:
        raise ValueError("Linewidth must be between 0.05 and 100 MHz.")
    if not 0.0 < config.contrast_percent <= 40.0:
        raise ValueError("Contrast must be greater than 0 and at most 40%.")
    if not 0.0 <= config.B_angle_deg <= 90.0:
        raise ValueError("B angle must be between 0 and 90 degrees.")
    if not 0.0 <= config.vertical_offset <= 1.0:
        raise ValueError("Vertical offset must be between 0 and 1.")
    if len(config.magnitude_scales) != len(config.B_values_mT):
        raise ValueError("One dip magnitude multiplier is required per B curve.")
    if any(value < 0.0 or value > 5.0 for value in config.magnitude_scales):
        raise ValueError("Dip magnitude multipliers must be between 0 and 5.")


def simulate_odmr_overlay(config: ODMRConfig) -> dict[str, object]:
    """Run the existing validated ODMR model once for every requested B value."""

    validate_odmr(config)
    spectra: list[np.ndarray] = []
    frequency: np.ndarray | None = None

    for B_mT in config.B_values_mT:
        params = nv.Params(
            B_min=B_mT,
            B_max=B_mT,
            n_B=1,
            theta_B=config.B_angle_deg,
            phi_B=0.0,
            geometry="single",
            auto_f=False,
            f_min=config.f_min_GHz,
            f_max=config.f_max_GHz,
            n_f=config.n_frequency,
            D_gs=config.D_GHz,
            E_gs=config.strain_E_MHz,
            phi_E=0.0,
            nuc="none",
            gamma_fwhm=config.linewidth_MHz,
            mw_ideal_perp=True,
            model="A",
            contrast=config.contrast_percent,
        )
        result = nv.simulate(params)
        if frequency is None:
            frequency = np.asarray(result["f_GHz"], dtype=float)
        spectra.append(np.asarray(result["PL"][0], dtype=float))

    assert frequency is not None
    return {
        "frequency_GHz": frequency,
        "B_mT": np.asarray(config.B_values_mT, dtype=float),
        "PL_norm": np.vstack(spectra),
    }


def displayed_odmr_spectra(data: dict[str, object], config: ODMRConfig) -> np.ndarray:
    """Apply adjustable dip magnitudes and successive vertical baseline offsets.

    Magnitude scales act on the resonance-induced fluorescence reduction rather
    than on the entire signal, so every unshifted physical baseline remains one:

        displayed_i = 1 + i*offset - scale_i*(1 - PL_i).
    """

    physical = np.asarray(data["PL_norm"], dtype=float)
    scale = np.asarray(config.magnitude_scales, dtype=float)[:, None]
    baseline = 1.0 + config.vertical_offset * np.arange(physical.shape[0])[:, None]
    return baseline - scale * (1.0 - physical)


def style_axis(axis) -> None:
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.tick_params(direction="in", length=4.0, width=0.8)
    axis.grid(False)


def draw_odmr(figure: Figure, data: dict[str, object], config: ODMRConfig) -> None:
    figure.clear()
    axis = figure.add_subplot(111)
    frequency = np.asarray(data["frequency_GHz"])
    B_values = np.asarray(data["B_mT"])
    spectra = displayed_odmr_spectra(data, config)
    denominator = max(len(B_values) - 1, 1)

    for index, (B_mT, spectrum) in enumerate(zip(B_values, spectra)):
        color = ODMR_CMAP(0.10 + 0.80 * index / denominator)
        axis.plot(frequency, spectrum, color=color, label=fr"{B_mT:g} mT")

    axis.set_xlabel("Microwave frequency (GHz)")
    axis.set_ylabel("Normalized fluorescence + offset")
    axis.set_xlim(config.f_min_GHz, config.f_max_GHz)
    ymin = float(np.nanmin(spectra))
    ymax = float(np.nanmax(spectra))
    padding = max(0.018, 0.06 * (ymax - ymin))
    axis.set_ylim(ymin - padding, ymax + padding)
    axis.legend(title=r"$|B|$", frameon=False, ncol=2, handlelength=2.1)
    style_axis(axis)
    figure.subplots_adjust(left=0.16, right=0.97, bottom=0.17, top=0.96)


def _json_ready(value):
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, tuple):
        return list(value)
    raise TypeError(f"Cannot serialize {type(value).__name__}")


def _metadata(config: ODMRConfig) -> dict[str, object]:
    return {
        "schema_version": 1,
        "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "script": SCRIPT_NAME,
        "figure": "odmr",
        "parameters": asdict(config),
        "notes": "CW-ODMR reuses weak-MW model A in nv_odmr_gui.py.",
        "display_equation": (
            "PL_displayed_i = 1 + i*vertical_offset "
            "- magnitude_scale_i*(1 - PL_physical_i)"
        ),
    }


def export_odmr_csv(path: Path, data: dict[str, object], config: ODMRConfig) -> None:
    frequency = np.asarray(data["frequency_GHz"])
    B_values = np.asarray(data["B_mT"])
    physical = np.asarray(data["PL_norm"])
    displayed = displayed_odmr_spectra(data, config)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        headers = ["frequency_GHz"]
        for value in B_values:
            headers.extend(
                [f"PL_physical_B_{value:g}_mT", f"PL_displayed_B_{value:g}_mT"]
            )
        writer.writerow(headers)
        columns: list[np.ndarray] = []
        for physical_curve, displayed_curve in zip(physical, displayed):
            columns.extend([physical_curve, displayed_curve])
        writer.writerows(zip(frequency, *columns))


def export_odmr_xlsx(path: Path, data: dict[str, object], config: ODMRConfig) -> None:
    """Export machine-readable spectra plus the simulation/display parameters."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Excel export requires openpyxl. Install it with: pip install openpyxl"
        ) from exc

    frequency = np.asarray(data["frequency_GHz"], dtype=float)
    B_values = np.asarray(data["B_mT"], dtype=float)
    physical = np.asarray(data["PL_norm"], dtype=float)
    displayed = displayed_odmr_spectra(data, config)

    workbook = Workbook()
    spectra_sheet = workbook.active
    spectra_sheet.title = "Spectra"
    spectra_sheet.sheet_view.showGridLines = False

    headers = ["Microwave frequency (GHz)"]
    for value in B_values:
        headers.extend(
            [
                f"Physical PL (B={value:g} mT)",
                f"Displayed PL (B={value:g} mT)",
            ]
        )
    spectra_sheet.append(headers)
    columns: list[np.ndarray] = []
    for physical_curve, displayed_curve in zip(physical, displayed):
        columns.extend([physical_curve, displayed_curve])
    for row in zip(frequency, *columns):
        spectra_sheet.append([float(value) for value in row])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for cell in spectra_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    spectra_sheet.freeze_panes = "A2"
    spectra_sheet.auto_filter.ref = spectra_sheet.dimensions
    spectra_sheet.column_dimensions["A"].width = 27
    for column_index in range(2, len(headers) + 1):
        spectra_sheet.column_dimensions[get_column_letter(column_index)].width = 25
    for column in spectra_sheet.iter_cols(
        min_row=2, max_row=spectra_sheet.max_row, max_col=len(headers)
    ):
        for cell in column:
            cell.font = Font(name="Arial", size=10)
            cell.number_format = "0.000000"

    parameters_sheet = workbook.create_sheet("Parameters")
    parameters_sheet.sheet_view.showGridLines = False
    parameter_rows = [
        ("Parameter", "Value", "Unit / definition"),
        ("Magnetic fields", ", ".join(f"{v:g}" for v in config.B_values_mT), "mT"),
        ("Magnitude scales", ", ".join(f"{v:g}" for v in config.magnitude_scales), "one per sorted B curve"),
        ("Frequency minimum", config.f_min_GHz, "GHz"),
        ("Frequency maximum", config.f_max_GHz, "GHz"),
        ("Frequency points", config.n_frequency, "count"),
        ("Zero-field splitting D", config.D_GHz, "GHz"),
        ("Transverse strain E", config.strain_E_MHz, "MHz"),
        ("Lorentzian FWHM", config.linewidth_MHz, "MHz"),
        ("Single-line contrast", config.contrast_percent, "%"),
        ("Angle B to NV axis", config.B_angle_deg, "degrees"),
        ("Vertical offset per curve", config.vertical_offset, "normalized fluorescence"),
        (
            "Displayed-curve equation",
            "PL_displayed_i = 1 + i*offset - scale_i*(1 - PL_physical_i)",
            "i starts at 0 in ascending B order",
        ),
        ("Generated UTC", datetime.now(timezone.utc).isoformat(timespec="seconds"), "ISO 8601"),
    ]
    for row in parameter_rows:
        parameters_sheet.append(row)
    for cell in parameters_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in parameters_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
    parameters_sheet.freeze_panes = "A2"
    parameters_sheet.auto_filter.ref = parameters_sheet.dimensions
    parameters_sheet.column_dimensions["A"].width = 29
    parameters_sheet.column_dimensions["B"].width = 72
    parameters_sheet.column_dimensions["C"].width = 33

    workbook.save(path)


class PlotWindow:
    """The independent ODMR plot window with toolbar and bundle export."""

    def __init__(
        self,
        app: "ReviewBoxApp",
        title: str,
        filename: str,
    ) -> None:
        self.app = app
        self.title = title
        self.filename = filename
        self.window: tk.Toplevel | None = None
        self.figure: Figure | None = None
        self.canvas: FigureCanvasTkAgg | None = None

    def open(self) -> None:
        if self.window is not None and self.window.winfo_exists():
            self.window.deiconify()
            self.window.lift()
            return

        self.window = tk.Toplevel(self.app.root)
        self.window.title(self.title)
        self.window.protocol("WM_DELETE_WINDOW", self._close)
        self.figure = Figure(figsize=(5.2, 3.8), dpi=110, facecolor="white")
        self.canvas = FigureCanvasTkAgg(self.figure, master=self.window)
        self.canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)

        footer = ttk.Frame(self.window, padding=(6, 2, 6, 6))
        footer.pack(fill=tk.X)
        toolbar = NavigationToolbar2Tk(self.canvas, footer, pack_toolbar=False)
        toolbar.update()
        toolbar.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Button(footer, text="Export figures + Excel data", command=self.export_bundle).pack(
            side=tk.RIGHT, padx=(8, 0)
        )
        self._place()

    def _place(self) -> None:
        assert self.window is not None
        screen_w = self.app.root.winfo_screenwidth()
        screen_h = self.app.root.winfo_screenheight()
        plot_w = max(460, min(720, (screen_w - 440) // 2))
        plot_h = max(430, min(610, screen_h - 120))
        x = 420
        if x + plot_w > screen_w:
            x = max(10, screen_w - plot_w - 10)
        y = 35
        self.window.geometry(f"{plot_w}x{plot_h}+{x}+{y}")

    def _close(self) -> None:
        if self.window is not None:
            self.window.destroy()
        self.window = None
        self.figure = None
        self.canvas = None

    def redraw(self) -> None:
        self.open()
        assert self.figure is not None and self.canvas is not None
        draw_odmr(self.figure, self.app.odmr_data, self.app.odmr_config)
        self.canvas.draw_idle()

    def export_bundle(self) -> None:
        if self.figure is None:
            return
        selected = filedialog.asksaveasfilename(
            parent=self.window,
            title=f"Export {self.title}",
            defaultextension=".png",
            initialfile=self.filename,
            filetypes=[("PNG image", "*.png"), ("All files", "*.*")],
        )
        if not selected:
            return
        stem = Path(selected).with_suffix("")
        paths = {
            "png": stem.with_suffix(".png"),
            "pdf": stem.with_suffix(".pdf"),
            "svg": stem.with_suffix(".svg"),
            "csv": stem.with_suffix(".csv"),
            "xlsx": stem.with_suffix(".xlsx"),
            "json": stem.with_suffix(".json"),
        }
        data = self.app.odmr_data
        config = self.app.odmr_config
        try:
            self.figure.savefig(paths["png"], dpi=300, bbox_inches="tight", facecolor="white")
            self.figure.savefig(paths["pdf"], bbox_inches="tight", facecolor="white")
            self.figure.savefig(paths["svg"], bbox_inches="tight", facecolor="white")
            export_odmr_csv(paths["csv"], data, config)
            export_odmr_xlsx(paths["xlsx"], data, config)
            payload = _metadata(config)
            payload["data"] = dict(data)
            payload["data"]["PL_displayed"] = displayed_odmr_spectra(data, config)
            with paths["json"].open("w", encoding="utf-8") as handle:
                json.dump(payload, handle, ensure_ascii=False, indent=2, default=_json_ready)
        except Exception:
            messagebox.showerror("Export failed", traceback.format_exc()[-1800:], parent=self.window)
            return
        messagebox.showinfo(
            "Export complete",
            "Created PNG, PDF, SVG, CSV, XLSX, and JSON files at:\n"
            + str(stem.parent.resolve()),
            parent=self.window,
        )


class ReviewBoxApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("NV-center ODMR splitting figure generator")
        self.root.geometry("405x650+10+20")
        self.root.minsize(390, 600)

        self.odmr_config = ODMRConfig()
        self.odmr_data: dict[str, object] = {}
        self.vars: dict[str, tk.Variable] = {}
        self.status = tk.StringVar(value="Ready")

        self.odmr_window = PlotWindow(
            self,
            "NV ODMR: fluorescence vs microwave frequency",
            "nv_odmr_B_overlay.png",
        )

        self._build_variables()
        self._build_layout()
        self.root.after(80, self.update_figure)

    def _build_variables(self) -> None:
        odmr = self.odmr_config
        defaults = {
            "B_values": ", ".join(f"{value:g}" for value in odmr.B_values_mT),
            "f_min_GHz": odmr.f_min_GHz,
            "f_max_GHz": odmr.f_max_GHz,
            "n_frequency": odmr.n_frequency,
            "D_GHz": odmr.D_GHz,
            "strain_E_MHz": odmr.strain_E_MHz,
            "linewidth_MHz": odmr.linewidth_MHz,
            "contrast_percent": odmr.contrast_percent,
            "B_angle_deg": odmr.B_angle_deg,
            "vertical_offset": odmr.vertical_offset,
            "magnitude_scales": ", ".join(f"{value:g}" for value in odmr.magnitude_scales),
        }
        for name, value in defaults.items():
            self.vars[name] = tk.StringVar(value=str(value))

    def _entry(self, parent, row: int, label: str, name: str, unit: str = "") -> None:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=2)
        entry = ttk.Entry(parent, textvariable=self.vars[name], width=15)
        entry.grid(row=row, column=1, sticky="ew", padx=(8, 5), pady=2)
        entry.bind("<Return>", lambda _event: self.update_figure())
        ttk.Label(parent, text=unit).grid(row=row, column=2, sticky="w", pady=2)

    def _build_layout(self) -> None:
        container = ttk.Frame(self.root, padding=10)
        container.pack(fill=tk.BOTH, expand=True)

        intro = ttk.Label(
            container,
            text=(
                "Publication-oriented ODMR splitting panel using the existing\n"
                "Hamiltonian model. Export includes an Excel data workbook."
            ),
            justify=tk.LEFT,
        )
        intro.pack(fill=tk.X, pady=(0, 8))

        odmr = ttk.LabelFrame(container, text="ODMR field overlay", padding=8)
        odmr.pack(fill=tk.X)
        self._entry(odmr, 0, "B values", "B_values", "mT")
        self._entry(odmr, 1, "Frequency min", "f_min_GHz", "GHz")
        self._entry(odmr, 2, "Frequency max", "f_max_GHz", "GHz")
        self._entry(odmr, 3, "Frequency points", "n_frequency")
        self._entry(odmr, 4, "Zero-field splitting D", "D_GHz", "GHz")
        self._entry(odmr, 5, "Transverse strain E", "strain_E_MHz", "MHz")
        self._entry(odmr, 6, "Lorentzian FWHM", "linewidth_MHz", "MHz")
        self._entry(odmr, 7, "Single-line contrast", "contrast_percent", "%")
        self._entry(odmr, 8, "Angle B to NV axis", "B_angle_deg", "deg")
        self._entry(odmr, 9, "Vertical offset / curve", "vertical_offset", "norm.")
        self._entry(odmr, 10, "Dip magnitude multipliers", "magnitude_scales")
        odmr.columnconfigure(1, weight=1)
        ttk.Label(
            odmr,
            text=(
                "Magnitude order follows the sorted B list; one value applies to all. "
                "Fixed: single NV axis, no hyperfine, weak-MW model A, transverse MW."
            ),
            foreground="#555555",
            wraplength=340,
        ).grid(row=11, column=0, columnspan=3, sticky="w", pady=(5, 0))

        buttons = ttk.Frame(container)
        buttons.pack(fill=tk.X, pady=(10, 5))
        ttk.Button(buttons, text="Update figure", command=self.update_figure).pack(
            side=tk.LEFT, fill=tk.X, expand=True
        )
        ttk.Button(buttons, text="Reopen plot window", command=self.reopen_window).pack(
            side=tk.LEFT, fill=tk.X, expand=True, padx=(6, 0)
        )
        ttk.Label(container, textvariable=self.status, foreground="#1b5e20", wraplength=360).pack(
            fill=tk.X
        )

    def _read_odmr_config(self) -> ODMRConfig:
        B_values = parse_B_values(self.vars["B_values"].get())
        return ODMRConfig(
            B_values_mT=B_values,
            f_min_GHz=float(self.vars["f_min_GHz"].get()),
            f_max_GHz=float(self.vars["f_max_GHz"].get()),
            n_frequency=int(self.vars["n_frequency"].get()),
            D_GHz=float(self.vars["D_GHz"].get()),
            strain_E_MHz=float(self.vars["strain_E_MHz"].get()),
            linewidth_MHz=float(self.vars["linewidth_MHz"].get()),
            contrast_percent=float(self.vars["contrast_percent"].get()),
            B_angle_deg=float(self.vars["B_angle_deg"].get()),
            vertical_offset=float(self.vars["vertical_offset"].get()),
            magnitude_scales=parse_magnitude_scales(
                self.vars["magnitude_scales"].get(), len(B_values)
            ),
        )

    def update_figure(self) -> None:
        self.root.configure(cursor="watch")
        self.status.set("Calculating...")
        self.root.update_idletasks()
        try:
            odmr_config = self._read_odmr_config()
            odmr_data = simulate_odmr_overlay(odmr_config)
        except Exception as exc:
            self.status.set("Parameters need attention.")
            messagebox.showerror("Could not update figure", str(exc), parent=self.root)
            return
        finally:
            self.root.configure(cursor="")

        self.odmr_config = odmr_config
        self.odmr_data = odmr_data
        self.odmr_window.redraw()
        self.status.set("Figure is current. Export from the plot window.")

    def reopen_window(self) -> None:
        if not self.odmr_data:
            self.update_figure()
            return
        self.odmr_window.redraw()


def main() -> None:
    root = tk.Tk()
    try:
        ttk.Style(root).theme_use("vista")
    except tk.TclError:
        pass
    ReviewBoxApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
