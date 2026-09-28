#!/usr/bin/env python3
"""Plot modeled iSCAT and Mie signals with a separate camera-noise benchmark.

Keep this file beside ``iscat_darkfield_mie_comparison_gui.py``, which supplies
the sphere-scattering model. The published camera values are contextual
benchmarks; they do not establish a simulated limit of detection.

Run ``python iscat_simplified_gui.py`` for the GUI or add ``--batch`` to export
the default figure and data without opening a window.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.colors import to_rgb
from openpyxl import Workbook
from openpyxl.styles import Font
import iscat_darkfield_mie_comparison_gui as core

OUTPUT_STEM = "iscat_simplified"
SIGNAL_MODES = ("Phase-optimized envelope", "Selected phase")
CAMERA = {
    "camera": "Hamamatsu ORCA-Flash4.0 LT",
    "source": "Wu, Tsai & Hsieh, J. Phys. Chem. C (2025)",
    "doi": "10.1021/acs.jpcc.4c07989",
    "location": "Main-text noise discussion; SI Fig. S4; Experimental Methods",
    "measurement_wavelength_nm": 532,
    "filter_bandwidth_nm": 18,
    "full_well_electrons": 30000,
    "operating_fraction_of_full_well": 0.5,
    "typical_exposure_ms": 10,
    "single_frame_shot_noise_percent": 0.8,
    "averaged_frames": 1000,
    "reported_averaged_noise_percent": 0.03,
    "meaning": "Literature noise benchmark, not a universal detection threshold",
}


def default_config():
    return dict(particles=list(core.DEFAULT_PARTICLES), wavelengths=[533., 632.],
                diameter_min=10., diameter_max=1000., points=1600,
                reference_ratio=0.1, phase=90., n_medium=1.333, n_glass=1.518,
                na=1.42, collection=0.3, model=core.DEFAULT_MODEL,
                signal_mode=SIGNAL_MODES[0], log_x=True, log_y=True)


def calculate(config):
    system = core.ComparisonSystem(
        n_medium=config["n_medium"], n_glass=config["n_glass"],
        numerical_aperture=config["na"], collection_efficiency=config["collection"],
        reference_intensity_ratio=config["reference_ratio"], net_phase_deg=config["phase"])
    args = (system, config["particles"], config["wavelengths"],
            config["diameter_min"], config["diameter_max"], config["points"])
    results = core.calculate_particle_wavelength_overlay(*args, model=config["model"])
    # Keep the Mie comparison at full-model fidelity when the iSCAT model changes.
    mie_results = (results if config["model"] == core.DEFAULT_MODEL else
                   core.calculate_particle_wavelength_overlay(*args, model=core.DEFAULT_MODEL))
    signal_key = ("iscat_envelope" if config["signal_mode"] == SIGNAL_MODES[0]
                  else "iscat_single_magnitude")
    return [dict(material=r["particle_name"], wavelength=r["system"].wavelength_vacuum_nm,
                 diameter=r["diameter_nm"], signal=100*r[signal_key],
                 mie=100*m["darkfield"])
            for r, m in zip(results, mie_results)]


def draw(config, results, figure=None):
    fig = figure if figure is not None else plt.figure(figsize=(11.5, 6.3))
    fig.clear()
    grid = fig.add_gridspec(1, 2, width_ratios=(3.5, 1.35), wspace=0.18,
                           left=0.09, right=0.98, bottom=0.18, top=0.91)
    ax = fig.add_subplot(grid[0])
    side = fig.add_subplot(grid[1])
    side.axis("off")
    handles = []
    for row in results:
        pi = config["particles"].index(row["material"])
        wi = config["wavelengths"].index(row["wavelength"])
        base = np.asarray(to_rgb(core.SERIES_COLORS[pi % len(core.SERIES_COLORS)]))
        tint = min(0.65, 0.30*wi)
        color = (1-tint)*base + tint
        # Undefined log(0) is omitted, never replaced by a made-up positive floor.
        signal = np.where(row["signal"] > 0, row["signal"], np.nan) if config["log_y"] else row["signal"]
        mie = np.where(row["mie"] > 0, row["mie"], np.nan) if config["log_y"] else row["mie"]
        ax.plot(row["diameter"], signal, color=color, lw=2.0)
        ax.plot(row["diameter"], mie, color=color, lw=1.35, ls="--")
        material = row["material"].replace(" bead", "").replace(" nanoparticle", "")
        material = material.replace(" dielectric (illustrative)", " (model)")
        handles.append(Line2D([], [], color=color, lw=2,
                              label=f"{material}, {row['wavelength']:g} nm"))
    ax.set_xscale("log" if config["log_x"] else "linear")
    ax.set_yscale("log" if config["log_y"] else "linear")
    ax.set_xlim(config["diameter_min"], config["diameter_max"])
    if not config["log_y"]:
        ax.set_ylim(bottom=0)
    ax.set_xlabel("Particle diameter (nm)")
    ax.set_ylabel("Reference-normalized signal (%)")
    ax.set_title("iSCAT and dark-field Mie", fontsize=13, loc="left", pad=12)
    ax.grid(True, which="major", color="0.88", lw=0.6)
    ax.legend(handles=handles, loc="upper left", frameon=False, fontsize=8)
    signal_label = ("iSCAT: phase-optimized" if config["signal_mode"] == SIGNAL_MODES[0]
                    else f"iSCAT: phase {config['phase']:g} deg")
    styles = [Line2D([], [], color="0.25", lw=2, label=signal_label),
              Line2D([], [], color="0.25", lw=1.35, ls="--", label="Dark-field: full Mie")]
    fig.legend(handles=styles, loc="lower left", bbox_to_anchor=(0.09, 0.065),
               frameon=False, fontsize=9, ncol=2)
    formula = (r"iSCAT: $2q+q^2$" if config["signal_mode"] == SIGNAL_MODES[0]
               else r"iSCAT: $|2q\cos\phi+q^2|$")
    fig.text(0.09, 0.033, formula + r"; dark-field: $q^2$; $q^2=P_{sca}/P_{ref}$.",
             fontsize=8, color="0.35")
    side.text(0, 1, "Camera data", fontsize=12, weight="bold", va="top")
    side.text(0, 0.90, "Hamamatsu\nORCA-Flash4.0 LT", fontsize=10, va="top", linespacing=1.5)
    camera_rows = [
        ("Full well", "30,000 e- / pixel"),
        ("Background level", "~50% full well"),
        ("Single-frame shot noise", "~0.8%"),
        ("Noise after 1000 frames", "~0.03%"),
    ]
    for index, (label, value) in enumerate(camera_rows):
        y = 0.74 - index*0.13
        side.text(0, y, label, fontsize=8.5, color="0.4", va="top")
        side.text(0, y-0.045, value, fontsize=11, va="top")
    side.text(0, 0.17, "Literature conditions: 532 nm\nTypical exposure: 10 ms/frame\n"
              "Wu et al. (2025), SI Fig. S4\nDOI: 10.1021/acs.jpcc.4c07989",
              fontsize=7.8, va="top", linespacing=1.5)
    side.text(0, -0.05, "Independent noise reference;\nnot a universal detection limit.",
              fontsize=8, color="0.35", va="top")
    return fig


def save_excel(path, config, results):
    """Export five curve columns and separate camera and model-setting sheets."""
    workbook = Workbook()
    curves = workbook.active
    curves.title = "Curves"
    curves.append(["material", "wavelength_nm", "diameter_nm",
                   "iscat_normalized_signal_percent", "darkfield_full_mie_percent"])
    for result in results:
        for diameter, signal, mie in zip(result["diameter"], result["signal"], result["mie"]):
            curves.append([result["material"], float(result["wavelength"]),
                           float(diameter), float(signal) if np.isfinite(signal) else None,
                           float(mie) if np.isfinite(mie) else None])
    curves.freeze_panes = "C2"
    curves.auto_filter.ref = curves.dimensions
    for column, width in {"A": 30, "B": 18, "C": 18, "D": 34, "E": 31}.items():
        curves.column_dimensions[column].width = width
    for cell in curves[1]:
        cell.font = Font(bold=True)

    camera = workbook.create_sheet("Camera benchmark")
    camera.append(["Parameter", "Literature value"])
    for key, value in CAMERA.items():
        camera.append([key, value])
    camera.column_dimensions["A"].width = 39
    camera.column_dimensions["B"].width = 72
    for cell in camera[1]:
        cell.font = Font(bold=True)

    settings = workbook.create_sheet("Model settings")
    settings.append(["Parameter", "Value"])
    for key, value in config.items():
        settings.append([key, json.dumps(value, ensure_ascii=True)
                         if isinstance(value, (list, tuple, dict)) else value])
    settings.column_dimensions["A"].width = 35
    settings.column_dimensions["B"].width = 72
    for cell in settings[1]:
        cell.font = Font(bold=True)
    workbook.save(path)


def save(stem, config, results, fig):
    stem = Path(stem)
    save_excel(stem.with_suffix(".xlsx"), config, results)
    fig.savefig(stem.with_suffix(".png"), dpi=250)
    fig.savefig(stem.with_suffix(".pdf"))
    # Only requested series data; camera and calculation settings are separate.
    with stem.with_suffix(".csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["material", "wavelength_nm", "diameter_nm",
                         "iscat_normalized_signal_percent", "darkfield_full_mie_percent"])
        for r in results:
            writer.writerows((r["material"], r["wavelength"], *values)
                             for values in zip(r["diameter"], r["signal"], r["mie"]))
    stem.with_suffix(".json").write_text(json.dumps(config, indent=2), encoding="utf-8")
    stem.with_name(stem.name + "_camera.json").write_text(
        json.dumps(CAMERA, indent=2), encoding="utf-8")


class App:
    def __init__(self, root):
        import tkinter as tk
        from tkinter import ttk
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
        self.root = root
        self.current = None
        root.title("iSCAT - simplified comparison")
        root.geometry("1480x850")
        root.minsize(1150, 700)
        shell = ttk.Frame(root, padding=8)
        shell.pack(fill="both", expand=True)
        left = ttk.Frame(shell, width=340)
        left.pack(side="left", fill="y", padx=(0, 10))
        tabs = ttk.Notebook(left)
        tabs.pack(fill="both", expand=True)
        basic, advanced = ttk.Frame(tabs, padding=8), ttk.Frame(tabs, padding=8)
        tabs.add(basic, text="Main")
        tabs.add(advanced, text="Optics")
        self.vars = {}
        defaults = default_config()
        for key, value in defaults.items():
            if key not in ("particles", "wavelengths"):
                self.vars[key] = (tk.BooleanVar(value=value) if isinstance(value, bool)
                                  else tk.StringVar(value=str(value)))
        self.vars["wavelengths"] = tk.StringVar(value="533, 632")
        ttk.Label(basic, text="Materials").pack(anchor="w")
        self.materials = {name: tk.BooleanVar(value=name in defaults["particles"])
                          for name in core.PARTICLE_PRESETS}
        for name, var in self.materials.items():
            ttk.Checkbutton(basic, text=name, variable=var).pack(anchor="w", pady=2)

        def field(parent, label, key, options=None):
            ttk.Label(parent, text=label).pack(anchor="w", pady=(9, 2))
            if options:
                entry = ttk.Combobox(parent, textvariable=self.vars[key],
                                     values=options, state="readonly", width=31)
            else:
                entry = ttk.Entry(parent, textvariable=self.vars[key], width=33)
            entry.pack(fill="x")

        field(basic, "Wavelengths (nm, comma-separated)", "wavelengths")
        field(basic, "Minimum diameter (nm)", "diameter_min")
        field(basic, "Maximum diameter (nm)", "diameter_max")
        field(basic, "Displayed iSCAT signal", "signal_mode", SIGNAL_MODES)
        ttk.Checkbutton(basic, text="Log diameter axis", variable=self.vars["log_x"]).pack(anchor="w", pady=(10, 0))
        ttk.Checkbutton(basic, text="Log signal axis", variable=self.vars["log_y"]).pack(anchor="w")
        ttk.Label(basic, text="Click Calculate & Plot after changing selections.",
                  wraplength=300, foreground="#555555").pack(anchor="w", pady=10)
        for label, key in [("Reference intensity ratio", "reference_ratio"),
                           ("Selected phase (deg)", "phase"),
                           ("Medium refractive index", "n_medium"),
                           ("Glass refractive index", "n_glass"),
                           ("Objective NA", "na"), ("Collection efficiency", "collection"),
                           ("Diameter points", "points")]:
            field(advanced, label, key)
        field(advanced, "iSCAT scattering model", "model", core.MODEL_OPTIONS)
        ttk.Label(advanced, text="Envelope uses the maximum over phase. Selected phase uses your phase value; at 90 deg it equals dark-field for the same model.",
                  wraplength=300, foreground="#555555").pack(pady=10)
        ttk.Button(left, text="Calculate & Plot", command=self.update).pack(fill="x", pady=(8, 3))
        ttk.Button(left, text="Save figure + data", command=self.export).pack(fill="x", pady=3)
        ttk.Button(left, text="Save Excel only", command=self.export_excel).pack(fill="x", pady=3)
        self.status = tk.StringVar()
        ttk.Label(left, textvariable=self.status, wraplength=320).pack(fill="x", pady=5)
        right = ttk.Frame(shell)
        right.pack(side="left", fill="both", expand=True)
        self.fig = Figure(figsize=(11.5, 6.3), dpi=100)
        self.canvas = FigureCanvasTkAgg(self.fig, master=right)
        self.canvas.get_tk_widget().pack(fill="both", expand=True)
        NavigationToolbar2Tk(self.canvas, right).update()
        self.update()

    def update(self):
        from tkinter import messagebox
        try:
            config = {key: var.get() for key, var in self.vars.items()}
            config["particles"] = [k for k, v in self.materials.items() if v.get()]
            if not config["particles"]:
                raise ValueError("Select at least one material.")
            config["wavelengths"] = list(core.parse_wavelengths(config["wavelengths"]))
            for key in ("diameter_min", "diameter_max", "reference_ratio", "phase", "n_medium", "n_glass", "na", "collection"):
                config[key] = float(config[key])
            config["points"] = int(config["points"])
            results = calculate(config)
            draw(config, results, self.fig)
            self.canvas.draw_idle()
            self.current = config, results
            self.status.set(f"Updated: {len(results)} material / wavelength pairs.")
        except Exception as exc:
            self.current = None
            self.status.set("Check input values.")
            messagebox.showerror("Input error", str(exc), parent=self.root)

    def export(self):
        from tkinter import filedialog, messagebox
        self.update()
        if self.current is None:
            return
        path = filedialog.asksaveasfilename(parent=self.root, defaultextension=".png",
                initialdir=Path(__file__).resolve().parent, initialfile=OUTPUT_STEM+".png",
                filetypes=[("PNG", "*.png")])
        if path:
            try:
                save(Path(path).with_suffix(""), *self.current, self.fig)
                self.status.set("Saved Excel curves, camera benchmark, and settings; PNG/PDF and CSV/JSON.")
            except Exception as exc:
                messagebox.showerror("Save error", str(exc), parent=self.root)

    def export_excel(self):
        from tkinter import filedialog, messagebox
        self.update()
        if self.current is None:
            return
        path = filedialog.asksaveasfilename(parent=self.root, defaultextension=".xlsx",
            initialdir=Path(__file__).resolve().parent, initialfile=OUTPUT_STEM+".xlsx",
            filetypes=[("Excel workbook", "*.xlsx")])
        if path:
            try:
                save_excel(path, *self.current)
                self.status.set("Saved Excel: five curve columns, with separate camera and settings sheets.")
            except Exception as exc:
                messagebox.showerror("Excel export error", str(exc), parent=self.root)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--batch", action="store_true")
    args = parser.parse_args()
    if args.batch:
        config = default_config()
        results = calculate(config)
        figure = draw(config, results)
        save(Path(__file__).resolve().parent / OUTPUT_STEM, config, results, figure)
        plt.close(figure)
        print("Saved simplified figure, curve data and separate camera data.")
    else:
        core.configure_tk_runtime()
        import tkinter as tk
        root = tk.Tk()
        App(root)
        root.mainloop()
