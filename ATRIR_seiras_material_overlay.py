#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Standalone Au/Ag/Cu ATR-SEIRAS material-overlay simulator.

The planar optical stack is

    prism | metal or metal/void effective layer | [surface adsorbate] | solution

and all angles are internal angles measured from the surface normal.  The top
panel reports the raw p-polarized reflectivity of a vibration-free reference.
When the adsorbate is enabled, the lower panel reports the experimentally
relevant differential absorbance

    Delta A(wn, theta) = -log10(R_with_vibration / R_reference).

The default adsorbate is a generic CO comparison oscillator represented by a
thin Lorentz layer centred at 2050 cm^-1.  Its parameters are deliberately held
identical for Au, Ag, and Cu so the curves compare electromagnetic effects, not
metal-dependent adsorption chemistry.  The reference retains the same background
layer but removes the oscillator, so Delta A isolates the CO band.

The GUI compares Au, Ag, and Cu on one set of axes.  All three materials use the
same user-entered nominal thicknesses (5 and 10 nm by default), the same
prism, solution, morphology, adsorbate, and angular beam.  Colour identifies
material and line style identifies thickness.  Keeping the thickness grid
shared prevents a material comparison from silently changing film thickness.

The Film coupling tab reports the fraction of incident p-polarized power that
is dissipated inside the finite metal/EMA film.  It is calculated from the
decrease of normal Poynting flux across that layer, rather than from 1-R, so
transmission or downstream absorption is not incorrectly assigned to the film.
The 60-degree power budget is shown explicitly for comparison with common
fixed-angle ATR-SEIRAS experiments.

The default morphology is a connected porous film with gap-localized adsorbate.
The bounded gap-field correction is semi-quantitative and is not a replacement
for an explicit FEM reconstruction of a measured film morphology.
The prism can be selected from Si/Ge/ZnSe or replaced by a constant custom
index.  Films can be treated either as continuous bulk metal or as a
Bruggeman effective medium containing solution-filled voids.  The EMA is a
far-field morphology baseline; it does not resolve local island hot spots.

The Ordal Drude constants are restricted to the infrared.  A constant custom
metal n+ik option is provided for reproducing a single-wavelength visible SPR
example with literature optical constants.

Usage:
  python ATRIR_seiras_material_overlay.py
  python ATRIR_seiras_material_overlay.py --selftest

Normal startup never writes output files.  Figures and Excel data are exported
only when the corresponding GUI button is clicked.
"""

from __future__ import annotations

import argparse
import math
import sys
from dataclasses import dataclass

import numpy as np


HEADLESS = "--selftest" in sys.argv
if HEADLESS:
    import matplotlib

    matplotlib.use("Agg")

from matplotlib.figure import Figure


EPS = 1.0e-14j


# Ordal et al., Applied Optics 24, 4493-4499 (1985), in cm^-1.
# These compact Drude fits are useful baselines.  Measured optical constants
# should replace them for quantitative work on a particular deposited film.
METAL_DRUDE = {
    "Au": (72800.0, 215.0),
    "Ag": (72700.0, 145.0),
    "Cu": (59600.0, 73.2),
    "Pt": (41500.0, 558.0),
    "Pd": (44000.0, 124.0),
}

CURVE_COLORS = (
    "#0072B2", "#D55E00", "#009E73", "#CC79A7",
    "#56B4E9", "#E69F00", "#000000", "#999999",
)

MATERIAL_COLORS = {
    "Au": "#D99A00",
    "Ag": "#4C78A8",
    "Cu": "#B5542D",
}

OVERLAY_LINESTYLES = ("-", "--", "-.", ":")

MORPHOLOGY_CONTINUOUS = "Smooth continuous film (planar TMM)"
MORPHOLOGY_ISOLATED = "Isolated islands (surface CO)"
MORPHOLOGY_CONNECTED = "Connected porous film (gap CO)"
MORPHOLOGY_THICKNESS_LINKED = "Thickness-linked semi-quantitative"
MORPHOLOGY_CHOICES = (
    MORPHOLOGY_THICKNESS_LINKED,
    MORPHOLOGY_ISOLATED,
    MORPHOLOGY_CONNECTED,
    MORPHOLOGY_CONTINUOUS,
)

# The auto model changes smoothly from isolated to connected over a fixed
# 6 nm-wide interval centred on the user-selected percolation thickness.  This
# compact assumption avoids pretending that nominal thickness alone uniquely
# determines a real deposited film's topology.
PERCOLATION_HALF_WIDTH_NM = 3.0


@dataclass(frozen=True)
class FilmCase:
    metal: str
    thickness_nm: float
    color: str
    linestyle: str = "-"

    @property
    def label(self):
        return "%s %g nm" % (self.metal, self.thickness_nm)


DEFAULTS = {
    "wavenumber_cm": 2050.0,
    "prism": "Si",
    "custom_prism_n": 1.515,
    "n_solution": 1.300,
    "k_solution": 0.0050,
    "angle_min_deg": 15.0,
    "angle_max_deg": 85.0,
    "divergence_fwhm_deg": 2.0,
    "multi_metal_comparison": True,
    "metal": "Au",
    "thicknesses_nm": (5.0, 10.0),
    "au_thicknesses_nm": (5.0, 15.0),
    "ag_thicknesses_nm": (5.0, 15.0),
    "cu_thicknesses_nm": (5.0, 15.0),
    "metal_optical_model": "Infrared Drude",
    "custom_metal_n": 0.18,
    "custom_metal_k": 3.43,
    "film_morphology": MORPHOLOGY_CONNECTED,
    "ema_fill_fraction": 0.70,
    "connected_fill_fraction": 0.90,
    "percolation_thickness_nm": 10.0,
    "gap_field_gain": 20.0,
    "include_adsorbate": True,
    "adsorbate_name": "CO oscillator (generic comparison)",
    "adsorbate_center_cm": 2050.0,
    "adsorbate_fwhm_cm": 20.0,
    "adsorbate_strength": 0.0020,
    "adsorbate_thickness_nm": 1.0,
    "adsorbate_n_background": 1.45,
}


def _csqrt(z):
    """Passive square-root branch: Im(q)>=0; if real, Re(q)>=0."""
    q = np.sqrt(np.asarray(z, dtype=complex) + 0j)
    flip = (q.imag < -1e-15) | ((np.abs(q.imag) <= 1e-15) & (q.real < 0))
    return np.where(flip, -q, q)


def drude_eps(wavenumber_cm, metal):
    """Infrared Drude dielectric function with Im(eps)>0 convention."""
    wn = np.asarray(wavenumber_cm, dtype=float)
    wp, gamma = METAL_DRUDE[metal]
    return 1.0 - wp**2 / (wn**2 + 1j * wn * gamma)


def metal_eps(wavenumber_cm, metal, optical_model="Infrared Drude",
              custom_n=0.18, custom_k=3.43):
    """Metal dielectric function at the selected observation wavenumber."""
    if optical_model == "Infrared Drude":
        return drude_eps(wavenumber_cm, metal)
    if optical_model == "Custom n,k":
        return complex(float(custom_n), float(custom_k)) ** 2
    raise ValueError("Unknown metal optical model: %s" % optical_model)


def bruggeman_eps(eps_metal, eps_host, metal_fill_fraction):
    """Symmetric 3D Bruggeman effective dielectric function.

    The physical analytical root is tracked continuously from the host at
    f=0 to the requested metal fraction.  This avoids a branch swap near the
    Bruggeman percolation region.
    """
    eps_metal = complex(eps_metal)
    eps_host = complex(eps_host)
    fraction = float(metal_fill_fraction)
    if not 0.0 <= fraction <= 1.0:
        raise ValueError("EMA metal fill fraction must be between 0 and 1.")
    if fraction <= 1e-15:
        return eps_host
    if fraction >= 1.0 - 1e-15:
        return eps_metal

    previous = eps_host
    for fi in np.linspace(0.0, fraction, 201)[1:]:
        coefficient = ((3.0 * fi - 1.0) * eps_metal
                       + (2.0 - 3.0 * fi) * eps_host)
        discriminant = complex(_csqrt(
            coefficient**2 + 8.0 * eps_metal * eps_host))
        candidates = ((coefficient + discriminant) / 4.0,
                      (coefficient - discriminant) / 4.0)
        passive = [candidate for candidate in candidates
                   if candidate.imag >= -1e-12]
        pool = passive if passive else list(candidates)
        previous = min(pool, key=lambda candidate: abs(candidate - previous))
    return complex(previous)


def canonical_morphology_name(name):
    """Map legacy saved settings to the current morphology names."""
    aliases = {
        "Continuous film": MORPHOLOGY_CONTINUOUS,
        "Island film (Bruggeman EMA)": MORPHOLOGY_ISOLATED,
    }
    return aliases.get(str(name), str(name))


def smooth_connection_fraction(thickness_nm, percolation_thickness_nm):
    """Semi-quantitative film connectedness in the closed interval [0, 1].

    The transition is a cubic smoothstep from ``percolation-3 nm`` to
    ``percolation+3 nm``.  It is a morphology control variable, not a universal
    thickness-to-percolation law.
    """
    lower = float(percolation_thickness_nm) - PERCOLATION_HALF_WIDTH_NM
    span = 2.0 * PERCOLATION_HALF_WIDTH_NM
    u = float(np.clip((float(thickness_nm) - lower) / span, 0.0, 1.0))
    return u * u * (3.0 - 2.0 * u)


def morphology_state(morphology, thickness_nm, isolated_fill_fraction,
                     connected_fill_fraction, percolation_thickness_nm):
    """Return the effective fill and gap participation for one film curve."""
    morphology = canonical_morphology_name(morphology)
    if morphology == MORPHOLOGY_CONTINUOUS:
        connection = 1.0
        fill = 1.0
        gap_participation = 0.0
    elif morphology == MORPHOLOGY_ISOLATED:
        connection = 0.0
        fill = float(isolated_fill_fraction)
        gap_participation = 0.0
    elif morphology == MORPHOLOGY_CONNECTED:
        connection = 1.0
        fill = float(connected_fill_fraction)
        gap_participation = 1.0
    elif morphology == MORPHOLOGY_THICKNESS_LINKED:
        connection = smooth_connection_fraction(
            thickness_nm, percolation_thickness_nm)
        fill = (float(isolated_fill_fraction)
                + connection * (float(connected_fill_fraction)
                                - float(isolated_fill_fraction)))
        gap_participation = connection
    else:
        raise ValueError("Unknown film morphology: %s" % morphology)
    return {
        "morphology": morphology,
        "connection_fraction": float(connection),
        "metal_fill_fraction": float(fill),
        "gap_participation": float(gap_participation),
    }


def pore_adsorbate_occupancy(state, thickness_nm, adsorbate_thickness_nm):
    """Convert fixed adsorbate amount into a bounded EMA-void volume fraction."""
    gap_fraction = float(state["gap_participation"])
    fill = float(state["metal_fill_fraction"])
    if gap_fraction <= 0.0 or fill >= 1.0 - 1e-12:
        return 0.0
    available_void_depth = float(thickness_nm) * (1.0 - fill)
    if available_void_depth <= 1e-12:
        return 0.0
    equivalent_gap_adsorbate = (
        gap_fraction * float(adsorbate_thickness_nm))
    return float(min(0.95, equivalent_gap_adsorbate / available_void_depth))


def prism_index(name, wavenumber_cm, custom_n=1.515):
    """Real prism index from compact Sellmeier expressions (lambda in um)."""
    lam = 1.0e4 / np.asarray(wavenumber_cm, dtype=float)
    l2 = lam**2
    if name == "Si":
        n2 = 11.6858 + 0.939816 / l2 + 0.00810461 * l2 / (l2 - 1.1071)
    elif name == "Ge":
        n2 = (9.28156 + 6.72880 * l2 / (l2 - 0.44105)
              + 0.21307 * l2 / (l2 - 3870.1))
    elif name == "ZnSe":
        n2 = (1.0
              + 4.45813734 * l2 / (l2 - 0.200859853**2)
              + 0.467216334 * l2 / (l2 - 0.391371166**2)
              + 2.89566290 * l2 / (l2 - 47.1362108**2))
    elif name == "Custom n":
        n2 = np.asarray(float(custom_n) ** 2)
    else:
        raise ValueError("Unknown prism: %s" % name)
    return np.sqrt(np.maximum(n2, 1.0))


def critical_angle_deg(prism, wavenumber_cm, n_solution, custom_prism_n=1.515):
    n_prism = float(prism_index(prism, wavenumber_cm, custom_prism_n))
    ratio = float(n_solution) / n_prism
    if not 0.0 < ratio < 1.0:
        return float("nan")
    return math.degrees(math.asin(ratio))


def planar_spp_angle_deg(prism, wavenumber_cm, metal, n_solution, k_solution,
                         custom_prism_n=1.515,
                         metal_optical_model="Infrared Drude",
                         custom_metal_n=0.18, custom_metal_k=3.43):
    """Approximate metal/solution planar-SPP phase-matching angle.

    This is a single-interface estimate.  It is a diagnostic rather than an
    assertion that a rough or discontinuous SEIRAS film supports that mode.
    """
    eps_m = complex(metal_eps(
        float(wavenumber_cm), metal, metal_optical_model,
        custom_metal_n, custom_metal_k))
    eps_s = complex(float(n_solution), float(k_solution)) ** 2
    n_spp = complex(_csqrt(eps_m * eps_s / (eps_m + eps_s)))
    n_prism = float(prism_index(prism, wavenumber_cm, custom_prism_n))
    ratio = n_spp.real / n_prism
    if not 0.0 < ratio < 1.0:
        return float("nan")
    return math.degrees(math.asin(ratio))


def lorentz_adsorbate_eps(wavenumber_cm, center_cm, fwhm_cm,
                          strength, n_background):
    """CO-layer dielectric function using one passive Lorentz oscillator.

    ``strength`` is the dimensionless Delta-epsilon coefficient.  At the band
    centre the oscillator is absorptive (positive Im(epsilon)).
    """
    wn = np.asarray(wavenumber_cm, dtype=float)
    wn0 = float(center_cm)
    gamma = float(fwhm_cm)
    denominator = wn0**2 - wn**2 - 1j * gamma * wn
    return float(n_background) ** 2 + float(strength) * wn0**2 / denominator


def rp_multilayer(theta_deg, wavenumber_cm, eps_layers, thicknesses_nm):
    """p-reflectivity of an isotropic planar multilayer stack.

    ``eps_layers`` contains incident medium, all finite layers, and substrate.
    ``thicknesses_nm`` contains one thickness for every finite internal layer.
    The recursive Rouard form is equivalent to a 2x2 transfer matrix.

    At exactly 90 degrees the ideal collimated grazing-incidence limit is set
    explicitly to R_p=1, avoiding a 0/0 numerical limit in the admittance.
    """
    theta_input = np.asarray(theta_deg, dtype=float)
    scalar_input = theta_input.ndim == 0
    theta_flat = np.atleast_1d(theta_input).astype(float)
    if len(eps_layers) < 2 or len(thicknesses_nm) != len(eps_layers) - 2:
        raise ValueError("Layer/thickness dimensions are inconsistent.")
    if np.any(theta_flat < 0.0) or np.any(theta_flat > 90.0):
        raise ValueError("Angles for the Fresnel model must be between 0 and 90 degrees.")

    grazing = np.isclose(theta_flat, 90.0, atol=1e-12, rtol=0.0)
    # A representable offset is needed here: nextafter(90, 0) becomes
    # indistinguishable from pi/2 after the degree-to-radian conversion and
    # can still make q0 exactly zero.  The returned grazing value is replaced
    # by its exact analytical limit below.
    theta_safe = np.where(grazing, 90.0 - 1.0e-5, theta_flat)
    theta = np.deg2rad(theta_safe)

    eps = [complex(value) for value in eps_layers]
    kx2 = eps[0] * np.sin(theta) ** 2
    q = [_csqrt(value - kx2) for value in eps]
    eta = [value / qj for value, qj in zip(eps, q)]
    interfaces = [
        (eta[j] - eta[j + 1]) / (eta[j] + eta[j + 1])
        for j in range(len(eps) - 1)
    ]

    r_eff = interfaces[-1]
    for j in range(len(eps) - 3, -1, -1):
        layer_index = j + 1
        beta = (2.0 * np.pi * float(wavenumber_cm)
                * float(thicknesses_nm[layer_index - 1]) * 1e-7
                * q[layer_index])
        phase = np.exp(2j * beta)
        r_eff = ((interfaces[j] + r_eff * phase)
                 / (1.0 + interfaces[j] * r_eff * phase))

    reflectivity = np.abs(r_eff) ** 2
    reflectivity = np.where(grazing, 1.0, reflectivity)
    reflectivity = np.real_if_close(reflectivity).astype(float)
    if scalar_input:
        return float(reflectivity[0])
    return reflectivity.reshape(theta_input.shape)


def p_multilayer_power_budget(theta_deg, wavenumber_cm, eps_layers,
                              thicknesses_nm):
    """Return the p-polarized power budget of an isotropic multilayer.

    The incident tangential electric-field amplitude is one.  For each finite
    layer, absorption is obtained from the decrease in the time-averaged normal
    Poynting flux between its entrance and exit interfaces.  Consequently this
    separates film dissipation from transmission/absorption in the semi-infinite
    solution; unlike ``1-R``, it remains meaningful below the critical angle.

    ``layer_absorptance`` is ordered like ``thicknesses_nm``.  For an EMA film,
    the first value is the total dissipation of the homogenized metal/void layer,
    not a constituent-resolved metal-only loss.
    """
    theta_input = np.asarray(theta_deg, dtype=float)
    scalar_input = theta_input.ndim == 0
    theta_flat = np.atleast_1d(theta_input).astype(float)
    if len(eps_layers) < 2 or len(thicknesses_nm) != len(eps_layers) - 2:
        raise ValueError("Layer/thickness dimensions are inconsistent.")
    if np.any(theta_flat < 0.0) or np.any(theta_flat > 90.0):
        raise ValueError(
            "Angles for the Fresnel model must be between 0 and 90 degrees.")

    grazing = np.isclose(theta_flat, 90.0, atol=1e-12, rtol=0.0)
    theta_safe = np.where(grazing, 90.0 - 1.0e-5, theta_flat)
    theta = np.deg2rad(theta_safe)
    eps = [complex(value) for value in eps_layers]
    kx2 = eps[0] * np.sin(theta) ** 2
    q = [_csqrt(value - kx2) for value in eps]
    eta = [value / qj for value, qj in zip(eps, q)]
    interfaces = [
        (eta[j] - eta[j + 1]) / (eta[j] + eta[j + 1])
        for j in range(len(eps) - 1)
    ]

    r_eff = interfaces[-1]
    for j in range(len(eps) - 3, -1, -1):
        layer_index = j + 1
        beta = (2.0 * np.pi * float(wavenumber_cm)
                * float(thicknesses_nm[layer_index - 1]) * 1e-7
                * q[layer_index])
        phase = np.exp(2j * beta)
        r_eff = ((interfaces[j] + r_eff * phase)
                 / (1.0 + interfaces[j] * r_eff * phase))

    incident_flux = np.real(eta[0])
    valid_flux = np.abs(incident_flux) > 1e-14
    denominator = np.where(valid_flux, incident_flux, 1.0)
    electric = 1.0 + r_eff
    magnetic = eta[0] * (1.0 - r_eff)
    layer_absorptance = []

    for layer_index in range(1, len(eps) - 1):
        flux_top = np.real(electric * np.conj(magnetic)) / denominator
        forward = 0.5 * (electric + magnetic / eta[layer_index])
        backward = 0.5 * (electric - magnetic / eta[layer_index])
        beta = (2.0 * np.pi * float(wavenumber_cm)
                * float(thicknesses_nm[layer_index - 1]) * 1e-7
                * q[layer_index])
        forward_bottom = forward * np.exp(1j * beta)
        backward_bottom = backward * np.exp(-1j * beta)
        electric = forward_bottom + backward_bottom
        magnetic = eta[layer_index] * (forward_bottom - backward_bottom)
        flux_bottom = np.real(electric * np.conj(magnetic)) / denominator
        absorbed = np.real(flux_top - flux_bottom)
        absorbed = np.where(grazing | ~valid_flux, 0.0, absorbed)
        # Roundoff around a passive zero-loss layer can produce tiny negatives.
        absorbed = np.where(
            (absorbed < 0.0) & (absorbed > -1e-10), 0.0, absorbed)
        layer_absorptance.append(absorbed)

    reflectivity = np.abs(r_eff) ** 2
    transmitted = np.real(electric * np.conj(magnetic)) / denominator
    reflectivity = np.where(grazing, 1.0, reflectivity)
    transmitted = np.where(grazing | ~valid_flux, 0.0, transmitted)
    closure = reflectivity + transmitted
    for absorbed in layer_absorptance:
        closure = closure + absorbed

    def shaped(values):
        values = np.real_if_close(values).astype(float)
        if scalar_input:
            return float(values[0])
        return values.reshape(theta_input.shape)

    return {
        "reflectivity": shaped(reflectivity),
        "transmitted_power": shaped(transmitted),
        "layer_absorptance": tuple(shaped(value)
                                   for value in layer_absorptance),
        "energy_closure": shaped(closure),
    }


def rp_stack(theta_deg, wavenumber_cm, prism, metal, thickness_nm,
             n_solution, k_solution, include_adsorbate=False,
             resonant_adsorbate=False, adsorbate_center_cm=2050.0,
             adsorbate_fwhm_cm=20.0, adsorbate_strength=0.002,
             adsorbate_thickness_nm=1.0, adsorbate_n_background=1.45,
             custom_prism_n=1.515,
             metal_optical_model="Infrared Drude",
             custom_metal_n=0.18, custom_metal_k=3.43,
             film_morphology=MORPHOLOGY_CONTINUOUS,
             ema_fill_fraction=0.45,
             connected_fill_fraction=0.90,
             percolation_thickness_nm=10.0,
             gap_field_gain=20.0, return_power_budget=False):
    """Reflectivity of the planar or semi-quantitative porous-film stack.

    In the isolated model the CO oscillator is placed in the explicit surface
    layer.  In the connected model the same equivalent adsorbate amount is
    moved into the EMA voids.  ``gap_field_gain`` scales only the resonant
    dielectric perturbation in those voids and is therefore an intensity-level
    proxy for unresolved inter-island hot spots.  It must not be interpreted as
    a fitted absolute enhancement factor.
    """
    eps_solution = complex(float(n_solution), float(k_solution)) ** 2
    eps_bulk_metal = complex(metal_eps(
        float(wavenumber_cm), metal, metal_optical_model,
        custom_metal_n, custom_metal_k))
    state = morphology_state(
        film_morphology, thickness_nm, ema_fill_fraction,
        connected_fill_fraction, percolation_thickness_nm)
    morphology = state["morphology"]
    gap_participation = float(state["gap_participation"])
    fill = float(state["metal_fill_fraction"])
    eps_ads_background = complex(float(adsorbate_n_background) ** 2)
    pore_occupancy = 0.0

    if morphology == MORPHOLOGY_CONTINUOUS:
        eps_film = eps_bulk_metal
    else:
        pore_occupancy = pore_adsorbate_occupancy(
            state, thickness_nm, adsorbate_thickness_nm)
        eps_void = eps_solution
        if include_adsorbate and pore_occupancy > 0.0:
            # The non-resonant molecular background is present in both sample
            # and reference.  Only the Lorentz perturbation receives the local
            # gap-field correction.
            eps_void_reference = (
                eps_solution
                + pore_occupancy * (eps_ads_background - eps_solution))
            eps_void = eps_void_reference
            if resonant_adsorbate:
                eps_ads_resonant = complex(lorentz_adsorbate_eps(
                    float(wavenumber_cm), adsorbate_center_cm,
                    adsorbate_fwhm_cm, adsorbate_strength,
                    adsorbate_n_background))
                eps_void = (
                    eps_void_reference
                    + pore_occupancy * float(gap_field_gain)
                    * (eps_ads_resonant - eps_ads_background))
        eps_film = bruggeman_eps(eps_bulk_metal, eps_void, fill)

    eps_layers = [
        complex(float(prism_index(
            prism, wavenumber_cm, custom_prism_n)) ** 2),
        eps_film,
    ]
    thicknesses = [float(thickness_nm)]
    if include_adsorbate:
        if resonant_adsorbate:
            eps_ads = complex(lorentz_adsorbate_eps(
                float(wavenumber_cm), adsorbate_center_cm,
                adsorbate_fwhm_cm,
                float(adsorbate_strength) * (1.0 - gap_participation),
                adsorbate_n_background))
        else:
            eps_ads = eps_ads_background
        eps_layers.append(eps_ads)
        thicknesses.append(float(adsorbate_thickness_nm))
    eps_layers.append(eps_solution)
    if return_power_budget:
        return p_multilayer_power_budget(
            theta_deg, wavenumber_cm, eps_layers, thicknesses)
    return rp_multilayer(theta_deg, wavenumber_cm, eps_layers, thicknesses)


def rp_three_layer(theta_deg, wavenumber_cm, prism, metal,
                   thickness_nm, n_solution, k_solution):
    """Backward-compatible prism|metal|solution wrapper."""
    return rp_stack(theta_deg, wavenumber_cm, prism, metal, thickness_nm,
                    n_solution, k_solution, include_adsorbate=False)


def gaussian_angular_average(values, step_deg, fwhm_deg):
    """Gaussian intensity average with normalized physical boundary handling.

    Both numerator and a valid-angle mask are convolved.  Renormalization avoids
    the old zero-padding artifact that forced reflectivity downward near 90 deg.
    A finite-divergence beam centred at 90 deg is a truncated average of rays
    below 90 deg, so only the ideal collimated curve must reach exactly one.
    """
    values = np.asarray(values, dtype=float)
    if fwhm_deg <= 1e-12:
        return values.copy()
    sigma = float(fwhm_deg) / (2.0 * math.sqrt(2.0 * math.log(2.0)))
    half = max(1, int(math.ceil(4.0 * sigma / float(step_deg))))
    x = np.arange(-half, half + 1, dtype=float) * float(step_deg)
    kernel = np.exp(-0.5 * (x / sigma) ** 2)
    kernel /= kernel.sum()
    numerator = np.convolve(values, kernel, mode="same")
    denominator = np.convolve(np.ones_like(values), kernel, mode="same")
    return numerator / np.maximum(denominator, 1e-300)


def parse_thicknesses(value):
    if isinstance(value, str):
        tokens = [part.strip() for part in value.replace(";", ",").split(",")]
        numbers = [float(token) for token in tokens if token]
    else:
        numbers = [float(item) for item in value]
    if not numbers:
        raise ValueError("Enter at least one metal thickness.")
    if len(numbers) > len(CURVE_COLORS):
        raise ValueError("Use at most %d thicknesses." % len(CURVE_COLORS))
    if any(not 0.0 < number <= 200.0 for number in numbers):
        raise ValueError("Metal thicknesses must be >0 and <=200 nm.")
    return tuple(numbers)


def film_cases(metal, thicknesses_nm, comparison_styles=False):
    thicknesses = parse_thicknesses(thicknesses_nm)
    if comparison_styles and len(thicknesses) != 2:
        raise ValueError("Enter exactly two thicknesses for each material.")
    material_color = MATERIAL_COLORS.get(str(metal))
    return tuple(
        FilmCase(
            str(metal), float(thickness),
            (material_color if comparison_styles and material_color
             else CURVE_COLORS[index]),
            ("-" if index == 0 else "--") if comparison_styles else "-")
        for index, thickness in enumerate(thicknesses))


def comparison_film_cases(params):
    """Build a fair material overlay on one shared thickness grid."""
    thicknesses = parse_thicknesses(params["thicknesses_nm"])
    if len(thicknesses) > len(OVERLAY_LINESTYLES):
        raise ValueError(
            "Material overlay supports at most %d common thicknesses."
            % len(OVERLAY_LINESTYLES))
    return tuple(
        FilmCase(metal, thickness, MATERIAL_COLORS[metal],
                 OVERLAY_LINESTYLES[index])
        for metal in ("Au", "Ag", "Cu")
        for index, thickness in enumerate(thicknesses)
    )


def _validate_params(params):
    wn = float(params["wavenumber_cm"])
    if not 100.0 <= wn <= 25000.0:
        raise ValueError("Wavenumber must be 100-25000 cm^-1.")
    if str(params["prism"]) not in ("Si", "Ge", "ZnSe", "Custom n"):
        raise ValueError("Unknown prism: %s" % params["prism"])
    if (str(params["prism"]) == "Custom n"
            and not 1.0 <= float(params["custom_prism_n"]) <= 5.0):
        raise ValueError("Custom prism n must be 1-5.")
    multi_metal = bool(params.get("multi_metal_comparison", False))
    if multi_metal:
        values = parse_thicknesses(params["thicknesses_nm"])
        if len(values) > len(OVERLAY_LINESTYLES):
            raise ValueError(
                "Use at most %d common thicknesses for the Au/Ag/Cu overlay."
                % len(OVERLAY_LINESTYLES))
    elif str(params["metal"]) not in tuple(METAL_DRUDE) + ("Custom",):
        raise ValueError("Unknown metal: %s" % params["metal"])
    optical_model = str(params["metal_optical_model"])
    if optical_model not in ("Infrared Drude", "Custom n,k"):
        raise ValueError("Unknown metal optical model: %s" % optical_model)
    if (not multi_metal and str(params["metal"]) == "Custom"
            and optical_model != "Custom n,k"):
        raise ValueError("Custom material requires the Custom n,k optical model.")
    if optical_model == "Infrared Drude" and wn > 10000.0:
        raise ValueError(
            "The bundled Drude constants are an infrared baseline only. "
            "For a visible SPR example, select Custom n,k and enter "
            "metal optical constants at the chosen wavelength.")
    if optical_model == "Custom n,k":
        if not 0.0 <= float(params["custom_metal_n"]) <= 20.0:
            raise ValueError("Custom metal n must be 0-20.")
        if not 0.0 <= float(params["custom_metal_k"]) <= 20.0:
            raise ValueError("Custom metal k must be 0-20.")
    morphology = canonical_morphology_name(params["film_morphology"])
    if morphology not in MORPHOLOGY_CHOICES:
        raise ValueError("Unknown film morphology: %s" % morphology)
    isolated_fill = float(params["ema_fill_fraction"])
    connected_fill = float(params["connected_fill_fraction"])
    if not 0.01 <= isolated_fill <= 0.98:
        raise ValueError("Isolated-film EMA fill must be 0.01-0.98.")
    if not 0.01 <= connected_fill <= 0.98:
        raise ValueError("Connected-film EMA fill must be 0.01-0.98.")
    if connected_fill < isolated_fill:
        raise ValueError(
            "Connected-film fill must be >= isolated-film fill.")
    if not 0.5 <= float(params["percolation_thickness_nm"]) <= 100.0:
        raise ValueError("Percolation midpoint must be 0.5-100 nm.")
    if not 1.0 <= float(params["gap_field_gain"]) <= 100.0:
        raise ValueError("Gap-field intensity gain must be 1-100.")
    ns, ks = float(params["n_solution"]), float(params["k_solution"])
    if not 0.1 <= ns <= 5.0 or not 0.0 <= ks <= 1.0:
        raise ValueError("Solution n or k is outside the supported range.")
    amin, amax = float(params["angle_min_deg"]), float(params["angle_max_deg"])
    if not 0.0 <= amin < amax <= 90.0:
        raise ValueError("Angles must satisfy 0 <= min < max <= 90 degrees.")
    if not 0.0 <= float(params["divergence_fwhm_deg"]) <= 20.0:
        raise ValueError("Angular FWHM must be 0-20 degrees.")
    if not multi_metal:
        parse_thicknesses(params["thicknesses_nm"])
    if bool(params["include_adsorbate"]):
        if not 100.0 <= float(params["adsorbate_center_cm"]) <= 25000.0:
            raise ValueError("Adsorbate band centre must be 100-25000 cm^-1.")
        if not 0.1 <= float(params["adsorbate_fwhm_cm"]) <= 1000.0:
            raise ValueError("Adsorbate FWHM must be 0.1-1000 cm^-1.")
        if not 0.0 < float(params["adsorbate_strength"]) <= 1.0:
            raise ValueError("Adsorbate oscillator strength must be >0 and <=1.")
        if not 0.01 <= float(params["adsorbate_thickness_nm"]) <= 100.0:
            raise ValueError("Adsorbate thickness must be 0.01-100 nm.")
        if not 1.0 <= float(params["adsorbate_n_background"]) <= 3.0:
            raise ValueError("Adsorbate background n must be 1-3.")


def simulate(params):
    """Calculate raw reflectivity and optional CO differential absorbance."""
    p = dict(DEFAULTS)
    p.update(params)
    p["thicknesses_nm"] = parse_thicknesses(p["thicknesses_nm"])
    for key in ("au_thicknesses_nm", "ag_thicknesses_nm",
                "cu_thicknesses_nm"):
        p[key] = parse_thicknesses(p[key])
    p["film_morphology"] = canonical_morphology_name(p["film_morphology"])
    _validate_params(p)

    wn = float(p["wavenumber_cm"])
    prism = str(p["prism"])
    ns, ks = float(p["n_solution"]), float(p["k_solution"])
    amin, amax = float(p["angle_min_deg"]), float(p["angle_max_deg"])
    fwhm = float(p["divergence_fwhm_deg"])
    include_ads = bool(p["include_adsorbate"])
    multi_metal = bool(p.get("multi_metal_comparison", False))
    cases = (comparison_film_cases(p) if multi_metal
             else film_cases(p["metal"], p["thicknesses_nm"]))

    # The full physical angular domain is used before convolution.  This makes
    # the 90-degree endpoint a normalized truncated angular average instead of
    # a zero-padded numerical cliff.
    step = 0.02
    theta_full = np.arange(0.0, 90.0 + 0.5 * step, step)
    theta = np.arange(amin, amax + 0.5 * step, step)
    theta[-1] = amax
    theta_c = critical_angle_deg(
        prism, wn, ns, float(p["custom_prism_n"]))
    spp_metals = ("Au", "Ag", "Cu") if multi_metal else (str(p["metal"]),)
    theta_spp_by_metal = {
        metal: planar_spp_angle_deg(
            prism, wn, metal, ns, ks,
            custom_prism_n=float(p["custom_prism_n"]),
            metal_optical_model=str(p["metal_optical_model"]),
            custom_metal_n=float(p["custom_metal_n"]),
            custom_metal_k=float(p["custom_metal_k"]))
        for metal in spp_metals
    }
    theta_spp = theta_spp_by_metal[spp_metals[0]]

    # Metal-free ATR reference with the same equivalent planar CO amount.  Its
    # ratio to the metal-containing DeltaA is exported as a diagnostic only;
    # absolute DeltaA remains the primary experimentally observable signal.
    if include_ads:
        eps_prism = complex(float(prism_index(
            prism, wn, float(p["custom_prism_n"]))) ** 2)
        eps_solution = complex(ns, ks) ** 2
        eps_ads_background = complex(float(p["adsorbate_n_background"]) ** 2)
        eps_ads_resonant = complex(lorentz_adsorbate_eps(
            wn, float(p["adsorbate_center_cm"]),
            float(p["adsorbate_fwhm_cm"]),
            float(p["adsorbate_strength"]),
            float(p["adsorbate_n_background"])))
        bare_ref_full = rp_multilayer(
            theta_full, wn,
            [eps_prism, eps_ads_background, eps_solution],
            [float(p["adsorbate_thickness_nm"])])
        bare_sample_full = rp_multilayer(
            theta_full, wn,
            [eps_prism, eps_ads_resonant, eps_solution],
            [float(p["adsorbate_thickness_nm"])])
        bare_ref_avg_full = gaussian_angular_average(
            bare_ref_full, step, fwhm)
        bare_sample_avg_full = gaussian_angular_average(
            bare_sample_full, step, fwhm)
        bare_delta_a_ideal = -np.log10(
            np.maximum(np.interp(theta, theta_full, bare_sample_full), 1e-300)
            / np.maximum(np.interp(theta, theta_full, bare_ref_full), 1e-300))
        bare_delta_a = -np.log10(
            np.maximum(np.interp(theta, theta_full, bare_sample_avg_full), 1e-300)
            / np.maximum(np.interp(theta, theta_full, bare_ref_avg_full), 1e-300))
    else:
        bare_delta_a_ideal = np.full_like(theta, np.nan)
        bare_delta_a = np.full_like(theta, np.nan)

    curves = []
    for case in cases:
        film_state = morphology_state(
            p["film_morphology"], case.thickness_nm,
            float(p["ema_fill_fraction"]),
            float(p["connected_fill_fraction"]),
            float(p["percolation_thickness_nm"]))
        pore_occupancy = pore_adsorbate_occupancy(
            film_state, case.thickness_nm,
            float(p["adsorbate_thickness_nm"])) if include_ads else 0.0
        common = dict(
            wavenumber_cm=wn,
            prism=prism,
            metal=case.metal,
            thickness_nm=case.thickness_nm,
            n_solution=ns,
            k_solution=ks,
            custom_prism_n=float(p["custom_prism_n"]),
            metal_optical_model=str(p["metal_optical_model"]),
            custom_metal_n=float(p["custom_metal_n"]),
            custom_metal_k=float(p["custom_metal_k"]),
            film_morphology=str(p["film_morphology"]),
            ema_fill_fraction=float(p["ema_fill_fraction"]),
            connected_fill_fraction=float(p["connected_fill_fraction"]),
            percolation_thickness_nm=float(p["percolation_thickness_nm"]),
            gap_field_gain=float(p["gap_field_gain"]),
            include_adsorbate=include_ads,
            adsorbate_center_cm=float(p["adsorbate_center_cm"]),
            adsorbate_fwhm_cm=float(p["adsorbate_fwhm_cm"]),
            adsorbate_strength=float(p["adsorbate_strength"]),
            adsorbate_thickness_nm=float(p["adsorbate_thickness_nm"]),
            adsorbate_n_background=float(p["adsorbate_n_background"]),
        )
        rp_ref_full = rp_stack(theta_full, resonant_adsorbate=False, **common)
        rp_ref_avg_full = gaussian_angular_average(rp_ref_full, step, fwhm)
        power_budget_full = rp_stack(
            theta_full, resonant_adsorbate=False,
            return_power_budget=True, **common)
        film_absorptance_ideal_full = np.asarray(
            power_budget_full["layer_absorptance"][0], dtype=float)
        film_absorptance_avg_full = gaussian_angular_average(
            film_absorptance_ideal_full, step, fwhm)
        transmitted_power_full = np.asarray(
            power_budget_full["transmitted_power"], dtype=float)
        other_layer_absorptance_full = np.zeros_like(theta_full)
        for layer_absorptance in power_budget_full["layer_absorptance"][1:]:
            other_layer_absorptance_full += np.asarray(
                layer_absorptance, dtype=float)
        downstream_power_avg_full = gaussian_angular_average(
            transmitted_power_full + other_layer_absorptance_full,
            step, fwhm)

        if include_ads:
            rp_sample_full = rp_stack(theta_full, resonant_adsorbate=True, **common)
            rp_sample_avg_full = gaussian_angular_average(rp_sample_full, step, fwhm)
        else:
            rp_sample_full = rp_ref_full.copy()
            rp_sample_avg_full = rp_ref_avg_full.copy()

        rp_ref_ideal = np.interp(theta, theta_full, rp_ref_full)
        rp_ref = np.interp(theta, theta_full, rp_ref_avg_full)
        rp_sample_ideal = np.interp(theta, theta_full, rp_sample_full)
        rp_sample = np.interp(theta, theta_full, rp_sample_avg_full)
        film_absorptance_ideal = np.interp(
            theta, theta_full, film_absorptance_ideal_full)
        film_absorptance = np.interp(
            theta, theta_full, film_absorptance_avg_full)
        for array in (rp_ref_ideal, rp_ref, rp_sample_ideal, rp_sample):
            np.clip(array, 1e-300, 1.0, out=array)
        # Passive numerical roundoff can be at the 1e-14 level.
        film_absorptance_ideal = np.maximum(film_absorptance_ideal, 0.0)
        film_absorptance = np.maximum(film_absorptance, 0.0)

        if include_ads:
            delta_a_ideal = -np.log10(rp_sample_ideal / rp_ref_ideal)
            delta_a = -np.log10(rp_sample / rp_ref)
            enhancement_factor_ideal = np.full_like(theta, np.nan)
            enhancement_factor = np.full_like(theta, np.nan)
            valid_ideal = np.abs(bare_delta_a_ideal) > 1e-12
            valid = np.abs(bare_delta_a) > 1e-12
            enhancement_factor_ideal[valid_ideal] = (
                delta_a_ideal[valid_ideal] / bare_delta_a_ideal[valid_ideal])
            enhancement_factor[valid] = delta_a[valid] / bare_delta_a[valid]
            strongest_index = int(np.argmax(np.abs(delta_a)))
            strongest_angle = float(theta[strongest_index])
            strongest_delta = float(delta_a[strongest_index])
        else:
            delta_a_ideal = np.full_like(theta, np.nan)
            delta_a = np.full_like(theta, np.nan)
            enhancement_factor_ideal = np.full_like(theta, np.nan)
            enhancement_factor = np.full_like(theta, np.nan)
            strongest_angle = float("nan")
            strongest_delta = float("nan")

        minimum_index = int(np.argmin(rp_ref))
        coupling_index = int(np.argmax(film_absorptance))
        coupling_max = float(film_absorptance[coupling_index])
        coupling_at_60 = float(np.interp(
            60.0, theta_full, film_absorptance_avg_full))
        reflected_at_60 = float(np.interp(60.0, theta_full, rp_ref_avg_full))
        downstream_at_60 = float(np.interp(
            60.0, theta_full, downstream_power_avg_full))
        coupling_efficiency_60 = float("nan")
        if (float(p["angle_min_deg"]) <= 60.0 <= float(p["angle_max_deg"])
                and coupling_max > 1e-15):
            coupling_efficiency_60 = coupling_at_60 / coupling_max
        curves.append({
            "case": case,
            "rp_ref_ideal": rp_ref_ideal,
            "rp_ref": rp_ref,
            "rp_sample_ideal": rp_sample_ideal,
            "rp_sample": rp_sample,
            "delta_a_ideal": delta_a_ideal,
            "delta_a": delta_a,
            "enhancement_factor_ideal": enhancement_factor_ideal,
            "enhancement_factor": enhancement_factor,
            "rp_min_angle_deg": float(theta[minimum_index]),
            "rp_min": float(rp_ref[minimum_index]),
            "rp_grazing_ideal": float(rp_ref_full[-1]),
            "rp_grazing_averaged": float(rp_ref_avg_full[-1]),
            "film_absorptance_ideal": film_absorptance_ideal,
            "film_absorptance": film_absorptance,
            "film_absorptance_60": coupling_at_60,
            "film_absorptance_max": coupling_max,
            "film_coupling_max_angle_deg": float(theta[coupling_index]),
            "film_coupling_efficiency_60": coupling_efficiency_60,
            "power_reflected_60": reflected_at_60,
            "power_downstream_60": downstream_at_60,
            "power_budget_max_closure_error": float(np.max(np.abs(
                np.asarray(power_budget_full["energy_closure"], dtype=float)
                - 1.0))),
            "strongest_delta_angle_deg": strongest_angle,
            "strongest_delta_a": strongest_delta,
            "connection_fraction": film_state["connection_fraction"],
            "effective_metal_fill": film_state["metal_fill_fraction"],
            "gap_participation": film_state["gap_participation"],
            "pore_adsorbate_occupancy": pore_occupancy,
        })

    return {
        "params": p,
        "theta_deg": theta,
        "theta_c_deg": theta_c,
        "theta_spp_deg": theta_spp,
        "theta_spp_deg_by_metal": theta_spp_by_metal,
        "bare_delta_a_ideal": bare_delta_a_ideal,
        "bare_delta_a": bare_delta_a,
        "curves": curves,
    }


def _add_overlay_legends(ax, params, material_loc="lower right",
                         thickness_loc="center right", fontsize=8.0):
    """Add separate colour/material and line-style/thickness legends."""
    material_handles = [
        ax.plot([], [], color=MATERIAL_COLORS[metal], lw=2.0,
                ls="-", label=metal)[0]
        for metal in ("Au", "Ag", "Cu")
    ]
    thickness_handles = [
        ax.plot([], [], color="0.25", lw=1.8,
                ls=OVERLAY_LINESTYLES[index], label="%g nm" % thickness)[0]
        for index, thickness in enumerate(
            parse_thicknesses(params["thicknesses_nm"]))
    ]
    material_legend = ax.legend(
        handles=material_handles, loc=material_loc, title="Material (colour)",
        fontsize=fontsize, title_fontsize=fontsize, framealpha=0.94)
    ax.add_artist(material_legend)
    return ax.legend(
        handles=thickness_handles, loc=thickness_loc,
        title="Thickness (line style)", fontsize=fontsize,
        title_fontsize=fontsize, framealpha=0.94)


def build_figure(result, figure=None):
    if figure is None:
        figure = Figure(figsize=(11.2, 7.6), dpi=110, constrained_layout=False)
    else:
        figure.clear()
        figure.set_layout_engine(None)
    ax_rp, ax_delta = figure.subplots(
        2, 1, sharex=True, gridspec_kw={"height_ratios": (1.15, 1.0)})
    theta = result["theta_deg"]
    p = result["params"]
    fwhm = float(p["divergence_fwhm_deg"])
    include_ads = bool(p["include_adsorbate"])
    multi_metal = bool(p.get("multi_metal_comparison", False))

    for curve in result["curves"]:
        case = curve["case"]
        ax_rp.plot(theta, curve["rp_ref"], lw=2.0, color=case.color,
                   ls=case.linestyle,
                   label=case.label)
        if include_ads:
            ax_delta.plot(theta, curve["delta_a"], lw=2.0,
                          color=case.color, ls=case.linestyle,
                          label=case.label)

    theta_c = result["theta_c_deg"]
    theta_spp = result["theta_spp_deg"]
    for ax in (ax_rp, ax_delta):
        if theta[0] < 65.0 and theta[-1] > 55.0:
            ax.axvspan(max(55.0, theta[0]), min(65.0, theta[-1]),
                       color="0.55", alpha=0.075, lw=0)
        if theta[0] <= 60.0 <= theta[-1]:
            ax.axvline(60.0, color="0.50", ls=(0, (2, 3)), lw=0.8)
        if np.isfinite(theta_c) and theta[0] <= theta_c <= theta[-1]:
            ax.axvline(theta_c, color="0.32", ls=":", lw=1.25)
        if (not multi_metal and np.isfinite(theta_spp)
                and theta[0] <= theta_spp <= theta[-1]):
            ax.axvline(theta_spp, color="0.12", ls="-.", lw=1.05)
        ax.grid(True, alpha=0.24)
        ax.set_xlim(float(p["angle_min_deg"]), float(p["angle_max_deg"]))

    if (not multi_metal and np.isfinite(theta_c) and np.isfinite(theta_spp)
            and theta[0] <= theta_c <= theta[-1]):
        spp_label = (r"bulk-metal planar $\theta_{SPP}$"
                     if p["film_morphology"] != MORPHOLOGY_CONTINUOUS
                     else r"planar $\theta_{SPP}$")
        ax_rp.text(
            theta_c + 0.5, 0.04,
            r"$\theta_c$=%.2f$^\circ$, %s $\approx$ %.2f$^\circ$"
            % (theta_c, spp_label, theta_spp),
            transform=ax_rp.get_xaxis_transform(), fontsize=8.5, va="bottom")
    elif not np.isfinite(theta_c):
        ax_rp.text(
            0.018, 0.055,
            "No TIR critical angle: prism n <= solution n",
            transform=ax_rp.transAxes, fontsize=9.0, color="#9C2F00",
            ha="left", va="bottom")
    elif multi_metal and theta[0] <= theta_c <= theta[-1]:
        ax_rp.text(
            theta_c + 0.35, 0.04, r"$\theta_c$=%.2f$^\circ$" % theta_c,
            transform=ax_rp.get_xaxis_transform(), fontsize=8.2,
            va="bottom", ha="left")

    min_rp = min(float(np.nanmin(curve["rp_ref"])) for curve in result["curves"])
    ax_rp.set_ylim(max(0.0, min_rp - 0.06), 1.025)
    ax_rp.set_ylabel(r"Reference p-reflectivity  $R_p$")
    ax_rp.set_title("Raw optical background (not adsorbate-specific absorption)", fontsize=10.5)
    if multi_metal:
        _add_overlay_legends(ax_rp, p)
    else:
        ax_rp.legend(loc="lower right", fontsize=8.5, ncol=2,
                     title=str(p["film_morphology"]))
    if multi_metal:
        style_text = ("angular average: FWHM %.1f deg\n"
                      "same thickness grid for Au, Ag, and Cu" % fwhm)
    else:
        style_text = ("angularly averaged beam (FWHM %.1f deg)" % fwhm
                      if fwhm > 1e-12 else "ideal collimated beam")
    ax_rp.text(0.018, 0.965, style_text, transform=ax_rp.transAxes,
               ha="left", va="top", fontsize=8.2, color="0.20")

    if float(p["angle_max_deg"]) >= 89.9:
        ax_rp.annotate(
            r"ideal $R_p(90^\circ)=1$",
            xy=(90.0, 1.0), xytext=(84.0, 0.975), fontsize=8.5,
            arrowprops={"arrowstyle": "->", "lw": 0.8, "color": "0.25"},
            ha="right", va="top")

    ax_delta.axhline(0.0, color="0.35", ls=":", lw=0.9)
    if include_ads:
        ax_delta.set_ylabel(r"CO differential absorbance  $\Delta A$")
        ax_delta.set_title(
            "Angularly averaged adsorbate-specific signal at %.0f cm$^{-1}$:  "
            r"$-\log_{10}(R_{CO}/R_{reference})$" % float(p["wavenumber_cm"]),
            fontsize=10.5)
        ax_delta.ticklabel_format(axis="y", style="sci", scilimits=(-3, 3))
        ax_delta.margins(y=0.12)
        if theta[0] <= 60.0 <= theta[-1]:
            ax_delta.text(
                54.5, 0.965, "common fixed-angle\nwindow (55-65 deg)",
                transform=ax_delta.get_xaxis_transform(), ha="right",
                va="top", fontsize=7.8, color="0.32")
    else:
        ax_delta.set_ylabel(r"Differential absorbance  $\Delta A$")
        ax_delta.text(
            0.5, 0.5,
            "Adsorbate layer disabled\nCO differential absorbance is not calculated",
            transform=ax_delta.transAxes, ha="center", va="center", fontsize=11)
        ax_delta.set_ylim(-1.0, 1.0)
        ax_delta.set_yticks([])
    ax_delta.set_xlabel("Internal incidence angle from surface normal (deg)")

    ads_text = ("CO model ON" if include_ads else "adsorbate model OFF")
    prism_text = ("custom n=%.4g" % float(p["custom_prism_n"])
                  if p["prism"] == "Custom n" else str(p["prism"]))
    if p["film_morphology"] == MORPHOLOGY_CONTINUOUS:
        morphology_text = "smooth continuous-film TMM"
    elif p["film_morphology"] == MORPHOLOGY_ISOLATED:
        morphology_text = ("isolated EMA, f=%.2f"
                           % float(p["ema_fill_fraction"]))
    elif p["film_morphology"] == MORPHOLOGY_CONNECTED:
        morphology_text = ("connected porous EMA, f=%.2f, gap gain=%.1f"
                           % (float(p["connected_fill_fraction"]),
                              float(p["gap_field_gain"])))
    else:
        morphology_text = ("thickness-linked EMA, f=%.2f->%.2f, "
                           "d_perc=%.1f nm, gap gain=%.1f"
                           % (float(p["ema_fill_fraction"]),
                              float(p["connected_fill_fraction"]),
                              float(p["percolation_thickness_nm"]),
                              float(p["gap_field_gain"])))
    model_title = ("ATR-SEIRAS" if include_ads else "ATR / SPR")
    material_title = ("Au / Ag / Cu, common thickness grid"
                      if multi_metal else "%s thickness series" % p["metal"])
    title = ("%s angle response: %s\n"
             "%s prism | %.0f cm$^{-1}$ | solution %.3f%+.4fi | %s | %s"
             % (model_title, material_title, prism_text, p["wavenumber_cm"],
                p["n_solution"], p["k_solution"], morphology_text, ads_text))
    figure.suptitle(title, fontsize=12, y=0.982)
    figure.subplots_adjust(left=0.105, right=0.985, top=0.875,
                           bottom=0.13, hspace=0.26)
    beam_text = ("Only the normalized %.1f-deg FWHM angular average is plotted."
                 % fwhm if fwhm > 1e-12
                 else "The plotted beam is ideal collimated (FWHM 0 deg).")
    model_note = (
        "  Semi-quantitative gap gain captures the reported angle trend; "
        "it is not a morphology-resolved FEM prediction."
        if p["film_morphology"] in (
            MORPHOLOGY_CONNECTED, MORPHOLOGY_THICKNESS_LINKED)
        else "")
    figure.text(
        0.5, 0.025, beam_text + model_note,
        ha="center", va="bottom", fontsize=8.2, color="0.25")
    return figure


def build_coupling_figure(result, figure=None):
    """Plot effective-film power dissipation and the usefulness of 60 degrees."""
    if figure is None:
        figure = Figure(figsize=(11.2, 7.6), dpi=110,
                        constrained_layout=False)
    else:
        figure.clear()
        figure.set_layout_engine(None)
    ax_angle, ax_eff = figure.subplots(
        2, 1, gridspec_kw={"height_ratios": (1.55, 1.0)})
    theta = result["theta_deg"]
    p = result["params"]
    multi_metal = bool(p.get("multi_metal_comparison", False))

    for curve in result["curves"]:
        case = curve["case"]
        ax_angle.plot(
            theta, curve["film_absorptance"], lw=2.0,
            color=case.color, ls=case.linestyle, label=case.label)
        ax_angle.plot(
            curve["film_coupling_max_angle_deg"],
            curve["film_absorptance_max"], marker="o", ms=3.5,
            color=case.color, mec=case.color, ls="none")

    theta_c = result["theta_c_deg"]
    if np.isfinite(theta_c) and theta[0] <= theta_c <= theta[-1]:
        ax_angle.axvline(theta_c, color="0.35", ls=":", lw=1.1)
        ax_angle.text(
            theta_c, 0.98, r"$\theta_c$", transform=ax_angle.get_xaxis_transform(),
            ha="right", va="top", fontsize=8.5, color="0.25")
    if theta[0] <= 60.0 <= theta[-1]:
        ax_angle.axvspan(
            max(55.0, theta[0]), min(65.0, theta[-1]),
            color="0.55", alpha=0.075, lw=0)
        ax_angle.axvline(60.0, color="0.45", ls=(0, (2, 3)), lw=1.0)
        ax_angle.text(
            60.0, 0.98, r"$60^\circ$", transform=ax_angle.get_xaxis_transform(),
            ha="center", va="top", fontsize=8.5, color="0.25")
    ax_angle.set_xlim(float(p["angle_min_deg"]), float(p["angle_max_deg"]))
    ymax = max(float(np.nanmax(curve["film_absorptance"]))
               for curve in result["curves"])
    ax_angle.set_ylim(0.0, min(1.02, max(0.05, 1.10 * ymax)))
    ax_angle.set_xlabel("Internal incidence angle from surface normal (deg)")
    ax_angle.set_ylabel(r"Effective-film absorptance  $A_{film}$")
    ax_angle.set_title(
        "Fraction of incident p-polarized power dissipated inside the finite film",
        fontsize=10.5)
    ax_angle.grid(True, alpha=0.24)
    if multi_metal:
        _add_overlay_legends(ax_angle, p)
    else:
        ax_angle.legend(loc="best", fontsize=8.0,
                        title=str(p["film_morphology"]))

    cases = [curve["case"] for curve in result["curves"]]
    film_at_60 = np.asarray([
        curve["film_absorptance_60"] for curve in result["curves"]],
        dtype=float)
    downstream_at_60 = np.asarray([
        curve["power_downstream_60"] for curve in result["curves"]],
        dtype=float)
    reflected_at_60 = np.asarray([
        curve["power_reflected_60"] for curve in result["curves"]],
        dtype=float)
    efficiencies = np.asarray([
        curve["film_coupling_efficiency_60"]
        for curve in result["curves"]], dtype=float)
    x = np.arange(len(cases), dtype=float)
    bars = ax_eff.bar(
        x, film_at_60, width=0.72,
        color=[case.color for case in cases], alpha=0.90,
        label=r"dissipated in film  $A_{film}$")
    ax_eff.bar(
        x, downstream_at_60, width=0.72, bottom=film_at_60,
        color="white", edgecolor="0.45", hatch="///", linewidth=0.7,
        label="downstream flux / other finite-layer loss")
    ax_eff.bar(
        x, reflected_at_60, width=0.72,
        bottom=film_at_60 + downstream_at_60,
        color="0.83", edgecolor="0.45", linewidth=0.7,
        label=r"reflected  $R_p$")
    ax_eff.axhline(1.0, color="0.35", ls=":", lw=0.9)
    ax_eff.set_ylim(0.0, 1.12)
    ax_eff.set_ylabel("Fraction of incident power")
    ax_eff.set_title(
        ("Reference-stack power budget at 60 degrees; labels give "
         r"$A_{film}(60^\circ)$ and $A_{film}(60^\circ)/A_{film,max}$"),
        fontsize=10.0)
    ax_eff.set_xticks(x)
    ax_eff.set_xticklabels(
        [case.label.replace(" ", "\n", 1) for case in cases],
        fontsize=7.8, rotation=0)
    ax_eff.grid(True, axis="y", alpha=0.22)
    ax_eff.legend(loc="upper center", bbox_to_anchor=(0.5, -0.24),
                  ncol=3, fontsize=7.3, frameon=False)
    for bar, curve, efficiency in zip(bars, result["curves"], efficiencies):
        if np.isfinite(efficiency):
            ax_eff.text(
                bar.get_x() + bar.get_width() / 2.0,
                1.025,
                "A=%.1f%%\neta60=%.0f%%" % (
                    100.0 * float(curve["film_absorptance_60"]),
                    100.0 * float(efficiency)),
                ha="center", va="bottom", fontsize=6.8)
        else:
            ax_eff.text(
                bar.get_x() + bar.get_width() / 2.0, 1.025,
                "A=%.1f%%\n60 outside scan" % (
                    100.0 * float(curve["film_absorptance_60"])),
                ha="center", va="bottom", fontsize=6.5)

    morphology_note = (
        "For EMA films this is total loss in the homogenized metal/void layer; "
        "it is not a metal-versus-void constituent decomposition."
        if p["film_morphology"] != MORPHOLOGY_CONTINUOUS else
        "For a smooth continuous film, this is the metal-film absorptance."
    )
    figure.suptitle(
        "Film-coupling power budget at %.0f cm$^{-1}$  |  %s"
        % (float(p["wavenumber_cm"]), str(p["film_morphology"])),
        fontsize=12, y=0.982)
    figure.text(
        0.5, 0.018,
        ("Computed from the Poynting-flux decrease across the film, not from "
         "$1-R_p$.  Downstream flux includes power entering the solution.  "
         + morphology_note),
        ha="center", va="bottom", fontsize=8.0, color="0.25")
    figure.subplots_adjust(
        left=0.105, right=0.985, top=0.90, bottom=0.19, hspace=0.34)
    return figure


def build_signal_export_figure(result, figure=None):
    """Build only the adsorbate-specific panel in a publication-ready 4:3 layout.

    The GUI deliberately retains both the optical-background and differential-
    signal panels.  Manual figure export uses this separate canvas so the saved
    file contains only the scientifically relevant second panel, with fonts and
    line weights sized for a 6.4-inch-wide manuscript figure.
    """
    if figure is None:
        figure = Figure(figsize=(6.4, 4.8), dpi=300,
                        constrained_layout=False)
    else:
        figure.clear()
        figure.set_size_inches(6.4, 4.8, forward=True)
        figure.set_layout_engine(None)

    ax = figure.subplots(1, 1)
    theta = result["theta_deg"]
    p = result["params"]
    include_ads = bool(p["include_adsorbate"])
    multi_metal = bool(p.get("multi_metal_comparison", False))

    if include_ads:
        for curve in result["curves"]:
            case = curve["case"]
            ax.plot(theta, curve["delta_a"], lw=1.8,
                    color=case.color, ls=case.linestyle,
                    label=case.label)
    else:
        ax.text(0.5, 0.5, "Adsorbate layer disabled\nNo CO differential signal",
                transform=ax.transAxes, ha="center", va="center",
                fontsize=10)
        ax.set_ylim(-1.0, 1.0)
        ax.set_yticks([])

    theta_c = result["theta_c_deg"]
    theta_spp = result["theta_spp_deg"]
    if theta[0] < 65.0 and theta[-1] > 55.0:
        ax.axvspan(max(55.0, theta[0]), min(65.0, theta[-1]),
                   color="0.55", alpha=0.075, lw=0)
    if theta[0] <= 60.0 <= theta[-1]:
        ax.axvline(60.0, color="0.50", ls=(0, (2, 3)), lw=0.8)
        ax.text(54.5, 0.965, "typical experimental\nwindow (55-65 deg)",
                transform=ax.get_xaxis_transform(), ha="right", va="top",
                fontsize=7.8, color="0.32")
    if np.isfinite(theta_c) and theta[0] <= theta_c <= theta[-1]:
        ax.axvline(theta_c, color="0.32", ls=":", lw=1.1)
        ax.text(theta_c, 0.985, r"$\theta_c$", transform=ax.get_xaxis_transform(),
                ha="right", va="top", fontsize=8.5, color="0.25")
    if (not multi_metal and np.isfinite(theta_spp)
            and theta[0] <= theta_spp <= theta[-1]):
        ax.axvline(theta_spp, color="0.12", ls="-.", lw=1.0)

    ax.axhline(0.0, color="0.38", ls=":", lw=0.8)
    ax.grid(True, color="0.82", alpha=0.55, lw=0.6)
    ax.set_xlim(float(p["angle_min_deg"]), float(p["angle_max_deg"]))
    ax.set_xlabel("Internal incidence angle from surface normal (deg)",
                  fontsize=10)
    ax.set_ylabel(r"CO differential absorbance  $\Delta A$", fontsize=10)
    title = ("Semi-quantitative CO response vs incidence angle"
             if p["film_morphology"] in (
                 MORPHOLOGY_CONNECTED, MORPHOLOGY_THICKNESS_LINKED)
             else "CO differential absorbance vs incidence angle")
    ax.set_title(title, fontsize=10.5, pad=8)
    ax.tick_params(axis="both", which="major", labelsize=8.5,
                   width=0.8, length=3.5)
    for spine in ax.spines.values():
        spine.set_linewidth(0.8)

    if include_ads:
        ax.ticklabel_format(axis="y", style="sci", scilimits=(-3, 3))
        ax.yaxis.get_offset_text().set_fontsize(8.5)
        ax.margins(y=0.12)
        if multi_metal:
            _add_overlay_legends(
                ax, p, material_loc="upper left",
                thickness_loc="upper right", fontsize=7.5)
        else:
            legend = ax.legend(loc="best", fontsize=7.8, ncol=2,
                               frameon=True, framealpha=0.94,
                               borderpad=0.45, handlelength=2.7,
                               columnspacing=1.0, handletextpad=0.55)
            legend.get_frame().set_linewidth(0.7)

    # Fixed margins preserve an exact 4:3 canvas; do not use bbox_inches="tight"
    # when saving this figure because it would change the exported aspect ratio.
    figure.subplots_adjust(left=0.14, right=0.975, bottom=0.145, top=0.91)
    return figure


def result_summary(result):
    p = result["params"]
    include_ads = bool(p["include_adsorbate"])
    multi_metal = bool(p.get("multi_metal_comparison", False))
    material_text = "Au / Ag / Cu comparison" if multi_metal else str(p["metal"])
    lines = [
        "ATR-SEIRAS incidence-angle calculation",
        "Stack: prism | %s %s | %s | solution" %
        (material_text, p["film_morphology"],
         "CO adsorbate" if include_ads else "no adsorbate"),
        "Angles are internal angles from the surface normal.",
        "",
        "Observation wavenumber = %.2f cm^-1" % p["wavenumber_cm"],
        "Prism = %s (n=%.5f)" % (
            p["prism"], float(prism_index(
                p["prism"], p["wavenumber_cm"], p["custom_prism_n"]))),
        "Metal optical model = %s" % p["metal_optical_model"],
        "Morphology model = %s" % p["film_morphology"],
        ("EMA fill isolated -> connected = %.3f -> %.3f"
         % (p["ema_fill_fraction"], p["connected_fill_fraction"])
         if p["film_morphology"] != MORPHOLOGY_CONTINUOUS
         else "EMA fill = not used for smooth continuous film"),
        ("Percolation midpoint = %.2f nm; gap-field intensity gain = %.2f"
         % (p["percolation_thickness_nm"], p["gap_field_gain"])
         if p["film_morphology"] in (
             MORPHOLOGY_CONNECTED, MORPHOLOGY_THICKNESS_LINKED)
         else "Gap-field correction = not used"),
        "theta_c = %.4f deg" % result["theta_c_deg"],
        (("bulk-interface theta_SPP: Au %.4f, Ag %.4f, Cu %.4f deg"
          % tuple(result["theta_spp_deg_by_metal"][metal]
                  for metal in ("Au", "Ag", "Cu")))
         if multi_metal else
         "bulk-metal/solution single-interface theta_SPP ~= %.4f deg"
         % result["theta_spp_deg"]),
        "angular FWHM = %.3f deg" % p["divergence_fwhm_deg"],
        "",
        "Film          connect  eff.fill  gap.CO   min-R angle   min R"
        + ("      strongest CO DeltaA" if include_ads else ""),
    ]
    for curve in result["curves"]:
        line = ("%-12s   %5.2f     %5.2f    %5.2f    %7.2f deg  %7.4f"
                % (curve["case"].label, curve["connection_fraction"],
                   curve["effective_metal_fill"],
                   curve["gap_participation"],
                   curve["rp_min_angle_deg"], curve["rp_min"]))
        if include_ads:
            line += "   %+.4e at %.2f deg" % (
                curve["strongest_delta_a"],
                curve["strongest_delta_angle_deg"])
        lines.append(line)
        efficiency = curve["film_coupling_efficiency_60"]
        efficiency_text = (
            "%.1f%%" % (100.0 * efficiency)
            if np.isfinite(efficiency) else "60 deg outside scan")
        lines.append(
            "             film coupling: A60=%.4f; max=%.4f at %.2f deg; "
            "60/max=%s; R60=%.4f; downstream60=%.4f" % (
                curve["film_absorptance_60"],
                curve["film_absorptance_max"],
                curve["film_coupling_max_angle_deg"],
                efficiency_text,
                curve["power_reflected_60"],
                curve["power_downstream_60"]))

    lines.extend([
        "",
        "Interpretation:",
        "- The upper Rp valley is total optical attenuation, not CO absorption.",
        "- The Film coupling tab reports Poynting-flux loss inside the finite film; it is not simply 1-Rp.",
        "- Afilm(60)/Afilm,max evaluates whether the common 60-deg setting is close to the best film-coupling angle in the plotted scan.",
        "- The lower DeltaA isolates the CO Lorentz oscillator by dividing by a vibration-free reference.",
        "- A genuine planar-SPP feature should be p-specific and lie just above theta_c.",
        "- Ideal Rp(90 deg)=1 is enforced from the Fresnel grazing-incidence limit.",
        "- Finite angular divergence is renormalized at the 0/90-deg boundaries; no zero-padding cliff is introduced.",
        "- Isolated curves place CO at the outer surface; connected curves move the same equivalent amount into EMA voids.",
        "- Gap-field gain is a bounded phenomenological intensity correction, not an absolute FEM enhancement factor.",
        "- Excel export contains only the angle and angularly averaged CO signal curves shown in the signal figure.",
        "- The thickness-linked connection fraction is a smooth morphology proxy and must be constrained by SEM/AFM for quantitative use.",
        "- If prism n <= solution n, no total-internal-reflection critical angle exists.",
        "- The bundled Drude constants and generic CO oscillator are baselines, not film-specific quantitative predictions.",
    ])
    return "\n".join(lines)


def angularly_averaged_signal_export_columns(result):
    """Return exactly the data plotted in the signal-only export figure."""
    theta = result["theta_deg"]
    headers = ["angle_deg"]
    columns = [theta]
    for curve in result["curves"]:
        stub = (curve["case"].label.lower().replace(" ", "_")
                .replace(".", "p"))
        headers.append("deltaA_CO_angularly_averaged_%s" % stub)
        columns.append(curve["delta_a"])
    return headers, columns


def export_xlsx(result, path):
    """Export one worksheet containing only angularly averaged signal data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    headers, columns = angularly_averaged_signal_export_columns(result)
    data = np.column_stack(columns)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Angularly averaged signal"
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for row in data:
        worksheet.append([
            float(value) if np.isfinite(value) else None for value in row
        ])
    worksheet.freeze_panes = "A2"
    for index, header in enumerate(headers, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            48, max(14, len(header) + 2))
    workbook.save(path)
    return path


def run_selftests(verbose=True):
    passed = 0
    failed = 0

    def check(name, condition, detail=""):
        nonlocal passed, failed
        if bool(condition):
            passed += 1
            if verbose:
                print("PASS", name)
        else:
            failed += 1
            print("FAIL", name, detail)

    tc = critical_angle_deg("Si", 2050.0, 1.30)
    check("T1 Si/solution critical angle", abs(tc - 22.3034) < 0.02, str(tc))

    theta = np.linspace(20.0, 90.0, 1401)
    rp = rp_three_layer(theta, 2050.0, "Si", "Pt", 20.0, 1.30, 0.005)
    check("T2 passive reflectivity bounds",
          np.all(np.isfinite(rp)) and np.min(rp) >= -1e-12
          and np.max(rp) <= 1.0 + 1e-9,
          "range=%r" % ((float(np.min(rp)), float(np.max(rp))),))

    check("T3 ideal grazing-incidence limit", abs(float(rp[-1]) - 1.0) < 1e-12,
          str(float(rp[-1])))
    rp_near = rp_three_layer(np.array([89.9, 89.99, 89.999]), 2050.0,
                             "Si", "Pt", 20.0, 1.30, 0.005)
    check("T4 approach to grazing reflection",
          rp_near[0] < rp_near[1] < rp_near[2] < 1.0,
          str(rp_near))

    # Zero metal thickness reduces to the direct prism/solution interface.
    theta_direct = np.linspace(30.0, 85.0, 551)
    rp0 = rp_three_layer(theta_direct, 2050.0, "Si", "Pt", 0.0, 1.30, 0.005)
    n0 = float(prism_index("Si", 2050.0))
    eps0, eps2 = n0**2, complex(1.30, 0.005)**2
    kx2 = eps0 * np.sin(np.deg2rad(theta_direct))**2
    q0, q2 = _csqrt(eps0 - kx2), _csqrt(eps2 - kx2)
    direct = np.abs((eps0 / q0 - eps2 / q2)
                    / (eps0 / q0 + eps2 / q2))**2
    check("T5 zero-thickness Fresnel limit",
          np.max(np.abs(rp0 - direct)) < 2e-10,
          str(float(np.max(np.abs(rp0 - direct)))))

    out = simulate(dict(DEFAULTS))
    labels = [curve["case"].label for curve in out["curves"]]
    check("T6 default Au/Ag/Cu common-thickness overlay",
          labels == [
              "Au 5 nm", "Au 10 nm",
              "Ag 5 nm", "Ag 10 nm",
              "Cu 5 nm", "Cu 10 nm"],
          str(labels))
    check("T7 CO differential signal is finite and nonzero",
          all(np.all(np.isfinite(curve["delta_a"]))
              and np.max(np.abs(curve["delta_a"])) > 1e-9
              for curve in out["curves"]))
    check("T8 all ideal default curves reach R(90)=1",
          all(abs(curve["rp_grazing_ideal"] - 1.0) < 1e-12
              for curve in out["curves"]))
    check("T9 finite-FWHM boundary has no zero-padding cliff",
          all(curve["rp_grazing_averaged"] > 0.85
              for curve in out["curves"]),
          str([curve["rp_grazing_averaged"] for curve in out["curves"]]))

    p_off = dict(DEFAULTS)
    p_off["include_adsorbate"] = False
    out_off = simulate(p_off)
    check("T10 adsorbate-off switch suppresses DeltaA",
          all(np.all(np.isnan(curve["delta_a"])) for curve in out_off["curves"]))

    p_zero_div = dict(DEFAULTS)
    p_zero_div["divergence_fwhm_deg"] = 0.0
    out_zero_div = simulate(p_zero_div)
    check("T11 zero-divergence ideal equality",
          all(np.max(np.abs(curve["rp_ref"] - curve["rp_ref_ideal"])) < 1e-12
              and np.max(np.abs(curve["delta_a"] - curve["delta_a_ideal"])) < 1e-12
              for curve in out_zero_div["curves"]))

    check("T12 custom thickness parser",
          parse_thicknesses("6, 12.5; 30") == (6.0, 12.5, 30.0))

    tc_bk7 = critical_angle_deg("Custom n", 1.0e7 / 633.0, 1.33, 1.515)
    check("T13 custom BK7-like prism critical angle",
          abs(tc_bk7 - 61.3886) < 0.01, str(tc_bk7))
    check("T14 no TIR when prism and solution indices match",
          np.isnan(critical_angle_deg("Custom n", 15798.0, 1.30, 1.30)))

    eps_m = complex(drude_eps(2050.0, "Pt"))
    eps_h = complex(1.30, 0.005) ** 2
    check("T15 Bruggeman endpoints",
          abs(bruggeman_eps(eps_m, eps_h, 0.0) - eps_h) < 1e-12
          and abs(bruggeman_eps(eps_m, eps_h, 1.0) - eps_m) < 1e-12)
    eps_ema = bruggeman_eps(eps_m, eps_h, 0.45)
    check("T16 passive finite island EMA",
          np.isfinite(eps_ema.real) and np.isfinite(eps_ema.imag)
          and eps_ema.imag >= 0.0, str(eps_ema))

    visible = dict(DEFAULTS)
    visible.update({
        "multi_metal_comparison": False,
        "wavenumber_cm": 1.0e7 / 633.0,
        "prism": "Custom n", "custom_prism_n": 1.515,
        "n_solution": 1.33, "k_solution": 0.0,
        "angle_min_deg": 60.0, "angle_max_deg": 90.0,
        "divergence_fwhm_deg": 0.0,
        "metal": "Au", "thicknesses_nm": (50.0,),
        "metal_optical_model": "Custom n,k",
        "custom_metal_n": 0.18, "custom_metal_k": 3.43,
        "film_morphology": "Continuous film",
        "include_adsorbate": False,
    })
    visible_out = simulate(visible)
    visible_curve = visible_out["curves"][0]
    visible_min_index = int(np.argmin(visible_curve["rp_ref"]))
    visible_min_angle = float(visible_out["theta_deg"][visible_min_index])
    check("T17 visible Au/BK7 SPR-like reflectivity minimum",
          70.0 < visible_min_angle < 74.0
          and float(visible_curve["rp_ref"][visible_min_index]) < 0.05,
          "angle=%.3f R=%.4g" % (
              visible_min_angle, visible_curve["rp_ref"][visible_min_index]))

    island = dict(DEFAULTS)
    island.update({
        "multi_metal_comparison": False,
        "include_adsorbate": False,
        "divergence_fwhm_deg": 0.0,
        "thicknesses_nm": (20.0,),
        "film_morphology": "Island film (Bruggeman EMA)",
        "ema_fill_fraction": 0.45,
    })
    island_rp = simulate(island)["curves"][0]["rp_ref"]
    continuous = dict(island)
    continuous["film_morphology"] = "Continuous film"
    continuous_rp = simulate(continuous)["curves"][0]["rp_ref"]
    check("T18 EMA morphology changes the optical response",
          np.max(np.abs(island_rp - continuous_rp)) > 0.05)

    export_figure = build_signal_export_figure(out)
    export_size = export_figure.get_size_inches()
    check("T19 signal-only export uses one 4:3 panel",
          len(export_figure.axes) == 1
          and abs(float(export_size[0] / export_size[1]) - 4.0 / 3.0) < 1e-12,
          "axes=%d, size=%s" % (len(export_figure.axes), export_size))

    check("T20 connected porous film is the default morphology",
          out["params"]["film_morphology"] == MORPHOLOGY_CONNECTED)
    check("T21 default curves use connected gap-CO morphology",
          all(curve["connection_fraction"] == 1.0
              and curve["gap_participation"] == 1.0
              and curve["pore_adsorbate_occupancy"] > 0.0
              for curve in out["curves"]))
    check("T22 pore occupancy remains physical",
          all(0.0 <= curve["pore_adsorbate_occupancy"] <= 0.95
              for curve in out["curves"]))
    check("T23 enhancement-factor diagnostic is finite at 60 deg",
          all(np.isfinite(np.interp(
              60.0, out["theta_deg"], curve["enhancement_factor"]))
              for curve in out["curves"]))

    custom = dict(DEFAULTS)
    custom.update({
        "multi_metal_comparison": False,
        "metal": "Custom", "thicknesses_nm": (40.0,),
        "metal_optical_model": "Custom n,k",
        "custom_metal_n": 0.20, "custom_metal_k": 3.20,
        "include_adsorbate": False,
    })
    custom_out = simulate(custom)
    check("T24 custom material and optical constants",
          len(custom_out["curves"]) == 1
          and custom_out["curves"][0]["case"].label == "Custom 40 nm"
          and np.all(np.isfinite(custom_out["curves"][0]["rp_ref"])))

    common_grid = {
        metal: tuple(curve["case"].thickness_nm for curve in out["curves"]
                     if curve["case"].metal == metal)
        for metal in ("Au", "Ag", "Cu")
    }
    check("T25 every material uses the same thickness grid",
          all(values == (5.0, 10.0)
              for values in common_grid.values()), str(common_grid))

    signal_headers, signal_columns = angularly_averaged_signal_export_columns(out)
    check("T30 data export contains averaged signal curves only",
          signal_headers == [
              "angle_deg",
              "deltaA_CO_angularly_averaged_au_5_nm",
              "deltaA_CO_angularly_averaged_au_10_nm",
              "deltaA_CO_angularly_averaged_ag_5_nm",
              "deltaA_CO_angularly_averaged_ag_10_nm",
              "deltaA_CO_angularly_averaged_cu_5_nm",
              "deltaA_CO_angularly_averaged_cu_10_nm",
          ] and len(signal_columns) == 7,
          str(signal_headers))

    budget_angles = np.linspace(15.0, 85.0, 701)
    budget = rp_stack(
        budget_angles, 2050.0, "Si", "Au", 20.0, 1.30, 0.0,
        include_adsorbate=False,
        film_morphology=MORPHOLOGY_CONTINUOUS,
        return_power_budget=True)
    budget_absorptance = np.asarray(
        budget["layer_absorptance"][0], dtype=float)
    budget_closure_error = float(np.max(np.abs(
        np.asarray(budget["energy_closure"], dtype=float) - 1.0)))
    check("T26 passive film power budget closes",
          budget_closure_error < 5e-10
          and np.min(budget_absorptance) >= -1e-10
          and np.max(budget_absorptance) <= 1.0 + 1e-10,
          "closure=%g A-range=%r" % (
              budget_closure_error,
              (float(np.min(budget_absorptance)),
               float(np.max(budget_absorptance)))))

    tir_budget = rp_stack(
        60.0, 2050.0, "Si", "Au", 20.0, 1.30, 0.0,
        include_adsorbate=False,
        film_morphology=MORPHOLOGY_CONTINUOUS,
        return_power_budget=True)
    tir_a = float(tir_budget["layer_absorptance"][0])
    check("T27 above-critical lossless stack has Afilm = 1-R",
          abs(tir_a - (1.0 - float(tir_budget["reflectivity"]))) < 5e-10
          and abs(float(tir_budget["transmitted_power"])) < 5e-10,
          str(tir_budget))

    check("T28 default 60-degree power-budget bars close",
          all(abs(curve["film_absorptance_60"]
                  + curve["power_reflected_60"]
                  + curve["power_downstream_60"] - 1.0) < 5e-9
              for curve in out["curves"]),
          str([(curve["case"].label,
                curve["film_absorptance_60"]
                + curve["power_reflected_60"]
                + curve["power_downstream_60"])
               for curve in out["curves"]]))
    check("T29 default 60-degree coupling efficiency is physical",
          all(0.0 <= curve["film_coupling_efficiency_60"] <= 1.0 + 1e-9
              for curve in out["curves"]))

    if verbose:
        print("\n%d passed, %d failed" % (passed, failed))
    return passed, failed


def launch_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    from tkinter.scrolledtext import ScrolledText

    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk

    root = tk.Tk()
    root.title("ATR-SEIRAS Au/Ag/Cu overlay and film-coupling budget")
    root.geometry("1580x1000")

    controls = ttk.Frame(root, padding=10)
    controls.pack(side="left", fill="y")
    plot_frame = ttk.Frame(root, padding=(0, 8, 8, 8))
    plot_frame.pack(side="right", fill="both", expand=True)

    variables = {
        "wavenumber_cm": tk.StringVar(value=str(DEFAULTS["wavenumber_cm"])),
        "prism": tk.StringVar(value=DEFAULTS["prism"]),
        "custom_prism_n": tk.StringVar(value=str(DEFAULTS["custom_prism_n"])),
        "n_solution": tk.StringVar(value=str(DEFAULTS["n_solution"])),
        "k_solution": tk.StringVar(value=str(DEFAULTS["k_solution"])),
        "angle_min_deg": tk.StringVar(value=str(DEFAULTS["angle_min_deg"])),
        "angle_max_deg": tk.StringVar(value=str(DEFAULTS["angle_max_deg"])),
        "divergence_fwhm_deg": tk.StringVar(value=str(DEFAULTS["divergence_fwhm_deg"])),
        "metal": tk.StringVar(value=DEFAULTS["metal"]),
        "thicknesses_nm": tk.StringVar(
            value=", ".join("%g" % value for value in DEFAULTS["thicknesses_nm"])),
        "metal_optical_model": tk.StringVar(value=DEFAULTS["metal_optical_model"]),
        "custom_metal_n": tk.StringVar(value=str(DEFAULTS["custom_metal_n"])),
        "custom_metal_k": tk.StringVar(value=str(DEFAULTS["custom_metal_k"])),
        "film_morphology": tk.StringVar(value=DEFAULTS["film_morphology"]),
        "ema_fill_fraction": tk.StringVar(value=str(DEFAULTS["ema_fill_fraction"])),
        "connected_fill_fraction": tk.StringVar(
            value=str(DEFAULTS["connected_fill_fraction"])),
        "percolation_thickness_nm": tk.StringVar(
            value=str(DEFAULTS["percolation_thickness_nm"])),
        "gap_field_gain": tk.StringVar(value=str(DEFAULTS["gap_field_gain"])),
        "include_adsorbate": tk.BooleanVar(value=DEFAULTS["include_adsorbate"]),
        "adsorbate_name": tk.StringVar(value=DEFAULTS["adsorbate_name"]),
        "adsorbate_center_cm": tk.StringVar(value=str(DEFAULTS["adsorbate_center_cm"])),
        "adsorbate_fwhm_cm": tk.StringVar(value=str(DEFAULTS["adsorbate_fwhm_cm"])),
        "adsorbate_strength": tk.StringVar(value=str(DEFAULTS["adsorbate_strength"])),
        "adsorbate_thickness_nm": tk.StringVar(value=str(DEFAULTS["adsorbate_thickness_nm"])),
        "adsorbate_n_background": tk.StringVar(value=str(DEFAULTS["adsorbate_n_background"])),
    }

    ttk.Label(controls, text="ATR-SEIRAS material overlay",
              font=("Segoe UI", 12, "bold")).grid(
                  row=0, column=0, columnspan=3, sticky="w", pady=(0, 8))
    row_index = 1

    def section(text):
        nonlocal row_index
        ttk.Label(controls, text=text, font=("Segoe UI", 10, "bold")).grid(
            row=row_index, column=0, columnspan=3, sticky="w", pady=(5, 2))
        row_index += 1

    def entry_row(label, key, unit="", width=13):
        nonlocal row_index
        ttk.Label(controls, text=label).grid(row=row_index, column=0, sticky="w")
        widget = ttk.Entry(controls, textvariable=variables[key], width=width)
        widget.grid(row=row_index, column=1, sticky="ew")
        if unit:
            ttk.Label(controls, text=unit).grid(row=row_index, column=2, sticky="w")
        row_index += 1
        return widget

    section("Optical geometry")
    entry_row("Observation wavenumber", "wavenumber_cm", "cm^-1")
    ttk.Label(controls, text="Prism").grid(row=row_index, column=0, sticky="w")
    prism_widget = ttk.Combobox(
        controls, textvariable=variables["prism"],
        values=("Si", "Ge", "ZnSe", "Custom n"),
        state="readonly", width=11)
    prism_widget.grid(row=row_index, column=1, sticky="ew")
    row_index += 1
    custom_prism_widget = entry_row("Custom prism n", "custom_prism_n")
    prism_state_label = ttk.Label(
        controls, text="Fixed prism dispersion is used", foreground="#777777")
    prism_state_label.grid(row=row_index, column=0, columnspan=3, sticky="w")
    row_index += 1
    entry_row("Solution n", "n_solution")
    entry_row("Solution k", "k_solution")
    entry_row("Angle min", "angle_min_deg", "deg")
    entry_row("Angle max", "angle_max_deg", "deg")
    entry_row("Beam divergence FWHM", "divergence_fwhm_deg", "deg")

    section("Au / Ag / Cu common-thickness overlay")
    ttk.Label(controls, text="Materials").grid(
        row=row_index, column=0, sticky="w")
    ttk.Label(controls, text="Au, Ag, Cu (fixed)", foreground="#006400").grid(
        row=row_index, column=1, columnspan=2, sticky="w")
    row_index += 1
    ttk.Label(controls, text="Optical constants").grid(
        row=row_index, column=0, sticky="w")
    ttk.Label(controls, text="Material-specific infrared Drude",
              foreground="#006400").grid(
                  row=row_index, column=1, columnspan=2, sticky="w")
    row_index += 1
    entry_row("Common thicknesses", "thicknesses_nm", "nm", width=18)
    ttk.Label(
        controls, text="Comma-separated; 1-4 values shared by all three metals",
        foreground="#555555").grid(
            row=row_index, column=0, columnspan=3, sticky="w")
    row_index += 1
    ttk.Label(controls, text="Film morphology").grid(
        row=row_index, column=0, sticky="w")
    morphology_widget = ttk.Combobox(
        controls, textvariable=variables["film_morphology"],
        values=MORPHOLOGY_CHOICES,
        state="readonly", width=36)
    morphology_widget.grid(row=row_index, column=1, columnspan=2, sticky="ew")
    row_index += 1
    ema_fill_widget = entry_row(
        "Isolated-film EMA fill", "ema_fill_fraction", "0-1")
    connected_fill_widget = entry_row(
        "Connected-film EMA fill", "connected_fill_fraction", "0-1")
    percolation_widget = entry_row(
        "Percolation midpoint", "percolation_thickness_nm", "nm")
    gap_gain_widget = entry_row(
        "Gap-field intensity gain", "gap_field_gain", "x")
    ema_state_label = ttk.Label(
        controls, text="Morphology controls not included",
        foreground="#777777")
    ema_state_label.grid(row=row_index, column=0, columnspan=3, sticky="w")
    row_index += 1
    section("Adsorbate layer")
    ttk.Checkbutton(controls, text="Include adsorbate layer",
                    variable=variables["include_adsorbate"]).grid(
                        row=row_index, column=0, columnspan=3, sticky="w")
    row_index += 1
    ttk.Label(controls, text="Species/model").grid(row=row_index, column=0, sticky="w")
    ads_widgets = []
    species_widget = ttk.Combobox(
        controls, textvariable=variables["adsorbate_name"],
        values=("CO oscillator (generic comparison)",),
        state="readonly", width=27)
    species_widget.grid(row=row_index, column=1, columnspan=2, sticky="ew")
    ads_widgets.append(species_widget)
    row_index += 1
    ads_widgets.append(entry_row("CO band centre", "adsorbate_center_cm", "cm^-1"))
    ads_widgets.append(entry_row("CO linewidth FWHM", "adsorbate_fwhm_cm", "cm^-1"))
    ads_widgets.append(entry_row("Lorentz strength", "adsorbate_strength"))
    ads_widgets.append(entry_row("Adsorbate thickness", "adsorbate_thickness_nm", "nm"))
    ads_widgets.append(entry_row("Adsorbate background n", "adsorbate_n_background"))
    ads_state_label = ttk.Label(controls, text="CO layer included in calculation",
                                foreground="#006400")
    ads_state_label.grid(row=row_index, column=0, columnspan=3, sticky="w", pady=(2, 5))
    row_index += 1

    ttk.Label(
        controls,
        text=("Every material uses the same nominal thickness grid.\n"
              "Colour identifies material; line style identifies thickness.\n"
              "Semi-quantitative gap gain reproduces trend only; it is not FEM."),
        foreground="#7A3E00", wraplength=340, justify="left").grid(
            row=row_index, column=0, columnspan=3, sticky="w", pady=(4, 7))
    row_index += 1

    notebook = ttk.Notebook(plot_frame)
    notebook.pack(fill="both", expand=True)
    optical_tab = ttk.Frame(notebook)
    coupling_tab = ttk.Frame(notebook)
    notebook.add(optical_tab, text="Optical response")
    notebook.add(coupling_tab, text="Film coupling / 60 deg")

    figure = Figure(figsize=(11.2, 7.6), dpi=100, constrained_layout=False)
    canvas = FigureCanvasTkAgg(figure, master=optical_tab)
    toolbar = NavigationToolbar2Tk(canvas, optical_tab, pack_toolbar=False)
    toolbar.update()
    toolbar.pack(side="bottom", fill="x")
    canvas.get_tk_widget().pack(side="top", fill="both", expand=True)

    coupling_figure = Figure(
        figsize=(11.2, 7.6), dpi=100, constrained_layout=False)
    coupling_canvas = FigureCanvasTkAgg(
        coupling_figure, master=coupling_tab)
    coupling_toolbar = NavigationToolbar2Tk(
        coupling_canvas, coupling_tab, pack_toolbar=False)
    coupling_toolbar.update()
    coupling_toolbar.pack(side="bottom", fill="x")
    coupling_canvas.get_tk_widget().pack(
        side="top", fill="both", expand=True)

    state = {"result": None}

    def update_prism_state(*_args):
        enabled = variables["prism"].get() == "Custom n"
        custom_prism_widget.configure(state="normal" if enabled else "disabled")
        if enabled:
            prism_state_label.configure(
                text="Custom constant n included in calculation",
                foreground="#006400")
        else:
            prism_state_label.configure(
                text="Fixed prism dispersion is used; custom n not included",
                foreground="#777777")

    def update_morphology_state(*_args):
        morphology = variables["film_morphology"].get()
        is_auto = morphology == MORPHOLOGY_THICKNESS_LINKED
        is_isolated = morphology == MORPHOLOGY_ISOLATED
        is_connected = morphology == MORPHOLOGY_CONNECTED
        ema_fill_widget.configure(
            state="normal" if (is_auto or is_isolated) else "disabled")
        connected_fill_widget.configure(
            state="normal" if (is_auto or is_connected) else "disabled")
        percolation_widget.configure(state="normal" if is_auto else "disabled")
        gap_gain_widget.configure(
            state="normal" if (is_auto or is_connected) else "disabled")
        if is_auto:
            ema_state_label.configure(
                text=("Each entered thickness maps from surface-CO to gap-CO "
                      "around the percolation midpoint"),
                foreground="#006400")
        elif is_isolated:
            ema_state_label.configure(
                text="Isolated EMA included; CO oscillator remains at surface",
                foreground="#006400")
        elif is_connected:
            ema_state_label.configure(
                text="Connected EMA included; CO oscillator is placed in voids",
                foreground="#006400")
        else:
            ema_state_label.configure(
                text="Smooth planar metal: EMA and gap correction not included",
                foreground="#777777")

    def update_adsorbate_state(*_args):
        enabled = bool(variables["include_adsorbate"].get())
        for widget in ads_widgets:
            if isinstance(widget, ttk.Combobox):
                widget.configure(state="readonly" if enabled else "disabled")
            else:
                widget.configure(state="normal" if enabled else "disabled")
        if enabled:
            ads_state_label.configure(
                text="CO amount is partitioned between surface and EMA voids",
                                      foreground="#006400")
        else:
            ads_state_label.configure(text="Adsorbate disabled: parameters not included",
                                      foreground="#777777")

    variables["include_adsorbate"].trace_add("write", update_adsorbate_state)
    variables["prism"].trace_add("write", update_prism_state)
    variables["film_morphology"].trace_add("write", update_morphology_state)
    update_adsorbate_state()
    update_prism_state()
    update_morphology_state()

    def read_params():
        prism_choice = variables["prism"].get()
        optical_choice = "Infrared Drude"
        morphology_choice = variables["film_morphology"].get()
        include_ads_choice = bool(variables["include_adsorbate"].get())
        return {
            "wavenumber_cm": float(variables["wavenumber_cm"].get()),
            "prism": prism_choice,
            "custom_prism_n": (
                float(variables["custom_prism_n"].get())
                if prism_choice == "Custom n" else DEFAULTS["custom_prism_n"]),
            "n_solution": float(variables["n_solution"].get()),
            "k_solution": float(variables["k_solution"].get()),
            "angle_min_deg": float(variables["angle_min_deg"].get()),
            "angle_max_deg": float(variables["angle_max_deg"].get()),
            "divergence_fwhm_deg": float(variables["divergence_fwhm_deg"].get()),
            "multi_metal_comparison": True,
            "metal": "Au",
            "thicknesses_nm": parse_thicknesses(variables["thicknesses_nm"].get()),
            "metal_optical_model": optical_choice,
            "custom_metal_n": DEFAULTS["custom_metal_n"],
            "custom_metal_k": DEFAULTS["custom_metal_k"],
            "film_morphology": morphology_choice,
            "ema_fill_fraction": float(
                variables["ema_fill_fraction"].get()),
            "connected_fill_fraction": float(
                variables["connected_fill_fraction"].get()),
            "percolation_thickness_nm": float(
                variables["percolation_thickness_nm"].get()),
            "gap_field_gain": float(variables["gap_field_gain"].get()),
            "include_adsorbate": include_ads_choice,
            "adsorbate_name": variables["adsorbate_name"].get(),
            "adsorbate_center_cm": (
                float(variables["adsorbate_center_cm"].get())
                if include_ads_choice else DEFAULTS["adsorbate_center_cm"]),
            "adsorbate_fwhm_cm": (
                float(variables["adsorbate_fwhm_cm"].get())
                if include_ads_choice else DEFAULTS["adsorbate_fwhm_cm"]),
            "adsorbate_strength": (
                float(variables["adsorbate_strength"].get())
                if include_ads_choice else DEFAULTS["adsorbate_strength"]),
            "adsorbate_thickness_nm": (
                float(variables["adsorbate_thickness_nm"].get())
                if include_ads_choice else DEFAULTS["adsorbate_thickness_nm"]),
            "adsorbate_n_background": (
                float(variables["adsorbate_n_background"].get())
                if include_ads_choice else DEFAULTS["adsorbate_n_background"]),
        }

    def calculate():
        try:
            params = read_params()
            result = simulate(params)
            build_figure(result, figure)
            build_coupling_figure(result, coupling_figure)
            canvas.draw_idle()
            coupling_canvas.draw_idle()
            summary.delete("1.0", "end")
            summary.insert("1.0", result_summary(result))
            state["result"] = result
        except Exception as exc:
            messagebox.showerror("Calculation error", str(exc))

    def save_signal_figure():
        if state["result"] is None:
            calculate()
        if state["result"] is None:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".png",
            initialfile="AuAgCu_common_thickness_overlay.png",
            filetypes=(("PNG (300 dpi)", "*.png"),
                       ("PDF (vector)", "*.pdf"),
                       ("All files", "*.*")))
        if path:
            export_figure = build_signal_export_figure(state["result"])
            export_figure.savefig(path, dpi=300, facecolor="white")

    def save_xlsx():
        if state["result"] is None:
            calculate()
        if state["result"] is None:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile="AuAgCu_angularly_averaged_signal.xlsx",
            filetypes=(("Excel", "*.xlsx"), ("All files", "*.*")))
        if path:
            export_xlsx(state["result"], path)

    button_bar = ttk.Frame(controls)
    button_bar.grid(row=row_index, column=0, columnspan=3, sticky="ew")
    ttk.Button(button_bar, text="Calculate", command=calculate).pack(side="left", padx=(0, 4))
    ttk.Button(button_bar, text="Save signal figure",
               command=save_signal_figure).pack(side="left", padx=4)
    ttk.Button(button_bar, text="Export averaged signal Excel",
               command=save_xlsx).pack(side="left", padx=4)
    ttk.Button(button_bar, text="Self-test",
               command=lambda: messagebox.showinfo(
                   "Self-test", "%d passed, %d failed" % run_selftests(False))).pack(
                       side="left", padx=4)
    row_index += 1

    summary = ScrolledText(controls, width=51, height=8, font=("Consolas", 8))
    summary.grid(row=row_index, column=0, columnspan=3, sticky="nsew", pady=(8, 0))
    controls.rowconfigure(row_index, weight=1)

    calculate()
    root.mainloop()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--selftest", action="store_true",
                        help="run numerical regression tests and exit")
    args = parser.parse_args(argv)

    if args.selftest:
        _, failed = run_selftests(True)
        return 1 if failed else 0
    launch_gui()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
